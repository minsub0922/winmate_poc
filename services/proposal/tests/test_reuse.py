"""기존 제안서 활용(PR1C → PRU2 → PRU2F → PRU3A|PRU3B → PRU4|PRU4B → PRU5) — AC-028~030 · AC-180~195."""
from __future__ import annotations

import io
import json
from typing import Any

from winmate_common.jobs import jobs

from test_sections_generate import make, standard_with_mi

# B 베이커리(다른 고객) 외부 PPTX — 표지 · 문제 제기 · 시장 2 · 과제 · 가치 2 · 솔루션 · 동선(역할 후보 2) · 사례 · 스펙 · 견적 · 일정
BAKERY = [
    ("B 베이커리 매장 사이니지 제안", ["B 베이커리 · 2024.09", "삼성전자 B2B 제안"]),
    ("오프라인 매장, 디지털 전환의 기로", ["매장 운영의 변화가 빨라지고 있어요", "다음 장의 시장 수치가 근거예요"]),
    ("국내 QSR 디지털 메뉴보드 도입 추이", ["2024년 기준 도입률 38%", "프랜차이즈 본사 주도 교체가 늘고 있어요"]),
    ("경쟁 브랜드 도입 현황", ["경쟁 5개 브랜드가 메뉴보드를 도입했어요"]),
    ("B 베이커리 매장 운영 과제 3가지", ["매장마다 메뉴 교체가 늦어요", "점주 인터뷰에서 나온 불편을 정리했어요",
                                      "B 베이커리 매출 1,200억 원 중 판촉비 비중이 커요"]),
    ("가치 제안: 본사 통제 · 운영 부담 ↓", ["본사가 콘텐츠를 한 번에 배포해 매장 운영 부담을 줄입니다", "매장별 운영 부담 ↓ · 본사 통제 ↑"]),
    ("기대 효과 Before / After", ["교체 소요 5일 → 당일", "가격 오류 30% 감소"]),
    ("솔루션 구성도 MagicINFO 8", ["MagicINFO 8 서버로 전 매장 일괄 배포", "QM55B 메뉴보드 2대 · 카운터 상단"]),
    ("매장 동선별 사이니지 배치", ["입구 · 카운터 · 대기 공간 동선"]),
    ("도입 사례 · C 베이커리", ["C 베이커리 120개 매장 도입 레퍼런스"]),
    ("제품 스펙 비교", ["QM55B · QB65B 사양 비교"]),
    ("견적 · 단가", ["QM55B 단가 1,450,000원 · VAT 별도", "할인 10% 적용 시 총액"]),
    ("추진 일정", ["2024.10 계약 · 2024.11 설치 · 2024.12 오픈"]),
]
SOURCE_SENTENCE = "본사가 콘텐츠를 한 번에 배포해 매장 운영 부담을 줄입니다"


def bakery_pptx() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    for title, lines in BAKERY:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = lines[0]
        for ln in lines[1:]:
            body.add_paragraph().text = ln
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def upload_bakery() -> str:
    from winmate_common.platform import save_file
    meta = await save_file("B베이커리_매장사이니지_제안.pptx", bakery_pptx(),
                           "application/vnd.openxmlformats-officedocument.presentationml.presentation", source="upload", confidential=True)
    return meta["id"]


async def reuse_view(client: Any, pid: str) -> dict[str, Any]:
    r = await client.get(f"/v1/proposals/{pid}/reuse")
    assert r.status_code == 200, r.text
    return r.json()


def reuse_calls(ctx: Any) -> list[dict[str, Any]]:
    return [c for c in ctx.stubs.ai_calls if str(c.get("task") or "").startswith("pr.reuse")]


async def start_borrow(client: Any, ctx: Any) -> tuple[str, str]:
    """A 커피(이번) ← B 베이커리 PPTX(다른 고객) 분석까지."""
    p = await make(client, title="전국 매장 디지털 메뉴보드 전환", customer={"name": "A 커피 프랜차이즈", "scale_text": "320개 매장"},
                   rq_ref={"rq_id": "rq_acoffee", "version": 1})
    pid = p["id"]
    fid = await upload_bakery()
    r = await client.post(f"/v1/proposals/{pid}/reuse", json={"sources": [{"kind": "file", "file_id": fid}], "mode_pref": "auto"})
    assert r.status_code == 202, r.text
    acc = r.json()
    assert acc["kind"] == "proposal.reuse" and acc["reuse_id"].startswith("pru_")
    v = await reuse_view(client, pid)
    assert v["status"] == "analyzing" and v["status_label"] == "기준별 분석 진행 중" and v["can_confirm"] is False     # AC-029
    await ctx.run_jobs()
    j = await jobs().get(acc["job_id"])
    assert j.status == "awaiting_input", (j.status, j.error)
    assert (j.input_request or {}).get("kind") == "reuse_confirm_analysis"
    return pid, fid


async def test_borrow_flow_analysis_plan_apply(client: Any, ctx: Any) -> None:
    pid, fid = await start_borrow(client, ctx)
    v = await reuse_view(client, pid)
    assert v["status"] == "awaiting_confirm", v
    assert v["sources"][0]["format"] == "pptx" and v["sources"][0]["meta_label"].startswith("PPTX · 13장")
    assert [ph["status"] for ph in v["phases"]] == ["done", "done", "done"]
    pages = {pg["no"]: pg for pg in v["pages"]}
    assert pages[1]["flow_role"] == "표지"
    assert pages[12]["excluded"] and pages[13]["excluded"]                                         # 견적 · 일정 자동 제외
    assert pages[9]["role_state"] == "need" and pages[9]["role_candidates"] == ["공간 시나리오", "솔루션 구성"]
    assert v["excluded_page_count"] == 2 and v["band_label"].startswith("원본 13장 · 외부 파일")
    crit = {c["no"]: c for c in v["criteria"]}
    assert crit[9]["state"] == "ok" and crit[9]["must"] and "매출 1건" in crit[9]["summary"]          # 비복제 · 매출 줄
    assert crit[2]["state"] == "need" and crit[4]["state"] == "need"
    assert "Why Samsung 고리 없음" in crit[2]["summary"]
    assert "단종 1종(QM55B → QM55C)" in crit[5]["summary"] and "MagicINFO 9" in crit[5]["summary"]
    assert v["recommendation"] == {"mode": "borrow", "reason": "다른 고객 제안서라 설득 구조만 가져와요", "badge": "추천 · 다른 고객"}
    assert v["context"]["source_customer"] == "B 베이커리" and v["context"]["same_customer"] is False
    assert any(s["dashed"] and s["name"] == "Why Samsung · 경쟁 비교" for s in v["flow"]["steps"])
    assert v["must_done"] == 1 and v["footer_label"].endswith("필수 확인 1 / 3 마침 · 2곳 남음")
    # 비복제 쪽 · 줄은 이후 LLM 입력에 없다(AC-181) · 모든 reuse 호출은 기밀(AC-195)
    calls = reuse_calls(ctx)
    assert calls and all(c.get("confidential") is True for c in calls)
    later = [json.dumps(c, ensure_ascii=False) for c in calls if c.get("task") != "pr.reuse_noncopy"]
    for secret in ("1,450,000", "VAT 별도", "매출 1,200억", "2024.11 설치"):
        assert not any(secret in s for s in later), secret
    # 필수 확인 전 → 409(AC-182)
    r = await client.post(f"/v1/proposals/{pid}/reuse:confirm-analysis")
    assert r.status_code == 409 and r.json()["error"]["code"] == "MUST_CONFIRM_PENDING"
    assert r.json()["error"]["details"]["missing"] == [2, 4]
    # PRU2F — p.10(여기선 p.9) 역할을 고치면 고정 · 재분석 뒤에도(AC-184)
    r = await client.put(f"/v1/proposals/{pid}/reuse/pages/9/role", json={"role": "솔루션 구성"})
    assert r.status_code == 200, r.text
    pg9 = next(pg for pg in r.json()["pages"] if pg["no"] == 9)
    assert pg9["role_state"] == "edited" and pg9["locked"] and pg9["role_state_label"] == "수정함"
    r = await client.put(f"/v1/proposals/{pid}/reuse/pages/12/role", json={"role": "가치 제안"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "VERDICT_LOCKED"
    r = await client.post(f"/v1/proposals/{pid}/reuse:reanalyze")
    assert r.status_code == 200 and r.json()["status"] == "analyzing"
    await ctx.run_jobs()
    v = await reuse_view(client, pid)
    assert v["status"] == "awaiting_confirm"
    pg9 = next(pg for pg in v["pages"] if pg["no"] == 9)
    assert pg9["flow_role"] == "솔루션 구성" and pg9["role_state"] == "edited"
    # 기준 상세(PRU2F)
    cd = (await client.get(f"/v1/proposals/{pid}/reuse/criteria/2")).json()
    assert cd["header_label"] == f"신뢰도 {cd['confidence']}% · 다음 단계로 가려면 확인이 필요해요"
    assert cd["footer_label"].startswith("논리 흐름 확인 · 13장 중 태그 확인") and len(cd["chips"]) == 9 and cd["flow"]["memos"]
    assert (await client.get(f"/v1/proposals/{pid}/reuse/criteria/12")).status_code == 404
    # 「이 흐름으로 확인」(AC-185) · 기준 4 확인
    r = await client.put(f"/v1/proposals/{pid}/reuse/criteria/2", json={"state": "ok"})
    assert r.status_code == 200 and all(pg["role_state"] != "need" for pg in r.json()["pages"])
    r = await client.put(f"/v1/proposals/{pid}/reuse/criteria/4", json={"state": "ok"})
    v = r.json()
    assert v["must_done"] == 3 and v["can_confirm"] is True
    # 확인 완료 → 추천(흐름 차용) 계획
    r = await client.post(f"/v1/proposals/{pid}/reuse:confirm-analysis")
    assert r.status_code == 200, r.text
    assert r.json()["route"].endswith("/reuse/plan?mode=borrow")
    v = await reuse_view(client, pid)
    pb = v["plan_borrow"]
    assert pb and pb["footer_label"].endswith("원본 내용 0줄 사용 · 1 / 6")
    kinds = {r_["src_step"]: r_["kind"] for r_ in pb["rows"]}
    assert kinds["견적 · 일정"] == "drop" and kinds["Why Samsung · 경쟁 비교"] == "new"
    assert pb["counts"]["drop"] == 2 and pb["sheets"] == pb["counts"]["flow"] + pb["counts"]["new"]
    assert [t["label"] for t in pb["not_take"]][:2] == ["고객 정보", "매장 사진"]
    await ctx.run_jobs()
    v = await reuse_view(client, pid)
    assert v["status"] == "awaiting_plan_confirm"
    # 방식 바꾸기 → 개선 계획(비복제 행 판정 잠금, AC-187)
    v = (await client.put(f"/v1/proposals/{pid}/reuse/mode", json={"mode": "improve"})).json()
    rows = [r_ for g in v["plan_improve"]["groups"] for r_ in g["rows"]]
    nc = next(r_ for r_ in rows if r_["noncopy"])
    assert nc["verdict"] == "drop" and nc["locked"]
    r = await client.put(f"/v1/proposals/{pid}/reuse/plan/rows/{nc['row_id']}", json={"verdict": "keep"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "VERDICT_LOCKED"
    v = (await client.put(f"/v1/proposals/{pid}/reuse/mode", json={"mode": "borrow"})).json()
    assert v["mode"] == "borrow"
    # 확정 → 흐름 가이드
    r = await client.post(f"/v1/proposals/{pid}/reuse/plan:confirm", json={"then": "sections"})
    assert r.status_code == 200, r.text
    assert r.json()["route"].endswith("?view=guide")
    await ctx.run_jobs()
    v = await reuse_view(client, pid)
    assert v["status"] == "applied", v.get("error")
    p = (await client.get(f"/v1/proposals/{pid}")).json()
    assert p["type"] == "standard" and p["derived_from"]["kind"] == "file" and p["derived_from"]["label"] == "B베이커리_매장사이니지_제안.pptx에서 파생"
    assert p["reuse"]["mode"] == "borrow" and p["stage"] == "sections"
    # 원본 대조 보기는 API 에서도 거부(AC-190)
    r = await client.get(f"/v1/proposals/{pid}/sections/vp/reuse-view", params={"view": "compare"})
    assert r.status_code == 403 and r.json()["error"]["code"] == "BORROW_MODE_CONTENT_HIDDEN"
    sv = (await client.get(f"/v1/proposals/{pid}/sections/vp")).json()
    assert sv["reuse_view"] == "guide"
    gv = (await client.get(f"/v1/proposals/{pid}/sections/vp/reuse-view")).json()
    assert gv["view"] == "guide" and gv["compare_disabled_reason"] == "흐름 차용 모드에선 원본 내용을 보여주지 않아요"
    assert gv["source_sheet"] is None and gv["guide"] and gv["guide"][0]["status"] == "writing"
    assert all("내용 안 가져옴" in g["from_source"] or "원본에" in g["from_source"] or "이번 유형" in g["from_source"] for g in gv["guide"])
    # 흐름 차용 Value Props → 플레이스홀더 [0]일 · [00]% · [0]명 확정 필요 등록, 원본 수치(30% · 5일)는 쓰지 않음(AC-191)
    items = (await client.get(f"/v1/proposals/{pid}/confirm-items", params={"status": "open"})).json()["items"]
    marks = {it["text"]["mark"] for it in items if it["section_key"] == "vp"}
    assert {"[0]일", "[00]%", "[0]명"} <= marks, marks
    shs = (await client.get(f"/v1/proposals/{pid}")).json()["sheets"]
    blob = ""
    for s in shs:
        blob += json.dumps((await client.get(f"/v1/proposals/{pid}/sheets/{s['id']}")).json()["display"], ensure_ascii=False)
    for secret in ("B 베이커리", "30% 감소", "1,450,000", SOURCE_SENTENCE, "점주 인터뷰"):
        assert secret not in blob, secret
    # PRU5(생성 전에도 계산) — 흐름 차용은 원본 내용을 쓰지 않으므로 전 시트가 「원본 대비 변경」
    sm_ = (await client.get(f"/v1/proposals/{pid}/reuse/summary")).json()
    assert [c_["label"] for c_ in sm_["counts"]] == ["유지", "갱신", "재작성", "신규", "제외"]
    assert sm_["counts"][4]["n"] == 2 and any(m["section"] == "견적 · 일정" for m in sm_["mapping"])
    assert sm_["traces"][0]["title"].startswith("새 제안서 각 시트 노트에 '흐름 차용")
    # 흐름 차용 초안 입력에 원본 문장이 없다(AC-190)
    drafts = [json.dumps(c, ensure_ascii=False) for c in ctx.stubs.ai_calls if c.get("task") == "pr.section_draft"
              and "reuse_borrow" in json.dumps(c, ensure_ascii=False)]
    assert drafts
    src_lines = [ln for _t, lines in BAKERY for ln in lines if len(ln) >= 8] + [t for t, _l in BAKERY[1:]]
    for d in drafts:
        assert all(ln not in d for ln in src_lines)
        assert '"confidential": true' in d


async def test_borrow_leak_check_rewrites_line(client: Any, ctx: Any) -> None:
    """모델이 원본 문장을 그대로 내면 저장 전에 지우고 다시 쓴다(AC-190 · §10.12)."""
    ctx.stubs.ai_overrides["pr.section_draft"] = {"json": {"sheets": [
        {"sheet_key": "VP", "title": "가치 제안 — 본사 통제", "points": [{"title": "본사 통제", "body": SOURCE_SENTENCE},
                                                                       {"title": "교체 주기", "body": "교체 주기를 당일로"}]}]}}
    pid, _fid = await start_borrow(client, ctx)
    for no in (2, 4):
        await client.put(f"/v1/proposals/{pid}/reuse/criteria/{no}", json={"state": "ok"})
    assert (await client.post(f"/v1/proposals/{pid}/reuse:confirm-analysis")).status_code == 200
    await ctx.run_jobs()
    assert (await client.post(f"/v1/proposals/{pid}/reuse/plan:confirm", json={"then": "sections"})).status_code == 200
    await ctx.run_jobs()
    assert (await reuse_view(client, pid))["status"] == "applied"
    vp = next(s for s in (await client.get(f"/v1/proposals/{pid}")).json()["sheets"] if s["role"] == "VP")
    disp = json.dumps((await client.get(f"/v1/proposals/{pid}/sheets/{vp['id']}")).json()["display"], ensure_ascii=False)
    assert SOURCE_SENTENCE not in disp
    assert "이번 고객 요구사항 기준으로 새로 쓴 문장이에요" in disp
    rew = [c for c in ctx.stubs.ai_calls if c.get("task") == "pr.reuse_line_rewrite"]
    assert rew and all(c.get("confidential") is True and SOURCE_SENTENCE not in json.dumps(c, ensure_ascii=False) for c in rew)


async def test_reuse_confidential_blocked_fails_job(client: Any, ctx: Any) -> None:
    """기밀 차단(403) → 잡 failed(POLICY_CONFIDENTIAL) · 화면 문구(AC-195), 다른 모델로 바꾸지 않음."""
    ctx.stubs.ai_overrides["pr.reuse_flow"] = {"error": {"status": 403, "code": "POLICY_CONFIDENTIAL", "message": "blocked"}}
    p = await make(client, title="전국 매장 디지털 메뉴보드 전환", customer={"name": "A 커피 프랜차이즈"})
    pid = p["id"]
    fid = await upload_bakery()
    r = await client.post(f"/v1/proposals/{pid}/reuse", json={"sources": [{"kind": "file", "file_id": fid}]})
    job = r.json()["job_id"]
    await ctx.run_jobs()
    j = await jobs().get(job)
    assert j.status == "failed" and j.error["code"] == "POLICY_CONFIDENTIAL"
    v = await reuse_view(client, pid)
    assert v["status"] == "failed" and v["error"]["message"] == "기밀 자료라 지금 설정된 모델로는 분석할 수 없어요"
    assert all(c.get("confidential") is True for c in reuse_calls(ctx))


async def test_reuse_start_validation(client: Any, ctx: Any) -> None:
    """PR1C — 50MB 초과 413 · PPTX/PDF 아님 415 · 원본 2개는 분석 1건(AC-028) · 목록 추천 ≤ 2(AC-030)."""
    from winmate_common.platform import save_file
    p = await make(client, title="t", customer={"name": "A 커피 프랜차이즈"})
    pid = p["id"]
    docx = await save_file("요구.docx", b"PK\x03\x04 not really", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           source="upload", confidential=True)
    r = await client.post(f"/v1/proposals/{pid}/reuse", json={"sources": [{"kind": "file", "file_id": docx["id"]}]})
    assert r.status_code == 415 and r.json()["error"]["code"] == "UNSUPPORTED_FILE_TYPE"
    from winmate_proposal import clients
    real = clients.file_meta

    async def big(fid: str) -> dict[str, Any]:
        m = dict(await real(fid) or {})
        m["size"] = 51 * 1024 * 1024
        return m
    clients.file_meta = big  # type: ignore[assignment]
    try:
        f1 = await upload_bakery()
        r = await client.post(f"/v1/proposals/{pid}/reuse", json={"sources": [{"kind": "file", "file_id": f1}]})
        assert r.status_code == 413 and r.json()["error"]["code"] == "FILE_TOO_LARGE"
    finally:
        clients.file_meta = real  # type: ignore[assignment]
    f1, f2 = await upload_bakery(), await upload_bakery()
    r = await client.post(f"/v1/proposals/{pid}/reuse", json={"sources": [{"kind": "file", "file_id": f1}, {"kind": "file", "file_id": f2}]})
    assert r.status_code == 202
    rid = r.json()["reuse_id"]
    v = await reuse_view(client, pid)
    assert v["id"] == rid and len(v["sources"]) == 2
    r = await client.post(f"/v1/proposals/{pid}/reuse", json={"sources": [{"kind": "proposal", "proposal_id": pid}]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "SOURCE_IS_SELF"
    # 원본 빼기(replace) · PR1C 라디오 저장만(mode_pref) · 마지막 원본 빼기(DELETE)
    r = await client.post(f"/v1/proposals/{pid}/reuse", json={"sources": [{"kind": "file", "file_id": f2}], "replace": True})
    assert r.status_code == 202 and r.json()["reuse_id"] == rid
    v = await reuse_view(client, pid)
    assert [s_["file_id"] for s_ in v["sources"]] == [f2]
    v = (await client.put(f"/v1/proposals/{pid}/reuse/mode", json={"mode_pref": "borrow"})).json()
    assert v["mode_pref"] == "borrow" and v["mode"] == "borrow"
    assert (await client.put(f"/v1/proposals/{pid}/reuse/mode", json={})).status_code == 422
    assert (await client.delete(f"/v1/proposals/{pid}/reuse")).status_code == 200
    r = await client.get(f"/v1/proposals/{pid}/reuse")
    assert r.status_code == 404 and r.json()["error"]["code"] == "REUSE_NOT_FOUND"
    c = (await client.get(f"/v1/proposals/{pid}/reuse/candidates")).json()
    assert len(c["items"]) <= 2


async def test_improve_same_customer_winmate(client: Any, ctx: Any) -> None:
    """A 커피 2025(같은 고객 Winmate) → 개선 · 수정(AC-183 · 186~189 · 192 · 193)."""
    # 원본: 표준 제안서 v1 → 시트 다듬기(QM55B · 280개 매장 · 단가 줄) → v2
    src = await standard_with_mi(client)
    await client.post(f"/v1/proposals/{src}/sections/mi:fill", json={})
    await ctx.run_jobs()
    r = await client.post(f"/v1/proposals/{src}:generate", json={"scope": "all", "infer_empty": True})
    assert r.status_code == 202
    await ctx.run_jobs()
    sp = (await client.get(f"/v1/proposals/{src}")).json()
    assert sp["version"] == 1
    sm = next(s for s in sp["sheets"] if s["section_key"] == "spaceProducts")
    cb = next(s for s in sp["sheets"] if s["role"] == "CB")
    r = await client.patch(f"/v1/proposals/{src}/sheets/{sm['id']}", json={"ops": [
        {"op": "set", "path": "/title", "value": "카운터 메뉴보드 QM55B 2대 · 바닥형 스탠드"},
        {"op": "set", "path": "/subtitle", "value": "QM55B 단가 1,450,000원 · VAT 별도"}]})
    assert r.status_code == 200, r.text
    r = await client.patch(f"/v1/proposals/{src}/sheets/{cb['id']}", json={"ops": [{"op": "set", "path": "/subtitle", "value": "전국 280개 매장 · 직영 · 가맹"}]})
    assert r.status_code == 200, r.text
    assert (await client.post(f"/v1/proposals/{src}/versions", json={"desc": "제출본"})).status_code == 201
    from winmate_proposal import versions as VER
    src_hash = (await VER.get_version(src, 2))["hash"]
    # 이번: 같은 고객 · 복제해서 시작(고객 이어받음 → 정의서도 원본 것으로)
    r = await client.post("/v1/proposals", json={"start_mode": "reuse", "source_proposal_id": src})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    assert r.json()["customer"]["name"] == "A 커피 프랜차이즈"
    c = (await client.get(f"/v1/proposals/{pid}/reuse/candidates")).json()
    assert c["items"][0]["proposal_id"] == src and c["items"][0]["why"] == "같은 고객" and c["label"].startswith("내 제안서 목록에서는")
    r = await client.post(f"/v1/proposals/{pid}/reuse", json={"sources": [{"kind": "proposal", "proposal_id": src}], "new_title": "2026 메뉴보드 확대"})
    assert r.status_code == 202, r.text
    await ctx.run_jobs()
    v = await reuse_view(client, pid)
    assert v["status"] == "awaiting_confirm", v.get("error")
    assert v["band_label"].startswith(f"원본 {len(v['pages'])}장 · Winmate 제안서 · v2")
    assert v["context"]["same_customer"] is True
    rec = v["recommendation"]
    assert rec["mode"] == "improve" and rec["badge"].startswith("추천 · 같은 고객 · 겹침 ")                     # AC-183
    assert (await client.get(f"/v1/proposals/{pid}")).json()["title"] == "2026 메뉴보드 확대"
    calls = reuse_calls(ctx)
    assert calls and all(c_.get("confidential") is True for c_ in calls)
    assert not any("1,450,000" in json.dumps(c_, ensure_ascii=False) for c_ in calls if c_.get("task") != "pr.reuse_noncopy")
    for no in (2, 4):
        await client.put(f"/v1/proposals/{pid}/reuse/criteria/{no}", json={"state": "ok"})
    r = await client.post(f"/v1/proposals/{pid}/reuse:confirm-analysis")
    assert r.status_code == 200 and r.json()["route"].endswith("/reuse/plan?mode=improve")
    v = await reuse_view(client, pid)
    pi = v["plan_improve"]
    rows = [r_ for g in pi["groups"] for r_ in g["rows"]]
    notes = [r_["note"] for r_ in rows]
    assert any(n.startswith("QM55B 단종 → QM55C") for n in notes), notes                                       # AC-193
    assert any(n.startswith("매장 280 → 320") for n in notes), notes
    assert pi["footer_label"].startswith(f"활용 계획 · 원본 {pi['source_total']}장 → 새 제안서 ")
    assert set(pi["totals"]) == {"keep", "update", "rewrite", "new", "drop"}
    # 판정 바꾸기는 고정 → 재분석 뒤에도 유지(AC-187)
    keep_row = next(r_ for r_ in rows if r_["verdict"] == "keep" and not r_["noncopy"])
    v = (await client.put(f"/v1/proposals/{pid}/reuse/plan/rows/{keep_row['row_id']}", json={"verdict": "rewrite"})).json()
    row2 = next(r_ for g in v["plan_improve"]["groups"] for r_ in g["rows"] if r_["row_id"] == keep_row["row_id"])
    assert row2["verdict"] == "rewrite" and row2["locked"]
    r = await client.put(f"/v1/proposals/{pid}/reuse/plan/rows/{keep_row['row_id']}", json={"verdict": "new"})
    assert r.status_code == 422
    await ctx.run_jobs()
    assert (await reuse_view(client, pid))["status"] == "awaiting_plan_confirm"
    r = await client.post(f"/v1/proposals/{pid}/reuse:reanalyze")
    assert r.status_code == 200
    await ctx.run_jobs()
    v = await reuse_view(client, pid)
    assert v["status"] == "awaiting_confirm"
    assert (await client.post(f"/v1/proposals/{pid}/reuse:confirm-analysis")).status_code == 200
    v = await reuse_view(client, pid)
    row2 = next(r_ for g in v["plan_improve"]["groups"] for r_ in g["rows"] if r_["row_id"] == keep_row["row_id"])
    assert row2["verdict"] == "rewrite" and row2["locked"]
    await ctx.run_jobs()
    # 확정 → 원본 대조(AC-188)
    r = await client.post(f"/v1/proposals/{pid}/reuse/plan:confirm", json={"then": "sections"})
    assert r.status_code == 200 and "?view=compare" in r.json()["route"], r.text
    await ctx.run_jobs()
    v = await reuse_view(client, pid)
    assert v["status"] == "applied", v.get("error")
    p = (await client.get(f"/v1/proposals/{pid}")).json()
    assert p["type"] == "standard" and p["derived_from"]["proposal_id"] == src and p["derived_from"]["label"].endswith("v2에서 파생")
    assert p["reuse"]["mode"] == "improve" and p["stage"] == "sections"
    sheets = {s["id"]: (await client.get(f"/v1/proposals/{pid}/sheets/{s['id']}")).json() for s in p["sheets"]}
    reused = [s for s in sheets.values() if (s.get("reuse") or {}).get("source_page")]
    assert reused and all(s["reuse"]["verdict"] in ("keep", "update", "rewrite") for s in reused)
    assert all(s["content"]["notes"].startswith(f"원본 p.{s['reuse']['source_page']} 기반 · ") for s in reused), [s["content"]["notes"] for s in reused]
    counter = next(s for s in sheets.values() if "QM55" in json.dumps(s["display"], ensure_ascii=False))
    disp = json.dumps(counter["display"], ensure_ascii=False)
    assert "QM55C" in disp and "QM55B" not in disp                                                              # AC-193
    blob = json.dumps([s["display"] for s in sheets.values()], ensure_ascii=False)
    assert "1,450,000" not in blob and "VAT 별도" not in blob                                                    # 비복제 줄(V10)
    assert "280개 매장" not in blob
    items = (await client.get(f"/v1/proposals/{pid}/confirm-items", params={"status": "open"})).json()["items"]
    tags = {it["tag"] for it in items if it["origin"] == "reuse" and it["category"] == "review"}
    assert {"단종 치환", "값 전파"} <= tags, tags
    # PRU4 원본 대조 보기
    key = counter["section_key"]
    rv = (await client.get(f"/v1/proposals/{pid}/sections/{key}/reuse-view", params={"sheet_id": counter["id"]})).json()
    assert rv["view"] == "compare" and rv["mode"] == "improve" and rv["source_sheet"]
    assert any(ln["mark"] == "drop" and ln["badge"] == "비복제" for ln in rv["source_sheet"]["lines"])
    assert rv["new_sheet"]["lines"] and rv["footer_label"].startswith("원본 대조 작성 · 바뀐 줄 ")
    assert any(s_["selected"] for s_ in rv["section_sheets"])
    # 원본 줄 전부 가져오기 — 비복제 줄은 가져오지 않는다(AC-189)
    r = await client.post(f"/v1/proposals/{pid}/sheets/{counter['id']}/reuse:pull-lines", json={"all": True})
    assert r.status_code == 200, r.text
    assert "1,450,000" not in json.dumps(r.json()["display"], ensure_ascii=False)
    nc_line = next(ln["id"] for ln in rv["source_sheet"]["lines"] if ln["badge"] == "비복제")
    r = await client.post(f"/v1/proposals/{pid}/sheets/{counter['id']}/reuse:pull-lines", json={"line_ids": [nc_line]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "LINE_NOT_AVAILABLE"
    # PPTX 생성 → PRU5(AC-192)
    r = await client.post(f"/v1/proposals/{pid}:generate", json={"scope": "all", "infer_empty": True})
    assert r.status_code == 202
    await ctx.run_jobs()
    vv = (await client.get(f"/v1/proposals/{pid}/versions")).json()
    assert vv["versions"][0]["derived_label"].endswith("v2에서 파생")
    sm_ = (await client.get(f"/v1/proposals/{pid}/reuse/summary")).json()
    assert [c_["label"] for c_ in sm_["counts"]] == ["유지", "갱신", "재작성", "신규", "제외"]
    assert sm_["mapping"][0]["section"] == "표지 · 목차 · 간지" and len(sm_["traces"]) == 3
    assert sm_["traces"][2]["title"] == "원본 제안서는 바뀌지 않음"
    assert sm_["review_total"] >= 2 and sm_["review_items"]
    assert sm_["footer_label"].startswith("완료 · ") and sm_["file_label"].endswith("원본 제안서는 바뀌지 않았어요.")
    assert (await VER.get_version(src, 2))["hash"] == src_hash                                                   # 원본 불변
    res = (await client.get(f"/v1/proposals/{pid}/result")).json()
    assert res["derived"] is True and res["summary_route"].endswith("/reuse/summary")
