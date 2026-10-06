"""DOCX — python-docx 로 보고서형 문서를 만든다.

문서: {title, subtitle?, meta?: [str] | str, sections: [{heading, level?, paragraphs?, bullets?, numbered?,
        table?: {columns, rows, col_widths?, highlight_rows?, caption?}, images?: [file_id | {file_id, width_cm?, caption?}],
        captions?: [str], sources?: [{label, url?}], page_break?}], page_size?: A4|Letter, orientation?}
`[확인 필요]` 같은 자리표시는 노란 형광 + 굵게(영어는 [TBD]).
"""
from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from .pptx_render import Assets, DictAssets
from .theme import FONT, Theme, rgb
from .values import (
    as_bullets,
    as_list,
    as_sources,
    as_table,
    cell_text,
    col_label,
    image_ref,
    label,
    localize,
    split_markers,
    to_text,
)

PAGE = {"A4": (Cm(21.0), Cm(29.7)), "LETTER": (Cm(21.59), Cm(27.94))}


class DocxError(ValueError):
    pass


def _rgb(hex_: str) -> RGBColor:
    r, g, b = rgb(hex_)
    return RGBColor(r, g, b)


def _font(run: Any, *, size: float | None = None, bold: bool | None = None, italic: bool | None = None,
          color: str | None = None) -> None:
    f = run.font
    f.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = _rgb(color)


def add_runs(par: Any, text: str, *, size: float | None = None, bold: bool = False, italic: bool = False,
             color: str | None = None) -> None:
    for seg, marked in split_markers(text):
        if not seg:
            continue
        r = par.add_run(seg)
        _font(r, size=size, bold=bold or marked, italic=italic, color="#9a3412" if marked else color)
        if marked:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _field(par: Any, instr: str) -> None:
    """PAGE · NUMPAGES 같은 필드."""
    r = par.add_run()
    _font(r, size=8.5, color="#8a91a0")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._r.append(fld_begin)
    r._r.append(instr_el)
    r._r.append(fld_end)


def _shade(cell: Any, hex_: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_.lstrip("#").upper())
    tc_pr.append(shd)


def _png_or_jpeg(data: bytes) -> bytes | None:
    try:
        im = Image.open(io.BytesIO(data))
        fmt = (im.format or "").upper()
        if fmt in ("PNG", "JPEG", "GIF", "BMP"):
            if max(im.size) <= 2400:
                return data
        im.thumbnail((2400, 2400))
        buf = io.BytesIO()
        if im.mode in ("RGBA", "LA", "P"):
            im.convert("RGBA").save(buf, "PNG")
        else:
            im.convert("RGB").save(buf, "JPEG", quality=88)
        return buf.getvalue()
    except Exception:  # noqa: BLE001
        return None


def build_docx(document: dict[str, Any], *, design: dict[str, Any] | None = None, lang: str = "ko",
               assets: Assets | None = None, confidential: bool = False) -> tuple[bytes, list[str]]:
    sections = document.get("sections")
    if not isinstance(sections, list):
        raise DocxError("sections 가 필요합니다")
    theme = Theme.from_design(design)
    assets = assets or DictAssets()
    warnings: list[str] = []
    doc = Document()
    sec = doc.sections[0]
    size = PAGE.get(str(document.get("page_size") or "A4").upper(), PAGE["A4"])
    landscape = str(document.get("orientation") or "").lower() == "landscape"
    sec.page_width, sec.page_height = (size[1], size[0]) if landscape else size
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    content_w = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)

    title = to_text(document.get("title"), lang)
    doc.core_properties.title = title
    doc.core_properties.author = "Winmate"
    # 머리 · 바닥
    hdr = sec.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(hdr, " · ".join(p for p in (title, label("confidential", lang) if confidential else "") if p), size=8.5, color="#8a91a0")
    ftr = sec.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(ftr, "PAGE")
    r = ftr.add_run(" / ")
    _font(r, size=8.5, color="#8a91a0")
    _field(ftr, "NUMPAGES")

    if title:
        p = doc.add_paragraph()
        add_runs(p, title, size=22, bold=True, color=theme.c("brand_text"))
        p.paragraph_format.space_after = Pt(4)
    sub = to_text(document.get("subtitle"), lang)
    if sub:
        p = doc.add_paragraph()
        add_runs(p, sub, size=12, color="#596170")
    meta = document.get("meta")
    meta_lines = [to_text(m, lang) for m in as_list(meta)] if meta else []
    if meta_lines:
        p = doc.add_paragraph()
        add_runs(p, " · ".join(m for m in meta_lines if m), size=9.5, color="#8a91a0")
    if title or sub:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

    for si, s in enumerate(sections):
        if not isinstance(s, dict):
            raise DocxError(f"sections[{si}] 형식이 잘못되었습니다")
        if s.get("page_break"):
            doc.add_page_break()
        heading = to_text(s.get("heading") or s.get("title"), lang)
        level = max(1, min(3, int(s.get("level") or 1)))
        if heading:
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            hp.paragraph_format.space_after = Pt(4)
            hp.paragraph_format.keep_with_next = True
            add_runs(hp, heading, size={1: 15, 2: 13, 3: 11.5}[level], bold=True,
                     color=theme.c("brand_text") if level == 1 else "#121417")
        for para in as_list(localize(s.get("paragraphs"), lang)):
            if isinstance(para, dict) and not ({"ko", "en"} & set(para)):
                p = doc.add_paragraph()
                add_runs(p, to_text(para.get("text"), lang), bold=bool(para.get("bold")), italic=bool(para.get("italic")))
            else:
                txt = to_text(para, lang)
                if txt:
                    p = doc.add_paragraph()
                    add_runs(p, txt)
        for b in as_bullets(s.get("bullets"), lang):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, b)
        for b in as_bullets(s.get("numbered"), lang):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, b)
        if s.get("table"):
            _table(doc, s["table"], lang, theme, content_w)
        images = as_list(s.get("images"))
        captions = as_list(s.get("captions"))
        for ii, img in enumerate(images):
            ref = image_ref(localize(img, lang))
            if not ref:
                continue
            a = assets.get(ref["file_id"])
            if a is None:
                warnings.append(f"image missing {ref['file_id']}")
                continue
            data = _png_or_jpeg(a.data)
            if data is None:
                warnings.append(f"image unreadable {ref['file_id']}")
                continue
            width = Cm(float(ref.get("width_cm"))) if ref.get("width_cm") else min(content_w, Cm(15.0))
            doc.add_picture(io.BytesIO(data), width=width)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = to_text(ref.get("caption") or (captions[ii] if ii < len(captions) else ""), lang)
            if a.ai_generated or ref.get("ai_generated"):
                cap = f"{cap} · {label('ai', lang)}" if cap else label("ai", lang)
            if cap:
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(cp, cap, size=9, color="#596170")
        srcs = as_sources(s.get("sources") or s.get("footnotes"), lang)
        if srcs:
            sp = doc.add_paragraph()
            add_runs(sp, f"{label('source', lang)}: " + " · ".join(
                f"{x['label']} ({x['url']})" if x.get("url") else x["label"] for x in srcs), size=8.5, color="#8a91a0")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), warnings


def _table(doc: Any, spec: Any, lang: str, theme: Theme, content_w: int) -> None:
    tbl = as_table(spec, lang)
    cols, rows = tbl.get("columns") or [], tbl.get("rows") or []
    n_cols = max(len(cols), max((len(r) for r in rows), default=0))
    if n_cols == 0:
        return
    t = doc.add_table(rows=(1 if cols else 0) + len(rows), cols=n_cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = tbl.get("col_widths")
    if not (isinstance(widths, list) and len(widths) == n_cols):
        widths = [1.4] + [1.0] * (n_cols - 1) if n_cols > 2 else [1.0] * n_cols
    tot = float(sum(widths))
    r0 = 0
    if cols:
        for ci in range(n_cols):
            c = t.cell(0, ci)
            c.width = int(content_w * widths[ci] / tot)
            c.text = ""
            add_runs(c.paragraphs[0], col_label(cols[ci], lang) if ci < len(cols) else "", size=9.5, bold=True,
                     color="#ffffff" if theme.c("on_brand") == "#ffffff" else "#121417")
            _shade(c, theme.brand)
        r0 = 1
        # 머리 행 반복
        tr_pr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        tr_pr.append(th)
    hl = set(tbl.get("highlight_rows") or [])
    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            c = t.cell(r0 + ri, ci)
            c.width = int(content_w * widths[ci] / tot)
            c.text = ""
            raw = row[ci] if ci < len(row) else ""
            add_runs(c.paragraphs[0], cell_text(raw, lang), size=9.5, bold=ci == 0 or ri in hl)
            if ri in hl:
                _shade(c, theme.c("tint"))
    cap = to_text(tbl.get("caption"), lang)
    if cap:
        p = doc.add_paragraph()
        add_runs(p, cap, size=9, color="#596170")
