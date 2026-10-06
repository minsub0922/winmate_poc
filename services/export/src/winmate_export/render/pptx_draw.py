"""python-pptx 그리기 도우미 — 좌표는 1280 × 720 캔버스 px, 슬라이드 크기에 맞춰 EMU 로 바꾼다."""
from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any

from lxml import etree
from PIL import Image
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import textfit
from .theme import Theme, rgb
from .values import split_markers

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def color(hex_: str) -> RGBColor:
    r, g, b = rgb(hex_)
    return RGBColor(r, g, b)


@dataclass
class Run:
    text: str
    size: float                 # px(캔버스 기준)
    bold: bool = False
    color: str = "#121417"
    number: bool = False        # 숫자 글꼴
    url: str | None = None
    italic: bool = False


@dataclass
class P:
    runs: list[Run]
    align: str = "left"
    line: float = 1.25
    space_before: float = 0.0   # px
    bullet: bool = False


class Canvas:
    """슬라이드 하나에 그리기."""

    def __init__(self, slide: Any, slide_w: int, slide_h: int, theme: Theme, lang: str):
        self.slide = slide
        self.sx = slide_w / 1280.0
        self.sy = slide_h / 720.0
        self.fscale = self.sx / 9525.0       # 13.333in 기준 1px = 9525 EMU
        self.theme = theme
        self.lang = lang

    # ── 좌표 ────────────────────────────────────────────
    def emu(self, x: float, y: float, w: float, h: float) -> tuple[Emu, Emu, Emu, Emu]:
        return Emu(int(round(x * self.sx))), Emu(int(round(y * self.sy))), Emu(max(1, int(round(w * self.sx)))), Emu(max(1, int(round(h * self.sy))))

    def pt(self, px: float) -> Pt:
        return Pt(round(px * 0.75 * self.fscale * 2) / 2)

    # ── 도형 ────────────────────────────────────────────
    def rect(self, x: float, y: float, w: float, h: float, *, fill: str | None = None, line: str | None = None,
             line_w: float = 1.0, radius: float = 0.0, alpha: float | None = None, shape: Any = None,
             dash: bool = False) -> Any:
        if w <= 0 or h <= 0:
            return None
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE)
        sp = self.slide.shapes.add_shape(kind, *self.emu(x, y, w, h))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                sp.adjustments[0] = max(0.0, min(0.5, radius / max(1.0, min(w, h))))
            except Exception:  # noqa: BLE001
                pass
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = color(self.theme.c(fill))
            if alpha is not None:
                set_alpha(sp.fill._xPr.find(qn("a:solidFill")), alpha)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = color(self.theme.c(line))
            sp.line.width = Emu(int(line_w * 9525 * self.fscale))
            if dash:
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            sp.line.fill.background()
        drop_style(sp)
        tf = sp.text_frame
        tf.text = ""
        return sp

    def line(self, x1: float, y1: float, x2: float, y2: float, *, col: str = "line2", width: float = 1.5,
             dash: bool = False) -> Any:
        c = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x1 * self.sx)), Emu(int(y1 * self.sy)),
                                            Emu(int(x2 * self.sx)), Emu(int(y2 * self.sy)))
        c.line.color.rgb = color(self.theme.c(col))
        c.line.width = Emu(int(width * 9525 * self.fscale))
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        drop_style(c)
        return c

    def ellipse(self, x: float, y: float, w: float, h: float, *, fill: str | None = None, line: str | None = None,
                line_w: float = 1.0) -> Any:
        return self.rect(x, y, w, h, fill=fill, line=line, line_w=line_w, shape=MSO_SHAPE.OVAL)

    def polygon(self, pts: Sequence[tuple[float, float]], *, fill: str | None = None, line: str | None = None,
                closed: bool = True, width: float = 1.5) -> Any:
        if len(pts) < 2:
            return None
        fb = self.slide.shapes.build_freeform(int(pts[0][0] * self.sx), int(pts[0][1] * self.sy), scale=1.0)
        fb.add_line_segments([(int(x * self.sx), int(y * self.sy)) for x, y in pts[1:]], close=closed)
        sp = fb.convert_to_shape()
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = color(self.theme.c(fill))
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = color(self.theme.c(line))
            sp.line.width = Emu(int(width * 9525 * self.fscale))
        else:
            sp.line.fill.background()
        drop_style(sp)
        return sp

    # ── 글자 ────────────────────────────────────────────
    def text(self, x: float, y: float, w: float, h: float, paras: list[P], *, anchor: str = "top",
             shape: Any = None, margin: tuple[float, float, float, float] = (0, 0, 0, 0), vertical: bool = False) -> Any:
        if w <= 2 or h <= 2 or not paras:
            return None
        sp = shape or self.slide.shapes.add_textbox(*self.emu(x, y, w, h))
        tf = sp.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        l, t, r, b = margin
        tf.margin_left, tf.margin_top = Emu(int(l * self.sx)), Emu(int(t * self.sy))
        tf.margin_right, tf.margin_bottom = Emu(int(r * self.sx)), Emu(int(b * self.sy))
        tf.vertical_anchor = ANCHOR.get(anchor, MSO_ANCHOR.TOP)
        if vertical:
            tf._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
        first = True
        for para in paras:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = ALIGN.get(para.align, PP_ALIGN.LEFT)
            # 줄 간격은 pt 로 고정 — 배수(%)는 글꼴 메트릭(Noto CJK ≈ 1.45em)을 곱해 화면 설계(CSS line-height)와 어긋난다
            size_max = max((r.size for r in para.runs if r.text), default=0.0)
            if size_max > 0:
                p.line_spacing = self.pt(size_max * para.line)
            if para.space_before:
                p.space_before = self.pt(para.space_before)
            for run in para.runs:
                for seg, marked in split_markers(run.text):
                    if not seg:
                        continue
                    r = p.add_run()
                    r.text = seg
                    self._font(r, run, marked)
                    if run.url:
                        try:
                            r.hyperlink.address = run.url
                        except Exception:  # noqa: BLE001
                            pass
        return sp

    def _font(self, r: Any, run: Run, marked: bool) -> None:
        f = r.font
        f.size = self.pt(run.size)
        f.bold = run.bold or marked
        f.italic = run.italic or None
        f.color.rgb = color(self.theme.c(self.theme.colors["mark_text"] if marked else run.color))
        face = self.theme.number_font if run.number else self.theme.font
        f.name = face
        rpr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rpr, qn(tag))
            el.set("typeface", face)
        if marked:
            hl = etree.SubElement(rpr, qn("a:highlight"))
            clr = etree.SubElement(hl, qn("a:srgbClr"))
            clr.set("val", self.theme.colors["mark_fill"].lstrip("#").upper())
            # a:highlight 는 a:latin 앞에 와야 한다(스키마 순서)
            latin = rpr.find(qn("a:latin"))
            if latin is not None:
                rpr.remove(hl)
                latin.addprevious(hl)

    def fit_paras(self, x: float, y: float, w: float, h: float, paras: list[textfit.Para], *, anchor: str = "top",
                  min_scale: float = 0.66) -> Any:
        """문단들을 상자에 맞춰(크기 단계 → 말줄임) 넣는다."""
        paras = [p for p in paras if p.text]
        if not paras:
            return None
        fitted = textfit.fit(paras, w, h, min_scale=min_scale)
        sc = fitted.scale
        out = []
        for i, fp in enumerate(fitted.paras):
            s = fp.size * (sc if fp.shrink else 1.0)
            out.append(P([Run(fp.text, s, fp.bold, fp.color, number=fp.number, url=fp.url)], align=fp.align,
                         line=fp.line, space_before=(fp.space_before * sc if i else 0)))
        return self.text(x, y, w, h, out, anchor=anchor)

    def fit_text(self, x: float, y: float, w: float, h: float, items: list[tuple[str, float, bool, str]], *,
                 align: str = "left", anchor: str = "top", line: float = 1.25, gap: float = 0.0,
                 number_first: bool = False, min_scale: float = 0.66, url: str | None = None) -> Any:
        """[(글, 크기, 굵게, 색)] 문단들을 상자에 맞춰 넣는다."""
        paras = [textfit.Para(t, s, b, line, space_before=(gap if i else 0), color=c, number=(number_first and i == 0),
                              align=align, url=url) for i, (t, s, b, c) in enumerate(it for it in items if it[0])]
        return self.fit_paras(x, y, w, h, paras, anchor=anchor, min_scale=min_scale)

    # ── 그림 ────────────────────────────────────────────
    def picture(self, data: bytes, x: float, y: float, w: float, h: float, *, fit: str = "cover",
                radius: float = 0.0, focus: tuple[float, float] = (0.5, 0.5)) -> Any:
        try:
            im = Image.open(io.BytesIO(data))
            iw, ih = im.size
            fmt = (im.format or "").upper()
        except Exception:  # noqa: BLE001
            return None
        if iw <= 0 or ih <= 0:
            return None
        if fmt not in ("PNG", "JPEG", "GIF", "BMP", "TIFF") or max(iw, ih) > MAX_IMAGE_PX:
            data = _reencode(im)
        if fit == "contain":
            s = min(w / iw, h / ih)
            dw, dh = iw * s, ih * s
            pic = self.slide.shapes.add_picture(io.BytesIO(data), *self.emu(x + (w - dw) / 2, y + (h - dh) / 2, dw, dh))
        else:
            pic = self.slide.shapes.add_picture(io.BytesIO(data), *self.emu(x, y, w, h))
            box_r, img_r = w / h, iw / ih
            if img_r > box_r:      # 가로가 넘침 → 좌우 자르기
                keep = box_r / img_r
                cut = 1 - keep
                pic.crop_left = cut * focus[0]
                pic.crop_right = cut * (1 - focus[0])
            elif img_r < box_r:    # 세로가 넘침 → 위아래 자르기
                keep = img_r / box_r
                cut = 1 - keep
                pic.crop_top = cut * focus[1]
                pic.crop_bottom = cut * (1 - focus[1])
            if radius > 0:
                round_picture(pic, radius / max(1.0, min(w, h)))
        return pic

    # ── 표 ──────────────────────────────────────────────
    def table(self, x: float, y: float, w: float, h: float, n_rows: int, n_cols: int, col_w: list[float],
              row_h: list[float]) -> Any:
        gf = self.slide.shapes.add_table(n_rows, n_cols, *self.emu(x, y, w, h))
        tbl = gf.table
        tblpr = tbl._tbl.tblPr
        # 기본 표 스타일(줄무늬 · 머리 굵게)을 끈다 — 직접 칠한다
        tblpr.set("firstRow", "0")
        tblpr.set("bandRow", "0")
        style = tblpr.find(qn("a:tableStyleId"))
        if style is None:
            style = etree.SubElement(tblpr, qn("a:tableStyleId"))
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # No Style, Table Grid → 테두리는 아래에서 덮어쓴다
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(cw * self.sx))
        for i, rh in enumerate(row_h):
            tbl.rows[i].height = Emu(int(rh * self.sy))
        return tbl

    def cell(self, cell: Any, runs: list[Run], *, fill: str | None = None, align: str = "left",
             anchor: str = "middle", pad: float = 8.0, borders: dict[str, tuple[str, float] | None] | None = None) -> None:
        if fill:
            cell.fill.solid()
            cell.fill.fore_color.rgb = color(self.theme.c(fill))
        else:
            cell.fill.background()
        cell.margin_left = cell.margin_right = Emu(int(pad * self.sx))
        cell.margin_top = cell.margin_bottom = Emu(int(3 * self.sy))
        cell.vertical_anchor = ANCHOR.get(anchor, MSO_ANCHOR.MIDDLE)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
        for run in runs:
            for seg, marked in split_markers(run.text):
                if not seg:
                    continue
                r = p.add_run()
                r.text = seg
                self._font(r, run, marked)
        set_cell_borders(cell, borders or {}, self.theme, self.fscale)

    # ── 차트 ────────────────────────────────────────────
    def chart(self, spec: dict[str, Any], x: float, y: float, w: float, h: float) -> Any:
        kind = (spec.get("type") or "bar").lower()
        cats = [str(c) for c in spec.get("categories") or []]
        series = [s for s in spec.get("series") or [] if isinstance(s, dict)]
        if not series:
            return None
        brand, mid, soft, gray = (self.theme.c("brand"), self.theme.c("mid"), self.theme.c("soft"), "#cfd5df")
        palette = [brand, mid, soft, "#8a91a0", "#d5d9e0", "#3d4452"]
        if kind in ("scatter", "bubble"):
            data = BubbleChartData() if kind == "bubble" else XyChartData()
            for s in series:
                ser = data.add_series(str(s.get("name") or ""))
                for p in s.get("points") or []:
                    if kind == "bubble":
                        ser.add_data_point(float(p.get("x", 0)), float(p.get("y", 0)), float(p.get("size", 10)))
                    else:
                        ser.add_data_point(float(p.get("x", 0)), float(p.get("y", 0)))
            ctype = XL_CHART_TYPE.BUBBLE if kind == "bubble" else XL_CHART_TYPE.XY_SCATTER
        else:
            data = CategoryChartData()
            data.categories = cats or [str(i + 1) for i in range(len(series[0].get("values") or []))]
            for s in series:
                vals = []
                for v in s.get("values") or []:
                    try:
                        vals.append(float(v))
                    except (TypeError, ValueError):
                        vals.append(None)
                data.add_series(str(s.get("name") or ""), vals)
            ctype = {
                "bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "hbar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
                "stacked": XL_CHART_TYPE.COLUMN_STACKED, "stacked100": XL_CHART_TYPE.BAR_STACKED_100,
                "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "donut": XL_CHART_TYPE.DOUGHNUT,
                "radar": XL_CHART_TYPE.RADAR_MARKERS, "area": XL_CHART_TYPE.AREA,
            }.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
        gf = self.slide.shapes.add_chart(ctype, *self.emu(x, y, w, h), data)
        ch = gf.chart
        ch.font.size = self.pt(12)
        ch.font.name = self.theme.font
        title = str(spec.get("title") or "").strip()
        ch.has_title = bool(title)
        if title:
            ch.chart_title.text_frame.text = title
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = self.pt(13)
            ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
        multi = len(series) > 1 or kind in ("pie", "doughnut", "donut", "stacked", "stacked100")
        ch.has_legend = multi
        if multi:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size = self.pt(12)
        try:
            plot = ch.plots[0]
        except Exception:  # noqa: BLE001
            return gf
        if kind in ("pie", "doughnut", "donut"):
            pts = plot.series[0].points
            for i in range(len(cats)):
                pts[i].format.fill.solid()
                pts[i].format.fill.fore_color.rgb = color(palette[i % len(palette)])
            plot.has_data_labels = True
            plot.data_labels.number_format = '0"%"' if spec.get("percent") else "General"
            plot.data_labels.font.size = self.pt(12)
            return gf
        try:
            plot.gap_width = 60
        except Exception:  # noqa: BLE001
            pass
        for i, s in enumerate(plot.series):
            col = palette[i % len(palette)] if multi else brand
            if kind in ("line", "radar"):
                s.format.line.color.rgb = color(col)
                s.format.line.width = Pt(2.25 * self.fscale)
                s.smooth = False
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = color(col)
        if not multi and kind in ("bar", "column", "hbar") and spec.get("highlight_last", True):
            ser = plot.series[0]
            n = len(cats)
            for i in range(n):
                pt_ = ser.points[i]
                pt_.format.fill.solid()
                pt_.format.fill.fore_color.rgb = color(brand if i == n - 1 else gray)
        if kind not in ("radar",):
            plot.has_data_labels = bool(spec.get("labels", not multi))
            if plot.has_data_labels:
                dl = plot.data_labels
                dl.font.size = self.pt(12)
                dl.font.color.rgb = color(self.theme.c("gray"))
                if spec.get("unit"):
                    dl.number_format = f'#,##0.##"{spec["unit"]}"'
                    dl.number_format_is_linked = False
                try:
                    dl.position = XL_LABEL_POSITION.OUTSIDE_END if kind in ("bar", "column", "hbar") else XL_LABEL_POSITION.ABOVE
                except Exception:  # noqa: BLE001
                    pass
        try:
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = color(self.theme.c("line"))
            va.format.line.fill.background()
            va.tick_labels.font.size = self.pt(11)
            va.tick_labels.font.color.rgb = color(self.theme.c("gray2"))
            ca = ch.category_axis
            ca.tick_labels.font.size = self.pt(12)
            ca.tick_labels.font.color.rgb = color(self.theme.c("gray"))
            ca.format.line.color.rgb = color(self.theme.c("line2"))
        except Exception:  # noqa: BLE001
            pass
        return gf


MAX_IMAGE_PX = 2400     # 슬라이드 한 장(13.3in)을 다 덮어도 180dpi — 더 크면 줄여 넣는다(파일 크기)


def _reencode(im: Image.Image) -> bytes:
    im = im.copy()
    if max(im.size) > MAX_IMAGE_PX:
        im.thumbnail((MAX_IMAGE_PX, MAX_IMAGE_PX), Image.LANCZOS)
    buf = io.BytesIO()
    has_alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    if has_alpha:
        im.convert("RGBA").save(buf, "PNG", optimize=True)
    else:
        im.convert("RGB").save(buf, "JPEG", quality=88, optimize=True)
    return buf.getvalue()


# ── XML 손질 ───────────────────────────────────────────────

def drop_style(shape: Any) -> None:
    """도형의 테마 스타일 참조(<p:style>)를 없앤다 — 테마 그림자 · 테두리가 끼어들지 않게(채우기 · 선은 직접 지정)."""
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)
    sp_pr = el.find(qn("p:spPr"))
    if sp_pr is not None and sp_pr.find(qn("a:effectLst")) is None:
        eff = etree.Element(qn("a:effectLst"))
        # 스키마 순서: xfrm, geom, fill, ln, effectLst, ...
        ln = sp_pr.find(qn("a:ln"))
        if ln is not None:
            ln.addnext(eff)
        else:
            sp_pr.append(eff)


def set_alpha(solid_fill: Any, alpha: float) -> None:
    if solid_fill is None:
        return
    clr = solid_fill[0]
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


def round_picture(pic: Any, frac: float) -> None:
    sp_pr = pic._element.spPr
    geom = sp_pr.find(qn("a:prstGeom"))
    if geom is None:
        return
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for g in list(av):
        av.remove(g)
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {int(max(0.0, min(0.5, frac)) * 100000)}")


def set_cell_borders(cell: Any, borders: dict[str, tuple[str, float] | None], theme: Theme, fscale: float) -> None:
    """borders: {'L'|'R'|'T'|'B': (색, 두께px) | None}. 지정하지 않은 변은 선 없음."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for side in ("L", "R", "T", "B"):
        tag = qn(f"a:ln{side}")
        old = tc_pr.find(tag)
        if old is not None:
            tc_pr.remove(old)
    # 스키마 순서: lnL, lnR, lnT, lnB 가 채우기(solidFill 등)보다 앞
    first_fill = None
    for child in tc_pr:
        if child.tag in (qn("a:solidFill"), qn("a:noFill"), qn("a:gradFill"), qn("a:blipFill"), qn("a:pattFill"), qn("a:grpFill")):
            first_fill = child
            break
    for side in ("L", "R", "T", "B"):
        spec = borders.get(side)
        ln = etree.Element(qn(f"a:ln{side}"))
        if spec:
            col, wpx = spec
            ln.set("w", str(int(wpx * 9525 * fscale)))
            sf = etree.SubElement(ln, qn("a:solidFill"))
            c = etree.SubElement(sf, qn("a:srgbClr"))
            c.set("val", theme.c(col).lstrip("#").upper())
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        if first_fill is not None:
            first_fill.addprevious(ln)
        else:
            tc_pr.append(ln)
