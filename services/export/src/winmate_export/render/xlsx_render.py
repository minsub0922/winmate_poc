"""XLSX — openpyxl 로 표 시트를 만든다.

문서: {title?, sheets: [{name, title?, subtitle?, columns: [{key, label, width?, format?, align?, wrap?}],
        rows: [{key: 값|칸객체} | [값...]], freeze?, merges?, notes?, header_style?, sources?,
        highlight_rows?, group_rows?, autofilter?, orientation?}]}
칸객체: {value, bold?, italic?, fill?, color?, comment?, format?, link?, align?}
"""
from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .theme import FONT, Theme, norm_hex
from .values import en_markers, excel_text, has_marker, label, localize, to_text

INVALID_SHEET = re.compile(r"[\[\]:*?/\\]")
FORMATS = {
    "text": "@", "int": "#,##0", "integer": "#,##0", "number": "#,##0.##", "decimal": "#,##0.00", "percent": "0.0%",
    "pct": "0.0%", "currency": "#,##0\"원\"", "krw": "#,##0\"원\"", "usd": "\"$\"#,##0.00", "date": "yyyy-mm-dd",
}
MARK_FILL = PatternFill("solid", fgColor="FFF2A8")
MARK_FONT_COLOR = "9A3412"
THIN = Side(style="thin", color="D5D9E0")
HEAD_LINE = Side(style="medium", color="121417")


class XlsxError(ValueError):
    pass


def _argb(v: str | None, theme: Theme) -> str | None:
    if not v:
        return None
    c = theme.c(v) if not v.startswith("#") else v
    h = norm_hex(c, "")
    return h[1:].upper() if h else None


def _sheet_name(name: Any, used: set[str], idx: int) -> str:
    base = INVALID_SHEET.sub(" ", excel_text(str(name or f"Sheet{idx + 1}"))).strip()[:31] or f"Sheet{idx + 1}"
    out, n = base, 2
    while out.lower() in used:
        suffix = f" ({n})"
        out = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(out.lower())
    return out


def _text(v: Any, lang: str) -> str:
    s = to_text(v, lang)
    return s


def _cell_value(raw: Any, lang: str) -> Any:
    """숫자는 숫자로 둔다(엑셀 계산 · 서식). 글은 언어 · 표시 변환."""
    v = localize(raw, lang)
    if isinstance(v, dict):
        v = localize(v.get("value", v.get("text")), lang)
    if isinstance(v, bool):
        return "●" if v else ""
    if isinstance(v, (int, float)):
        return v
    if v is None:
        return None
    if isinstance(v, (list, tuple)):
        return excel_text(" · ".join(_text(x, lang) for x in v if x not in (None, "")))
    s = excel_text(str(v))
    if lang == "en":
        s = en_markers(s)
    if s.startswith("="):
        s = "'" + s      # 수식 주입 방지 — 값으로만 넣는다
    return s


def _width_of(s: str) -> float:
    w = 0.0
    for ch in s:
        w += 2.0 if ord(ch) > 0x2E80 else 1.1
    return w


def build_xlsx(document: dict[str, Any], *, design: dict[str, Any] | None = None, lang: str = "ko",
               confidential: bool = False) -> bytes:
    sheets = document.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise XlsxError("sheets 가 비어 있습니다")
    theme = Theme.from_design(design)
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    title = to_text(document.get("title"), lang)
    wb.properties.title = title or None
    wb.properties.creator = "Winmate"
    if confidential:
        wb.properties.keywords = "CONFIDENTIAL"
    for si, spec in enumerate(sheets):
        render_sheet(wb, spec, si, theme=theme, lang=lang, confidential=confidential, used=used)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_sheet(wb: Workbook, spec: Any, si: int, *, theme: Theme, lang: str, confidential: bool, used: set[str]) -> Any:
    """표 시트 하나를 wb 에 우리 디자인으로 만든다(used = 이미 쓴 시트 이름 소문자 — 고객사 양식 뒤에 덧붙일 때도 쓴다)."""
    brand = theme.brand[1:].upper()
    if not isinstance(spec, dict):
        raise XlsxError(f"sheets[{si}] 형식이 잘못되었습니다")
    ws = wb.create_sheet(_sheet_name(to_text(spec.get("name"), lang), used, si))
    cols = spec.get("columns") or []
    rows = spec.get("rows") or []
    if not cols and rows and isinstance(rows[0], dict):
        cols = [{"key": k, "label": k} for k in rows[0].keys()]
    if not cols and rows and isinstance(rows[0], list):
        cols = [{"key": str(i), "label": ""} for i in range(len(rows[0]))]
    cols = [c if isinstance(c, dict) else {"key": str(c), "label": c} for c in cols]
    n_cols = max(1, len(cols))
    r = 1
    stitle = to_text(spec.get("title"), lang)
    if stitle:
        ws.cell(r, 1, _cell_value(stitle, lang)).font = Font(name=FONT, size=14, bold=True)
        if n_cols > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n_cols)
        r += 1
    sub = to_text(spec.get("subtitle"), lang)
    if sub:
        ws.cell(r, 1, _cell_value(sub, lang)).font = Font(name=FONT, size=10, color="596170")
        if n_cols > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n_cols)
        r += 1
    if stitle or sub:
        r += 1
    header_row = r
    hs = spec.get("header_style") or {}
    h_fill = _argb(hs.get("fill"), theme) or brand
    h_color = _argb(hs.get("color"), theme) or ("FFFFFF" if theme.c("on_brand") == "#ffffff" else "121417")
    h_bold = hs.get("bold", True)
    widths = [0.0] * n_cols
    if cols and any(to_text(c.get("label"), lang) for c in cols):
        for ci, c in enumerate(cols):
            txt = to_text(c.get("label") if c.get("label") is not None else c.get("key"), lang)
            cell = ws.cell(r, ci + 1, _cell_value(txt, lang))
            cell.font = Font(name=FONT, size=10, bold=h_bold, color=h_color)
            cell.fill = PatternFill("solid", fgColor=h_fill)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=HEAD_LINE)
            widths[ci] = max(widths[ci], _width_of(txt))
        ws.row_dimensions[r].height = 22
        r += 1
    else:
        header_row = 0
    first_data = r
    hl_rows = set(spec.get("highlight_rows") or [])
    grp_rows = set(spec.get("group_rows") or [])
    for ri, row in enumerate(rows):
        if isinstance(row, dict):
            cells = [row.get(c.get("key")) for c in cols]
            row_opts = row.get("_style") or {}
        elif isinstance(row, (list, tuple)):
            cells = list(row) + [None] * (n_cols - len(row))
            row_opts = {}
        else:
            cells = [row] + [None] * (n_cols - 1)
            row_opts = {}
        for ci in range(n_cols):
            raw = cells[ci] if ci < len(cells) else None
            c = cols[ci] if ci < len(cols) else {}
            val = _cell_value(raw, lang)
            cell = ws.cell(r, ci + 1, val)
            bold = ri in grp_rows or bool(row_opts.get("bold"))
            color = None
            fill = None
            if ri in hl_rows:
                fill, bold, color = theme.c("tint")[1:].upper(), True, theme.c("brand_text")[1:].upper()
            if ri in grp_rows:
                fill = "F5F6F8"
            fmt = FORMATS.get(str(c.get("format") or "").lower(), c.get("format")) if c.get("format") else None
            align = c.get("align") or ("right" if isinstance(val, (int, float)) and not isinstance(val, bool) else "left")
            if isinstance(raw, dict) and not ({"ko", "en"} & set(raw)):
                bold = bold or bool(raw.get("bold"))
                fill = _argb(raw.get("fill"), theme) or fill
                color = _argb(raw.get("color"), theme) or color
                if raw.get("format"):
                    fmt = FORMATS.get(str(raw["format"]).lower(), raw["format"])
                if raw.get("align"):
                    align = raw["align"]
                if raw.get("comment"):
                    cell.comment = Comment(excel_text(_text(raw["comment"], lang))[:2000], "Winmate")
                if raw.get("link"):
                    cell.hyperlink = str(raw["link"])
                    color = color or "1428A0"
                italic = bool(raw.get("italic"))
            else:
                italic = False
            if isinstance(val, str) and has_marker(val):
                fill, bold, color = "FFF2A8", True, MARK_FONT_COLOR
            cell.font = Font(name=FONT, size=10, bold=bold, italic=italic, color=color or "121417")
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if fmt and isinstance(val, (int, float)):
                cell.number_format = fmt
            elif fmt == "@":
                cell.number_format = "@"
            cell.alignment = Alignment(horizontal=align, vertical="top",
                                       wrap_text=c.get("wrap", True) if isinstance(val, str) else False)
            cell.border = Border(bottom=THIN)
            if val is not None:
                widths[ci] = max(widths[ci], min(60.0, max(_width_of(line) for line in str(val).split("\n"))))
        r += 1
    last_data = r - 1
    # 열 폭
    for ci, c in enumerate(cols):
        w = c.get("width")
        ws.column_dimensions[get_column_letter(ci + 1)].width = float(w) if isinstance(w, (int, float)) else max(8.0, min(60.0, widths[ci] + 2))
    # 병합
    for m in spec.get("merges") or []:
        try:
            if isinstance(m, str):
                ws.merge_cells(m)
            elif isinstance(m, dict):
                ws.merge_cells(start_row=int(m["row"]) + first_data, start_column=int(m["col"]) + 1,
                               end_row=int(m.get("row_end", m["row"])) + first_data,
                               end_column=int(m.get("col_end", m["col"])) + 1)
        except Exception as exc:  # noqa: BLE001
            raise XlsxError(f"병합 범위가 잘못되었습니다: {m}") from exc
    # 틀 고정
    fr = spec.get("freeze", True)
    if fr is True and header_row:
        ws.freeze_panes = ws.cell(header_row + 1, 1)
    elif isinstance(fr, str):
        ws.freeze_panes = fr
    elif isinstance(fr, dict):
        ws.freeze_panes = ws.cell(int(fr.get("row", 1)) + 1, int(fr.get("col", 0)) + 1)
    if spec.get("autofilter") and header_row and last_data >= first_data:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{last_data}"
    # 메모 · 출처(표 아래)
    notes = spec.get("notes")
    notes = [notes] if isinstance(notes, (str, dict)) else (notes or [])
    srcs = spec.get("sources") or []
    if notes or srcs:
        r += 1
    for n in notes:
        ws.cell(r, 1, _cell_value(_text(n, lang), lang)).font = Font(name=FONT, size=9, color="596170")
        r += 1
    if srcs:
        parts = []
        for s in srcs:
            if isinstance(s, dict):
                parts.append(_text(s.get("label") or s.get("url"), lang) + (f" ({s['url']})" if s.get("url") and s.get("label") else ""))
            else:
                parts.append(_text(s, lang))
        ws.cell(r, 1, excel_text(f"{label('source', lang)}: " + " · ".join(p for p in parts if p))).font = Font(name=FONT, size=9, color="8A91A0")
        r += 1
    # 인쇄
    ws.page_setup.orientation = spec.get("orientation") or ("landscape" if n_cols > 5 else "portrait")
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"
    if confidential:
        ws.oddHeader.right.text = "CONFIDENTIAL"
    ws.oddFooter.center.text = "&P / &N"
    return ws
