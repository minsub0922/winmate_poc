"""고객사 XLSX 양식 채우기 — 고객 파일 사본의 칸에 값만 써 넣는다(서식 · 병합 · 수식 · 열 너비 · 인쇄 설정 유지).

쓰는 곳: `POST /v1/exports {format: "xlsx", base_file_id, document?, fills?, form?}` → service.build(mode="form").
모델 호출 없이 결정적으로 맞춘다. 같은 입력이면 같은 칸에 쓴다.

입력 두 가지(함께 써도 된다 — document 먼저, fills 가 나중에 덮는다)
1. `document.sheets[]` — 우리 표 문서(xlsx_render 와 같은 모양). 행 이름(행 이름 열, 기본 = 글이 든 첫 열 · `form.label_key`) ×
   열 머리(`columns[].label`)로 양식 칸을 찾는다. 시트 짝: `sheets[i].form.sheet` → 이름이 같은 양식 시트 → 첫 문서 시트는 첫 보이는 시트.
2. `fills[]` — `{sheet?, cell? | row?, column?, value, note?}` 한 칸씩. `cell`(D7 · 'Sheet'!D7)은 그 칸, `row` 는 행 번호(1부터) 또는
   행 이름, `column` 은 열 번호 · 머리 글(못 찾으면 열 글자 D 로도 본다), `column` 이 없으면 행 이름 칸 바로 오른쪽.

맞추기
- 이름 비교 단계(앞 단계에서 맞으면 멈춘다): exact(NFKC · 대소문자 · 공백 · 끝의 `:` `*`) → normalized(번호 매김 「1.」「①」「가.」「-」 ·
  기호 · 공백 무시) → base(괄호 · 대괄호 안 무시: 「밝기 (cd/m²)」=「밝기」) → fuzzy(오타 수준 0.88 이상이고 2등과 차이가 날 때만, `form.fuzzy=false` 로 끔).
  한 칸에 두 언어(「밝기 / Brightness」 · 줄바꿈 · 「밝기 (Brightness)」)가 있으면 각각으로도 비교. 수식 칸은 저장된 계산값으로 비교. 문서의 {ko, en} 이름은 둘 다 쓴다.
- 머리 행: 위 60행 중 문서 열 이름이 가장 많이 맞는 행(같으면 위쪽). 여러 줄 머리(병합 머리 아래 제품명)는 머리 행 근처(-3..+2)도 본다.
- 행 이름: 머리 행 아래, 값 열이 아닌 열에서 찾는다. 같은 이름이 여럿이면 앞서 맞춘 행 다음 것부터(문서 순서 = 양식 순서). 한 양식 행은 한 번만.
- 값 열이 하나뿐인 문서인데 그 열 이름이 양식에 없으면: 맞춘 행들에서 모두 비어 있고 쓸 수 있는 열 중 「비고 · 메모 · 참고」를 빼고
  「제안 · 회신 · 응답 · 기재」 머리가 하나면 그 열, 후보가 하나뿐이면 그 열(match = inferred). 아니면 후보를 보고(unmatched_columns[].candidates).
- 방향: `form.orientation=auto`(기본)면 문서 그대로(행 = 항목)와 돌린 것(행 = 제품) 중 쓸 칸이 많은 쪽. 명시 힌트가 있으면 그대로(rows).
- 명시 힌트가 늘 먼저: `form.header_row` · `form.label_column` · `form.column_map{열 key: 'D'|4|'머리 글'}` ·
  `form.row_map{행 이름 · _key: 7|'양식 이름'}` · 행 dict 의 `_form_row`(번호) · `_form_label`(양식 이름) · 열의 `form_column`.
  (`sheet` · `header_row` · `label_column` 은 첫 문서 시트와 fills 기본 시트에만, 나머지는 모든 문서 시트에. 시트마다 `sheets[i].form` 으로 덮는다.)

쓰기 규칙
- 병합 칸 안쪽이면 같은 행의 병합 머리칸(왼쪽 위)에 쓴다(redirected_from). 다른 행에 걸친 병합이면 건너뛴다(skipped: merged).
- 수식 칸은 건너뛴다(formula — `form.overwrite_formulas=true` 로 덮기). 보호된 시트의 잠긴 칸은 건너뛴다(locked).
- 이미 값이 있는 칸: 이름으로 맞춘 칸은 비었거나 자리표시(`-` · `(기재)` · `[확인 필요]` 등)일 때만 쓴다(not_empty —
  `form.overwrite="always"` 로 덮기). fills 의 칸 주소(cell · 행 번호 + 열 번호/글자)는 덮는다(replaced 로 알림).
- 값: 숫자는 숫자로(셀 서식 그대로), 숫자 모양 글('1,200')도 숫자로(텍스트 서식 칸 · 0 으로 시작 · 15자리 넘음 제외). `=` 로 시작하는 글은
  수식이 아닌 글로. 목록은 「 · 」로 잇는다. null · 빈 글은 쓰지 않는다. 목록 유효성(드롭다운) 밖 값은 쓰고 알린다(cells[].warning).
- [확인 필요] 같은 자리표시가 든 칸은 노란 칠 + 굵게(`form.highlight_marks=false` 로 끔). note(또는 값 {value, note|comment})는 셀 메모.
- 못 맞춘 것은 보고: unmatched_rows · unmatched_columns · unmatched_sheets. 양식에만 있는 행은 form_only_rows.
  `form.unmatched_rows="append"` 면 양식 맨 아래(마지막 값 행 + 2)에 「추가 항목」 줄과 함께 마지막으로 맞춘 행 서식을 복사해 덧붙인다.
  `form.extra_sheets="append"` 면 짝 없는 문서 시트를 우리 디자인 시트로 뒤에 덧붙인다.
- 저장: 열 때 다시 계산(fullCalcOnLoad). openpyxl 이 다시 쓰지 못하는 양식 요소(도형 · 글상자 · 양식 컨트롤 · 머리글 그림 · x14 확장
  유효성 · 조건부 서식 · 스파크라인 · 슬라이서 등)는 저장 전후 부품 수를 비교해 lost 와 경고로 알린다(몰래 잃지 않는다).
"""
from __future__ import annotations

import contextlib
import datetime
import difflib
import io
import re
import unicodedata
import zipfile
from collections import Counter
from collections.abc import Callable
from copy import copy
from dataclasses import dataclass, field, replace
from typing import Any, ClassVar

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string

from .theme import Theme
from .values import MARKER, en_markers, excel_text, localize, to_text
from .xlsx_render import render_sheet

Pos = tuple[int, int]

MAX_SCAN_ROWS = 5000
MAX_SCAN_COLS = 200
HEADER_SCAN_ROWS = 60
EXCEL_MAX_ROW, EXCEL_MAX_COL = 1_048_576, 16_384
FUZZY_MIN = 0.88
FUZZY_MARGIN = 0.03
REPORT_CAP = 300
MARK_FILL = PatternFill("solid", fgColor="FFF2A8")
MARK_COLOR = "9A3412"
LEVELS = ("exact", "normalized", "base")
LEVEL_WEIGHT = {"exact": 1.0, "normalized": 0.95, "base": 0.85, "fuzzy": 0.6}
# 약한 순서(칸 하나의 match 는 행 · 열 중 약한 쪽)
STRENGTH = {"cell": 0, "hint": 0, "number": 0, "letter": 0, "exact": 1, "normalized": 2, "next_to_label": 2, "base": 3,
            "inferred": 4, "appended": 4, "fuzzy": 5}
PLACEHOLDERS = {"", "-", "--", "---", "–", "—", "―", "·", ".", "…", "_", "__", "___", "( )", "()", "[ ]", "[]", "(기재)", "(입력)",
                "(작성)", "(기입)", "기재", "입력", "작성", "기입", "tbd", "(tbd)", "n/a?"}
NEGATIVE_HEADERS = ("비고", "메모", "참고", "remark", "note", "comment", "memo")
POSITIVE_HEADERS = ("제안", "회신", "응답", "답변", "기재", "입력", "작성", "proposal", "proposed", "response", "answer", "offer",
                    "bid", "vendor", "supplier")
APPEND_TITLE = {"ko": "추가 항목", "en": "Additional items"}

_WS = re.compile(r"\s+")
_TRAIL = re.compile(r"[\s:：*※]+$")
_NONWORD = re.compile(r"[\W_]+")
_SPLIT = re.compile(r"\s*(?:\n|\s/\s|\|)\s*")
_ENUM = re.compile(
    r"^\s*(?:"
    r"[-–—―•·∙*※▪■□◆◇○●◦▶▷►☞→>]+"                     # 글머리 기호
    r"|[(（]?\d{1,3}(?:[.\-]\d{1,3})*[.)．）](?!\d)"      # 1.  1)  (1)  1-2.
    r"|\d{1,3}(?:\.\d{1,3})+(?=\s)"                        # 1.1 (공백 앞)
    r"|[①-⑳㉠-㉭㉮-㉻]"                                      # 원문자
    r"|[가나다라마바사아자차카타파하][.)．）]"                # 가.  나)
    r"|[A-Za-z][.)](?=\s)"                                 # A.  b)
    r")\s*")
_BRACKETS = re.compile(r"\([^()]*\)|\[[^\[\]]*\]|\{[^{}]*\}|<[^<>]*>|【[^【】]*】|「[^「」]*」|『[^『』]*』")
_PAREN_PAIR = re.compile(r"(.+?)\s*[(（]([^()（）]+)[)）]\s*")
_HANGUL = re.compile(r"[가-힣]")
_LATIN = re.compile(r"[A-Za-z]")
_NUM = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")
_CELL = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d{1,7})$")
_LETTERS = re.compile(r"^\$?([A-Za-z]{1,3})$")
_ISO_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")


class FormError(Exception):
    def __init__(self, code: str, message: str, details: dict[str, Any] | None = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.details = details or {}


# ── 이름 정규화 ───────────────────────────────────────────

def _nfkc(s: str) -> str:
    return unicodedata.normalize("NFKC", s)


def _strip_enum(s: str) -> str:
    for _ in range(3):
        t = _ENUM.sub("", s, count=1)
        if t == s:
            break
        s = t
    return s


def key_exact(s: str) -> str:
    return _TRAIL.sub("", _WS.sub(" ", _nfkc(s).casefold()).strip())


def key_norm(s: str) -> str:
    return _NONWORD.sub("", _nfkc(_strip_enum(s)).casefold())


def key_base(s: str) -> str:
    t = _nfkc(_strip_enum(s))
    for _ in range(3):
        u = _BRACKETS.sub(" ", t)
        if u == t:
            break
        t = u
    return _NONWORD.sub("", t.casefold())


def _keys(s: str) -> tuple[str, str, str]:
    """(exact, normalized, base) — key_exact · key_norm · key_base 와 같고 중간 결과를 같이 쓴다(색인이 빠르게)."""
    exact = _TRAIL.sub("", _WS.sub(" ", _nfkc(s).casefold()).strip())
    stripped = _nfkc(_strip_enum(s)).casefold()
    norm = _NONWORD.sub("", stripped)
    t = stripped
    for _ in range(3):
        u = _BRACKETS.sub(" ", t)
        if u == t:
            break
        t = u
    return exact, norm, (norm if t == stripped else _NONWORD.sub("", t))


def _script(s: str) -> str | None:
    h, latin = bool(_HANGUL.search(s)), bool(_LATIN.search(s))
    return "ko" if h and not latin else ("en" if latin and not h else None)


def variants(text: str) -> list[str]:
    """한 칸 글의 비교 후보 — 통째 + 줄바꿈 · ' / ' · '|' 로 나눈 것 + 두 언어 괄호(「밝기 (Brightness)」) · 붙은 「밝기/Brightness」."""
    s = str(text).strip()
    if not s:
        return []
    if not any(ch in s for ch in "\n/|(（"):
        return [s]                                                 # 대부분의 칸 — 나눌 것이 없다
    out = [s]
    parts = [p for p in _SPLIT.split(s) if p.strip()]
    if len(parts) > 1:
        out += parts
    for p in [s, *parts]:
        m = _PAREN_PAIR.fullmatch(p)
        if (m and _script(m.group(1)) and _script(m.group(2)) and _script(m.group(1)) != _script(m.group(2))
                and len(_NONWORD.sub("", m.group(2))) >= 3):            # 「(W)」「(kg)」 같은 단위는 따로 보지 않는다
            out += [m.group(1), m.group(2)]
        if "/" in p and " / " not in p:
            a, _, b = p.partition("/")
            if _script(a) and _script(b) and _script(a) != _script(b):
                out += [a, b]
    return list(dict.fromkeys(x.strip() for x in out if x.strip()))


def _plain(v: Any) -> str:
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:g}"
    if isinstance(v, str):
        return v.strip()
    return to_text(v, "ko").strip()


def labels_of(v: Any) -> list[str]:
    """문서 이름 값({ko, en} · 칸 객체 · 글 · 숫자) → 비교 후보(ko · en)."""
    if isinstance(v, dict) and not ({"ko", "en"} & set(v)):
        v = v.get("value", v.get("text", v.get("label")))
    out: list[str] = []
    if isinstance(v, dict):
        for k in ("ko", "en"):
            t = _plain(v.get(k))
            if t:
                out.append(t)
    else:
        t = _plain(v)
        if t:
            out.append(t)
    return list(dict.fromkeys(out))


def _has_kw(text: str, words: tuple[str, ...]) -> bool:
    t = key_exact(text or "")
    return any(w in t for w in words)


def is_blank(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() == ""
    if isinstance(v, dict) and not ({"ko", "en"} & set(v)):
        return is_blank(v.get("value", v.get("text")))
    if isinstance(v, dict):
        return all(is_blank(x) for x in v.values())
    if isinstance(v, list):
        return all(is_blank(x) for x in v)
    return False


def is_placeholder(v: Any) -> bool:
    """비었거나 자리표시('-' · '(기재)' · '[확인 필요]')인 칸 — 이름으로 맞춘 값은 여기에만 쓴다."""
    if v is None:
        return True
    if not isinstance(v, str):
        return False
    s = key_exact(v)
    if s in PLACEHOLDERS or re.fullmatch(r"[-–—―_.·…\s]*", s):
        return True
    return bool(MARKER.fullmatch(v.strip()))


# ── 이름 찾기 ─────────────────────────────────────────────

def _grams(s: str) -> frozenset[str]:
    return frozenset(s[i:i + 2] for i in range(len(s) - 1)) or frozenset([s])


class LabelIndex:
    """양식 칸 글 → 위치(병합이면 머리칸). 단계별 키 사전 + 오타 비교용 키(위치별).

    단계: exact = 칸 글 통째의 exact 키만 · normalized = 통째의 normalized 키 + 나눈 후보(두 언어 · 괄호 바깥 이름)의 exact/normalized 키 ·
    base = 통째 · 후보의 base 키. exact 가 아닌 단계에서 서로 다른 칸 글 여럿이 맞으면 고르지 않는다(ambiguous — 후보로 보고).
    """

    def __init__(self) -> None:
        self.maps: list[dict[str, list[Pos]]] = [{} for _ in LEVELS]
        self.norm_at: dict[Pos, list[str]] = {}
        self.whole_at: dict[Pos, str] = {}

    def _put(self, lvl: int, key: str, pos: Pos) -> None:
        if key:
            bucket = self.maps[lvl].setdefault(key, [])
            if not bucket or bucket[-1] != pos:
                bucket.append(pos)

    def add(self, pos: Pos, text: str) -> None:
        vs = variants(text)
        if not vs:
            return
        whole = _keys(vs[0])
        self.whole_at[pos] = whole[0]
        for lvl, k in enumerate(whole):
            self._put(lvl, k, pos)
        norms = [whole[1]] if len(whole[1]) >= 4 else []
        for v in vs[1:]:
            k = _keys(v)
            self._put(1, k[0], pos)
            self._put(1, k[1], pos)
            self._put(2, k[2], pos)
            if len(k[1]) >= 4:
                norms.append(k[1])
        if norms:
            self.norm_at[pos] = list(dict.fromkeys(norms))

    @staticmethod
    def wanted(labels: list[str]) -> list[set[str]]:
        out: list[set[str]] = [set() for _ in LEVELS]
        for lab in labels:
            vs = variants(lab)
            if not vs:
                continue
            for lvl, k in enumerate(_keys(vs[0])):
                out[lvl].add(k)
            for v in vs[1:]:
                k = _keys(v)
                out[1].update((k[0], k[1]))
                out[2].add(k[2])
        return out

    def find(self, labels: list[str], allow: Callable[[Pos], bool]) -> tuple[str | None, list[Pos]]:
        """exact → normalized → base 단계로 찾는다(앞 단계에서 맞으면 멈춘다). 서로 다른 글 여럿이면 ('ambiguous', 후보)."""
        for lvl, keys in enumerate(self.wanted(labels)):
            hits = sorted({p for k in keys if k for p in self.maps[lvl].get(k, ()) if allow(p)})
            if not hits:
                continue
            if lvl > 0 and len({self.whole_at.get(p) for p in hits}) > 1:
                return "ambiguous", hits
            return LEVELS[lvl], hits
        return None, []

    def fuzzy_many(self, queries: dict[int, list[str]], pool: list[Pos]) -> dict[int, tuple[list[Pos], float]]:
        """오타 수준 비교를 한꺼번에 — 번호마다 가장 비슷한 키 하나(FUZZY_MIN 이상, 2등과 FUZZY_MARGIN 이상 차이)의 위치.

        pool(비교할 양식 위치)을 좁혀 부르고, 길이 · 두 글자 묶음(Dice) 거르기 뒤에만 SequenceMatcher 를 돌린다.
        """
        keys: dict[str, list[Pos]] = {}
        for p in pool:
            for k in self.norm_at.get(p, ()):
                keys.setdefault(k, []).append(p)
        if not keys or not queries:
            return {}
        grams = [(k, len(k), _grams(k)) for k in keys]
        sm = difflib.SequenceMatcher(autojunk=False)
        out: dict[int, tuple[list[Pos], float]] = {}
        for i, labels in queries.items():
            norms = {n for lab in labels for v in variants(lab) if len(n := key_norm(v)) >= 4}
            scored: dict[str, float] = {}
            for n in norms:
                ln, gn = len(n), _grams(n)
                sm.set_seq2(n)
                for k, lk, gk in grams:
                    if 2 * min(ln, lk) < FUZZY_MIN * (ln + lk) or 4 * len(gn & gk) < len(gn) + len(gk):
                        continue
                    sm.set_seq1(k)
                    if sm.quick_ratio() < FUZZY_MIN:
                        continue
                    r = sm.ratio()
                    if r >= FUZZY_MIN and r > scored.get(k, 0.0):
                        scored[k] = r
            if not scored:
                continue
            ranked = sorted(scored.items(), key=lambda kv: -kv[1])
            if len(ranked) > 1 and ranked[0][1] - ranked[1][1] < FUZZY_MARGIN:
                continue                                           # 비슷한 후보가 둘 — 고르지 않는다
            out[i] = (sorted(keys[ranked[0][0]]), ranked[0][1])
        return out

    def find_one(self, labels: list[str], allow: Callable[[Pos], bool], *, fuzzy: bool) -> tuple[str | None, list[Pos]]:
        lvl, hits = self.find(labels, allow)
        if lvl or not fuzzy:
            return lvl, hits
        got = self.fuzzy_many({0: labels}, [p for p in self.norm_at if allow(p)])
        return ("fuzzy", got[0][0]) if 0 in got else (None, [])


# ── 양식 시트 ─────────────────────────────────────────────

class FormSheet:
    """양식 시트 하나 — 병합 지도 · 칸 글 사전 · 이름 색인 · 쓴 범위."""

    def __init__(self, ws: Any, cached: Callable[[str, int, int], Any] | None = None):
        self.ws = ws
        self.title = ws.title
        self._cached = cached
        self.protected = bool(getattr(ws.protection, "sheet", False))
        self.dim_row, self.dim_col = ws.max_row, ws.max_column     # 서식까지 든 시트 범위(쓰기 전) — 「범위 밖」 판단
        self._refresh_merges()
        self.text: dict[Pos, str] = {}
        self.last_row = self.last_col = 0
        self.truncated = False
        for (r, c), cell in list(ws._cells.items()):
            v = cell.value
            if v is None:
                continue
            if r > MAX_SCAN_ROWS or c > MAX_SCAN_COLS:
                self.truncated = True
                continue
            self.last_row, self.last_col = max(self.last_row, r), max(self.last_col, c)
            t = self._label_text(r, c, cell)
            if t:
                self.text[(r, c)] = t
        for b in self.span.values():
            if b[0] <= self.last_row:
                self.last_row = max(self.last_row, min(b[2], MAX_SCAN_ROWS))
        self.index = LabelIndex()
        for p, t in sorted(self.text.items()):
            self.index.add(p, t)
        self.choices = self._list_validations()

    def _refresh_merges(self) -> None:
        self.span: dict[Pos, tuple[int, int, int, int]] = {}
        self._inside: dict[Pos, Pos] = {}
        self._big: list[tuple[int, int, int, int]] = []
        for mr in self.ws.merged_cells.ranges:
            b = (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
            self.span[(b[0], b[1])] = b
            if (b[2] - b[0] + 1) * (b[3] - b[1] + 1) > 4000:
                self._big.append(b)
                continue
            for r in range(b[0], b[2] + 1):
                for c in range(b[1], b[3] + 1):
                    if (r, c) != (b[0], b[1]):
                        self._inside[(r, c)] = (b[0], b[1])

    def _label_text(self, r: int, c: int, cell: Any) -> str | None:
        v = cell.value
        if cell.data_type == "f":
            v = self._cached(self.title, r, c) if self._cached else None     # 수식 칸은 저장된 계산값으로
        if isinstance(v, str):
            return v.strip() or None
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return _plain(v)
        return None

    def _list_validations(self) -> list[tuple[Any, list[str]]]:
        out = []
        for dv in getattr(self.ws.data_validations, "dataValidation", []) or []:
            f = (dv.formula1 or "").strip()
            if dv.type == "list" and len(f) >= 2 and f[0] == '"' and f[-1] == '"':
                out.append((dv.sqref, [x.strip() for x in f[1:-1].split(",")]))
        return out

    def anchor_of(self, r: int, c: int) -> Pos:
        a = self._inside.get((r, c))
        if a:
            return a
        for b in self._big:
            if b[0] <= r <= b[2] and b[1] <= c <= b[3]:
                return (b[0], b[1])
        return (r, c)

    def bounds(self, p: Pos) -> tuple[int, int, int, int]:
        return self.span.get(p, (p[0], p[1], p[0], p[1]))

    def text_at(self, r: int, c: int) -> str:
        return self.text.get(self.anchor_of(r, c), "")

    def header_text(self, header_row: int | None, c: int) -> str:
        if not header_row:
            return ""
        for r in (header_row, header_row - 1, header_row - 2):
            if r >= 1 and self.text_at(r, c):
                return self.text_at(r, c)
        return ""

    def value(self, r: int, c: int) -> Any:
        cell = self.ws._cells.get((r, c))
        return None if cell is None else cell.value

    def is_formula(self, r: int, c: int) -> bool:
        cell = self.ws._cells.get((r, c))
        return cell is not None and cell.data_type == "f"      # 읽은 파일의 data_type 은 정확하다('=' 로 시작하는 글은 s)

    def _dim_style(self, r: int, c: int) -> Any | None:
        """빈 칸(셀 객체 없음)에 엑셀이 입히는 서식 — 행 서식(customFormat) → 열 서식(<col min..max> 묶음 포함)."""
        rd = self.ws.row_dimensions.get(r)
        if rd is not None and rd.has_style:
            return rd
        cd = self.ws.column_dimensions.get(get_column_letter(c))
        if cd is not None and cd.has_style:
            return cd
        for cd in self.ws.column_dimensions.values():
            if cd.has_style and cd.min and cd.max and cd.min <= c <= cd.max:
                return cd
        return None

    def is_locked(self, r: int, c: int) -> bool:
        if not self.protected:
            return False
        cell = self.ws._cells.get((r, c))
        if cell is not None:
            return bool(cell.protection.locked)
        d = self._dim_style(r, c)                                       # 입력 열을 열 서식으로 풀어 둔 양식이 많다
        return bool(d.protection.locked) if d is not None else True     # 보호 시트의 칸은 기본이 잠김

    def number_format(self, r: int, c: int) -> str:
        cell = self.ws._cells.get((r, c))
        if cell is not None:
            return cell.number_format
        d = self._dim_style(r, c)
        return d.number_format if d is not None else "General"

    def ensure_cell(self, r: int, c: int) -> Any:
        """쓸 칸 — 없던 칸은 행 · 열 서식을 입혀 만든다(엑셀에서 빈 칸에 입력할 때와 같게)."""
        cell = self.ws._cells.get((r, c))
        if cell is None:
            d = self._dim_style(r, c)
            cell = self.ws.cell(r, c)
            if d is not None:
                cell._style = copy(d._style)
        return cell

    def writable_empty(self, r: int, c: int) -> bool:
        """이름으로 맞춘 값을 쓸 수 있는 칸인가(같은 행 병합은 머리칸으로)."""
        a = self.anchor_of(r, c)
        if a[0] != r:
            return False
        return not self.is_formula(*a) and not self.is_locked(*a) and is_placeholder(self.value(*a))

    def bottom(self) -> int:
        """값 · 병합이 있는 마지막 행(덧붙일 자리 계산)."""
        last = 0
        for (r, _c), cell in self.ws._cells.items():
            if cell.value is not None:
                last = max(last, r)
        for b in self.span.values():
            last = max(last, b[2])
        return last


# ── 옵션 ─────────────────────────────────────────────────

def column_number(v: Any, *, field_name: str = "label_column") -> int:
    if isinstance(v, bool):
        raise FormError("VALIDATION_FAILED", f"{field_name} 은 열 글자(B) 또는 번호(1부터)입니다")
    if isinstance(v, int):
        if 1 <= v <= EXCEL_MAX_COL:
            return v
    elif isinstance(v, str):
        m = _LETTERS.match(v.strip())
        if m:
            try:
                return column_index_from_string(m.group(1).upper())
            except ValueError:
                pass
        if v.strip().isdigit():
            return column_number(int(v.strip()), field_name=field_name)
    raise FormError("VALIDATION_FAILED", f"{field_name} 은 열 글자(B) 또는 번호(1부터)입니다: {v!r}")


@dataclass
class Opts:
    sheet: str | int | None = None
    header_row: int | None = None
    label_column: int | None = None
    label_key: str | None = None
    column_map: dict[str, Any] = field(default_factory=dict)
    row_map: dict[str, Any] = field(default_factory=dict)
    orientation: str = "auto"
    overwrite: str = "empty"
    overwrite_formulas: bool = False
    fuzzy: bool = True
    unmatched_rows: str = "report"
    extra_sheets: str = "drop"
    highlight_marks: bool = True

    CHOICES: ClassVar[dict[str, tuple[str, ...]]] = {"orientation": ("auto", "rows", "columns"), "overwrite": ("empty", "always"),
               "unmatched_rows": ("report", "append"), "extra_sheets": ("drop", "append")}

    @classmethod
    def parse(cls, raw: Any, base: Opts | None = None, *, first: bool = True) -> Opts:
        o = replace(base) if base else cls()
        if not first:   # sheet · header_row · label_column 은 첫 문서 시트(와 fills)에만
            o.sheet, o.header_row, o.label_column = None, None, None
        if raw is None:
            return o
        if not isinstance(raw, dict):
            raise FormError("VALIDATION_FAILED", "form 은 객체여야 합니다")
        for k, v in raw.items():
            if v is None:
                continue
            if k == "sheet":
                if not isinstance(v, (str, int)) or isinstance(v, bool):
                    raise FormError("VALIDATION_FAILED", "form.sheet 은 시트 이름 또는 번호(0부터)입니다")
                o.sheet = v
            elif k == "header_row":
                if not isinstance(v, int) or isinstance(v, bool) or not 1 <= v <= EXCEL_MAX_ROW:
                    raise FormError("VALIDATION_FAILED", "form.header_row 는 1 이상의 행 번호입니다")
                o.header_row = v
            elif k == "label_column":
                o.label_column = column_number(v, field_name="form.label_column")
            elif k == "label_key":
                o.label_key = str(v)
            elif k in ("column_map", "row_map"):
                if not isinstance(v, dict):
                    raise FormError("VALIDATION_FAILED", f"form.{k} 는 객체여야 합니다")
                setattr(o, k, {**getattr(o, k), **{str(kk): vv for kk, vv in v.items() if vv not in (None, "")}})
            elif k in cls.CHOICES:
                if v not in cls.CHOICES[k]:
                    raise FormError("VALIDATION_FAILED", f"form.{k} 는 {' · '.join(cls.CHOICES[k])} 중 하나입니다: {v}")
                setattr(o, k, v)
            elif k in ("overwrite_formulas", "fuzzy", "highlight_marks"):
                setattr(o, k, bool(v))
        return o

    def has_hints(self) -> bool:
        return bool(self.header_row or self.label_column or self.column_map or self.row_map)


# ── 문서 표 ───────────────────────────────────────────────

@dataclass
class DocTable:
    source: str
    name: str
    row_labels: list[list[str]]
    row_hints: list[Any]
    col_keys: list[str]
    col_labels: list[list[str]]
    col_hints: list[Any]
    cells: list[list[Any]]
    label_header: list[str]
    has_hints: bool = False

    def display(self, labels: list[str], fallback: str = "") -> str:
        return labels[0] if labels else fallback

    def has_values(self, i: int) -> bool:
        return any(not is_blank(v) for v in self.cells[i])

    def transposed(self) -> DocTable:
        n = len(self.row_labels)
        return DocTable(source=self.source, name=self.name, row_labels=self.col_labels, row_hints=[None] * len(self.col_keys),
                        col_keys=[f"#{i}" for i in range(n)], col_labels=self.row_labels, col_hints=[None] * n,
                        cells=[[self.cells[i][j] for i in range(n)] for j in range(len(self.col_keys))], label_header=[])


def _textish(values: list[Any]) -> bool:
    vals = [v for v in values if not is_blank(v)]
    if not vals:
        return False
    texty = sum(1 for v in vals if labels_of(v) and not _NUM.match(labels_of(v)[0]))
    return texty * 2 >= len(vals)


def doc_table(spec: dict[str, Any], si: int, opts: Opts) -> DocTable:
    cols = spec.get("columns") or []
    rows = [r for r in (spec.get("rows") or []) if r is not None]
    if not cols and rows and isinstance(rows[0], dict):
        cols = [{"key": k, "label": k} for k in rows[0] if not str(k).startswith("_")]
    if not cols and rows and isinstance(rows[0], (list, tuple)):
        cols = [{"key": str(i), "label": ""} for i in range(len(rows[0]))]
    cols = [c if isinstance(c, dict) else {"key": str(c), "label": c} for c in cols]
    keys = [str(c.get("key") if c.get("key") is not None else i) for i, c in enumerate(cols)]
    if not keys:
        raise FormError("VALIDATION_FAILED", f"sheets[{si}] 에 columns · rows 가 없습니다")

    def cell(row: Any, ci: int) -> Any:
        if isinstance(row, dict):
            return row.get(keys[ci])
        if isinstance(row, (list, tuple)):
            return row[ci] if ci < len(row) else None
        return row if ci == 0 else None

    if opts.label_key is not None:
        if opts.label_key not in keys:
            raise FormError("VALIDATION_FAILED", f"form.label_key '{opts.label_key}' 가 sheets[{si}].columns 에 없습니다",
                            {"columns": keys})
        li = keys.index(opts.label_key)
    else:   # 글이 든 첫 열(번호 열 No 는 건너뜀)
        li = next((ci for ci in range(len(keys)) if _textish([cell(r, ci) for r in rows])), 0)
    value_idx = [ci for ci in range(len(keys)) if ci != li]
    row_labels, row_hints, cells = [], [], []
    for r in rows:
        labs = labels_of(cell(r, li))
        hint = None
        if isinstance(r, dict):
            if isinstance(r.get("_form_row"), int) and not isinstance(r.get("_form_row"), bool):
                hint = r["_form_row"]
            elif isinstance(r.get("_form_label"), str) and r["_form_label"].strip():
                hint = r["_form_label"]
            else:
                for k in [r.get("_key"), r.get("row_key"), *labs]:
                    if k is None:
                        continue
                    hint = opts.row_map.get(str(k), next((v for kk, v in opts.row_map.items() if key_exact(kk) == key_exact(str(k))), None))
                    if hint is not None:
                        break
        elif labs:
            hint = next((v for kk, v in opts.row_map.items() if any(key_exact(kk) == key_exact(x) for x in labs)), None)
        if isinstance(hint, bool) or (isinstance(hint, int) and not 1 <= hint <= EXCEL_MAX_ROW):
            raise FormError("VALIDATION_FAILED", f"양식 행 번호가 잘못되었습니다: {hint!r}")
        row_labels.append(labs)
        row_hints.append(hint)
        cells.append([cell(r, ci) for ci in value_idx])
    col_labels, col_hints = [], []
    for ci in value_idx:
        c = cols[ci]
        lab = c.get("label") if c.get("label") not in (None, "") else c.get("key")
        col_labels.append(labels_of(lab))
        col_hints.append(opts.column_map.get(keys[ci], c.get("form_column")))
    lab_col = cols[li]
    has_hints = any(h is not None for h in row_hints) or any(h is not None for h in col_hints) or opts.has_hints()
    return DocTable(source=f"sheets[{si}]", name=_plain(localize(spec.get("name"), "ko")) or f"Sheet{si + 1}",
                    row_labels=row_labels, row_hints=row_hints, col_keys=[keys[ci] for ci in value_idx], col_labels=col_labels,
                    col_hints=col_hints, cells=cells, label_header=labels_of(lab_col.get("label")), has_hints=has_hints)


# ── 표 맞추기 계획 ─────────────────────────────────────────

@dataclass
class TablePlan:
    orientation: str
    table: DocTable
    header_row: int | None = None
    label_col: int | None = None
    col_of: dict[int, tuple[int, str, str]] = field(default_factory=dict)   # 문서 열 → (양식 열, match, 머리 글)
    row_of: dict[int, tuple[int, str, str]] = field(default_factory=dict)   # 문서 행 → (양식 행, match, 양식 이름)
    col_candidates: dict[int, list[str]] = field(default_factory=dict)
    row_candidates: dict[int, list[str]] = field(default_factory=dict)      # 이름이 서로 다른 양식 행 여럿에 맞아 고르지 않은 것
    writes: list[tuple[int, int, int, int]] = field(default_factory=list)   # (문서 행, 문서 열, 양식 행, 양식 열)


def letters_column(fs: FormSheet, s: str) -> int | None:
    """'D' 같은 열 글자 → 번호. 한 글자는 늘, 두 · 세 글자는 쓰인 열 + 2 안쪽만('No' · 'Qty' 같은 머리 글을 먼 열로 읽지 않게)."""
    m = _LETTERS.match(s.strip())
    if not m:
        return None
    try:
        c = column_index_from_string(m.group(1).upper())
    except ValueError:
        return None
    return c if len(m.group(1)) == 1 or c <= fs.last_col + 2 else None


def describe(fs: FormSheet, poss: list[Pos], limit: int = 8) -> list[str]:
    """후보 칸 → ['B5 소비 전력 (최대)', …] (보고용)."""
    return [f"{get_column_letter(p[1])}{p[0]} {fs.text.get(p, '')}".strip() for p in poss[:limit] if p[1]]


def _column_from_hint(fs: FormSheet, hint: Any, header_row: int | None, fuzzy: bool) -> tuple[int | None, str]:
    if isinstance(hint, int) and not isinstance(hint, bool):
        return (hint, "hint") if 1 <= hint <= EXCEL_MAX_COL else (None, "")
    if not isinstance(hint, str) or not hint.strip():
        return None, ""
    rows = range(max(1, header_row - 3), header_row + 3) if header_row else range(1, min(fs.last_row, HEADER_SCAN_ROWS) + 1)
    allowed = set(rows)
    lvl, poss = fs.index.find_one([hint], lambda p: p[0] in allowed, fuzzy=fuzzy)
    if lvl and lvl != "ambiguous":
        p = min(poss, key=lambda p: (abs(p[0] - header_row) if header_row else p[0], p[1]))
        return p[1], "hint"
    c = letters_column(fs, hint)
    return (c, "hint") if c else (None, "")


def plan_table(fs: FormSheet, t: DocTable, opts: Opts, orientation: str) -> TablePlan:
    plan = TablePlan(orientation=orientation, table=t)
    idx = fs.index
    n_rows, n_cols = len(t.row_labels), len(t.col_keys)
    used_cols: set[int] = set()
    label_hint = opts.label_column

    # 1. 열 — 힌트 → 머리 행 고르기 → 머리 근처에서 열 정하기
    for j in range(n_cols):
        if t.col_hints[j] is None:
            continue
        c, how = _column_from_hint(fs, t.col_hints[j], opts.header_row, opts.fuzzy)
        if c and c not in used_cols:
            plan.col_of[j] = (c, how, fs.header_text(opts.header_row, c))
            used_cols.add(c)
    head_rows = {opts.header_row} if opts.header_row else set(range(1, min(fs.last_row, HEADER_SCAN_ROWS) + 1))
    hits: dict[int, tuple[str, list[Pos]]] = {}

    def in_head(p: Pos) -> bool:
        return p[0] in head_rows and p[1] not in used_cols and p[1] != label_hint

    for j in range(n_cols):
        if j in plan.col_of or not t.col_labels[j]:
            continue
        lvl, poss = idx.find(t.col_labels[j], in_head)
        if lvl == "ambiguous":
            plan.col_candidates[j] = describe(fs, poss)
        elif lvl:
            hits[j] = (lvl, poss)
    if opts.fuzzy:
        rest = {j: t.col_labels[j] for j in range(n_cols)
                if j not in plan.col_of and j not in hits and j not in plan.col_candidates and t.col_labels[j]}
        if rest:
            for j, (poss, _r) in idx.fuzzy_many(rest, [p for p in idx.norm_at if in_head(p)]).items():
                hits[j] = ("fuzzy", poss)
    score: Counter[int] = Counter()
    for lvl, poss in hits.values():
        for r in {p[0] for p in poss}:
            score[r] += LEVEL_WEIGHT[lvl]
    label_head: list[Pos] = []
    if t.label_header:
        lvl, poss = idx.find(t.label_header, lambda p: p[0] in head_rows)
        if lvl and lvl != "ambiguous":
            label_head = poss
            for r in {p[0] for p in poss}:
                score[r] += 0.5
    header_row = opts.header_row or (min(score, key=lambda r: (-score[r], r)) if score else None)
    plan.header_row = header_row

    def near(r: int) -> int:
        return abs(r - header_row) if header_row else r

    for j in sorted(hits, key=lambda j: (LEVELS.index(hits[j][0]) if hits[j][0] in LEVELS else 9, j)):
        lvl, poss = hits[j]
        cands = [p for p in poss if p[1] not in used_cols and (not header_row or header_row - 3 <= p[0] <= header_row + 2)]
        if cands:
            p = min(cands, key=lambda p: (near(p[0]), p[1]))
            plan.col_of[j] = (p[1], lvl, fs.text.get(p, ""))
            used_cols.add(p[1])
    value_cols = {c for c, _, _ in plan.col_of.values()}
    label_col = label_hint
    if not label_col and label_head:
        cands = [p for p in label_head if p[1] not in value_cols]
        if cands:
            label_col = min(cands, key=lambda p: (near(p[0]), p[1]))[1]

    # 2. 행 — 머리 아래, 값 열이 아닌 칸에서 이름 찾기(힌트 먼저)
    min_row = header_row + 1 if header_row else 1

    def in_area(p: Pos) -> bool:
        return p[0] >= min_row and p[1] not in value_cols

    cand_rows: dict[int, tuple[str, list[Pos]]] = {}
    for i in range(n_rows):
        h = t.row_hints[i]
        if isinstance(h, int):
            cand_rows[i] = ("hint", [(h, 0)])
            continue
        lvl, poss = None, []
        if isinstance(h, str):
            lvl, poss = idx.find([h], in_area)
            lvl = "hint" if lvl and lvl != "ambiguous" else None
        if not lvl and t.row_labels[i]:
            lvl, poss = idx.find(t.row_labels[i], in_area)
        if lvl == "ambiguous":
            plan.row_candidates[i] = describe(fs, poss)
        elif lvl:
            cand_rows[i] = (lvl, poss)
    if not label_col:
        cnt: Counter[int] = Counter()
        for _lvl, poss in cand_rows.values():
            cs = {p[1] for p in poss if p[1]}
            if len(cs) == 1:
                cnt[next(iter(cs))] += 1
        if not cnt:
            cnt = Counter(p[1] for _lvl, poss in cand_rows.values() for p in poss if p[1])
        if cnt:
            label_col = min(cnt, key=lambda c: (-cnt[c], c))
    plan.label_col = label_col
    if opts.fuzzy:   # 못 찾은 이름만 한꺼번에 — 행 이름 열(알면)에서
        rest = {i: [*([t.row_hints[i]] if isinstance(t.row_hints[i], str) else []), *t.row_labels[i]]
                for i in range(n_rows) if i not in cand_rows and i not in plan.row_candidates
                and (t.row_labels[i] or isinstance(t.row_hints[i], str))}
        if rest:
            pool = [p for p in idx.norm_at if in_area(p) and (not label_col or p[1] == label_col)]
            for i, (poss, _r) in idx.fuzzy_many(rest, pool).items():
                cand_rows[i] = ("fuzzy", poss)
    used_rows: set[int] = set()
    last = 0
    for i in range(n_rows):
        if i not in cand_rows:
            continue
        lvl, poss = cand_rows[i]
        free = [p for p in poss if p[0] not in used_rows]
        if not free:
            continue
        p = min(free, key=lambda p: (0 if p[0] > last else 1, 0 if (not label_col or p[1] in (0, label_col)) else 1, p[0], p[1]))
        plan.row_of[i] = (p[0], lvl, fs.text.get(p, "") if p[1] else "")
        used_rows.add(p[0])
        last = p[0]

    # 3. 값 열이 하나뿐인데 이름이 안 맞으면 — 맞춘 행들에서 비어 있는(값이 하나도 없는) 열 하나로
    if n_cols == 1 and 0 not in plan.col_of and plan.row_of:
        rows_m = [r for r, _, _ in plan.row_of.values()]
        left = max([label_col or 0] + [fs.bounds(fs.anchor_of(r, label_col))[3] for r in rows_m if label_col])
        cands: list[tuple[int, str]] = []
        for c in range(left + 1, max(fs.last_col, left + 1) + 1):
            if c in used_cols:
                continue
            good = bad = 0
            for r in rows_m:
                a = fs.anchor_of(r, c)
                if a != (r, c) or fs.is_formula(r, c) or fs.is_locked(r, c):
                    continue                                       # 병합 · 수식 · 잠김 칸은 판단에서 뺀다
                if is_placeholder(fs.value(r, c)):
                    good += 1
                else:
                    bad += 1
            if bad or good * 2 < len(rows_m):
                continue
            head = fs.header_text(header_row, c)
            if _has_kw(head, NEGATIVE_HEADERS):
                continue
            cands.append((c, head))
        pos = [x for x in cands if _has_kw(x[1], POSITIVE_HEADERS)]
        pick = pos[0] if len(pos) == 1 else (cands[0] if len(cands) == 1 else None)
        if pick is None and not header_row and cands and cands[0][0] == left + 1:
            pick = cands[0]                                          # 머리 없는 「이름 | 값」 양식
        if pick:
            plan.col_of[0] = (pick[0], "inferred", pick[1])
        elif cands:
            plan.col_candidates[0] = [f"{get_column_letter(c)} {h}".strip() for c, h in cands[:8]]

    for i, (r, _rl, _rt) in plan.row_of.items():
        for j, (c, _cl, _ct) in plan.col_of.items():
            if not is_blank(t.cells[i][j]):
                plan.writes.append((i, j, r, c))
    return plan


# ── 보고 ─────────────────────────────────────────────────

@dataclass
class Report:
    sheets: list[dict[str, Any]] = field(default_factory=list)
    cells: list[dict[str, Any]] = field(default_factory=list)
    unmatched_rows: list[dict[str, Any]] = field(default_factory=list)
    unmatched_columns: list[dict[str, Any]] = field(default_factory=list)
    unmatched_sheets: list[str] = field(default_factory=list)
    skipped: list[dict[str, Any]] = field(default_factory=list)
    form_only_rows: list[dict[str, Any]] = field(default_factory=list)
    appended_rows: list[dict[str, Any]] = field(default_factory=list)
    appended_sheets: list[str] = field(default_factory=list)
    lost: list[dict[str, Any]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    filled: int = 0

    def to_dict(self) -> dict[str, Any]:
        out: dict[str, Any] = {"filled": self.filled, "sheets": self.sheets}
        for k in ("cells", "unmatched_rows", "unmatched_columns", "skipped", "form_only_rows", "appended_rows"):
            items = getattr(self, k)
            out[k] = items[:REPORT_CAP]
            if len(items) > REPORT_CAP:
                out[f"{k}_total"] = len(items)
        out.update(unmatched_sheets=self.unmatched_sheets, appended_sheets=self.appended_sheets, lost=self.lost)
        return out

    def warnings(self) -> list[str]:
        w = list(self.notes)

        def few(items: list[str]) -> str:
            return ", ".join(items[:5]) + (f" 외 {len(items) - 5}개" if len(items) > 5 else "")

        if self.unmatched_rows:
            appended = {(x["sheet"], x.get("source")) for x in self.appended_rows}
            tail = (" — 양식 아래 「추가 항목」에 덧붙였습니다" if appended and all((x["sheet"], x["source"]) in appended for x in self.unmatched_rows)
                    else " — form.row_map · 행의 _form_row 로 지정하거나 form.unmatched_rows='append' 로 덧붙일 수 있어요")
            w.append(f"양식에서 찾지 못한 행 {len(self.unmatched_rows)}개: {few([str(x['label']) for x in self.unmatched_rows])}{tail}")
        if self.unmatched_columns:
            w.append(f"양식에서 찾지 못한 열 {len(self.unmatched_columns)}개: {few([str(x['label']) for x in self.unmatched_columns])}"
                     " — form.column_map 으로 지정할 수 있어요")
        if self.unmatched_sheets:
            w.append(f"양식에 짝이 없는 문서 시트: {few(self.unmatched_sheets)}(form.extra_sheets='append' 로 덧붙이기)")
        if self.skipped:
            by = Counter(x["reason"] for x in self.skipped)
            names = {"formula": "수식", "merged": "병합", "not_empty": "값 있음", "locked": "잠김", "duplicate": "겹침",
                     "out_of_bounds": "범위 밖", "invalid_cell": "주소 오류", "sheet_not_found": "시트 없음", "column_required": "열 없음"}
            w.append(f"쓰지 못한 칸 {len(self.skipped)}개(" + " · ".join(f"{names.get(k, k)} {v}" for k, v in by.most_common()) + ")")
        fuzzy = [c for c in self.cells if c.get("match") == "fuzzy"]
        if fuzzy:
            w.append(f"비슷한 이름으로 맞춘 칸 {len(fuzzy)}개 — 확인해 주세요: "
                     + few([f"{c['cell']} {c.get('row_label') or ''}".strip() for c in fuzzy]))
        outside = [c["cell"] for c in self.cells if c.get("outside")]
        if outside:
            w.append(f"양식 범위 밖에 쓴 칸 {len(outside)}개: {few(outside)}")
        moved = [f"{c['redirected_from']}→{c['cell']}" for c in self.cells if c.get("redirected_from") and c.get("match") == "cell"]
        if moved:
            w.append(f"병합 칸 안쪽 주소라 병합 머리칸에 썼습니다: {few(moved)}")
        bad_choice = [c for c in self.cells if c.get("warning")]
        if bad_choice:
            w.append(f"양식 목록(드롭다운)에 없는 값 {len(bad_choice)}칸: {few([c['cell'] for c in bad_choice])}")
        for x in self.lost:
            w.append(f"양식의 {x['name']} {x['base']}개 중 {x['base'] - x['output']}개가 유지되지 않았습니다(openpyxl 한계) — 원본과 비교해 주세요")
        return w


# ── 쓰기 ─────────────────────────────────────────────────

def _convert(raw: Any, lang: str, number_format: str = "General") -> tuple[Any, str | None]:
    """문서 값 → (셀 값, 메모). 숫자는 숫자로, 숫자 모양 글도 숫자로(텍스트 서식 · 0 시작 · 15자리 초과 제외)."""
    v = localize(raw, lang)
    note = None
    if isinstance(v, dict) and not ({"ko", "en"} & set(v)):
        n = v.get("note", v.get("comment"))
        note = to_text(n, lang) if n not in (None, "") else None
        v = localize(v.get("value", v.get("text")), lang)
    if v is None or v is False:
        return None, note
    if v is True:
        return "●", note
    if isinstance(v, (int, float)):
        return v, note
    if isinstance(v, (list, tuple)):
        v = " · ".join(to_text(x, lang) for x in v if x not in (None, ""))
    s = excel_text(to_text(v, lang) if isinstance(v, dict) else str(v))
    if lang == "en":
        s = en_markers(s)
    if not s.strip():
        return None, note
    t = s.strip()
    m = _ISO_DATE.match(t)
    if m and is_date_format(number_format):                       # 날짜 서식 칸이면 날짜 값으로
        try:
            return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3))), note
        except ValueError:
            pass
    digits = re.sub(r"\D", "", t)
    if _NUM.match(t) and number_format != "@" and not re.match(r"^[+-]?0\d", t) and len(digits) <= 15:
        num = float(t.replace(",", ""))
        return (int(num) if "." not in t else num), note
    return s, note


class Writer:
    def __init__(self, opts: Opts, lang: str, rep: Report):
        self.opts, self.lang, self.rep = opts, lang, rep
        self.written: dict[tuple[str, Pos], str] = {}

    def skip(self, fs: FormSheet | None, r: int | None, c: int | None, reason: str, source: str, **extra: Any) -> None:
        cell = f"{get_column_letter(c)}{r}" if r and c and 1 <= c <= EXCEL_MAX_COL else None
        self.rep.skipped.append({"sheet": fs.title if fs else extra.pop("sheet", None), "cell": cell, "reason": reason,
                                 "source": source, **{k: v for k, v in extra.items() if v is not None}})

    def write(self, fs: FormSheet, r: int, c: int, raw: Any, *, source: str, match: str, explicit: bool,
              note: str | None = None, row_label: str | None = None, column_label: str | None = None,
              row_match: str | None = None, column_match: str | None = None, fresh: bool = False) -> bool:
        if not (1 <= r <= EXCEL_MAX_ROW and 1 <= c <= EXCEL_MAX_COL):
            self.skip(fs, r, c, "out_of_bounds", source)
            return False
        redirected = None
        a = fs.anchor_of(r, c)
        if a != (r, c):
            if a[0] != r:
                self.skip(fs, r, c, "merged", source, anchor=f"{get_column_letter(a[1])}{a[0]}")
                return False
            redirected = f"{get_column_letter(c)}{r}"
            r, c = a
        coord = f"{get_column_letter(c)}{r}"
        if not fresh:
            if fs.is_formula(r, c) and not self.opts.overwrite_formulas:
                self.skip(fs, r, c, "formula", source)
                return False
            if fs.is_locked(r, c):
                self.skip(fs, r, c, "locked", source)
                return False
        prev_src = self.written.get((fs.title, (r, c)))
        if prev_src and not explicit:
            self.skip(fs, r, c, "duplicate", source, by=prev_src)
            return False
        value, cell_note = _convert(raw, self.lang, fs.number_format(r, c))
        if value is None:
            return False
        cur = fs.value(r, c)
        replaced = None
        if not fresh and not prev_src and not is_placeholder(cur):
            if not explicit and self.opts.overwrite != "always":
                self.skip(fs, r, c, "not_empty", source, current=str(cur)[:80])
                return False
            replaced = str(cur)[:80]
        cell = fs.ensure_cell(r, c)                                  # 쓸 때만 칸을 만든다(행 · 열 서식 상속)
        cell.value = value
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"                                     # 수식이 아니라 글로
        text = note or cell_note
        if text:
            old = cell.comment.text if cell.comment else ""
            body = f"{old}\n— Winmate: {text}" if old else text
            cell.comment = Comment(excel_text(body)[:2000], "Winmate")
        if self.opts.highlight_marks and isinstance(value, str) and MARKER.search(value):
            f = copy(cell.font)
            cell.font = Font(name=f.name, size=f.size, bold=True, italic=f.italic, underline=f.underline, strike=f.strike,
                             vertAlign=f.vertAlign, color=MARK_COLOR)
            cell.fill = MARK_FILL
        entry: dict[str, Any] = {"sheet": fs.title, "cell": coord, "match": match, "source": source}
        for k, v in (("row_label", row_label), ("column_label", column_label), ("row_match", row_match),
                     ("column_match", column_match), ("redirected_from", redirected), ("replaced", replaced)):
            if v:
                entry[k] = v
        if isinstance(value, str):
            for rng, allowed in fs.choices:
                try:
                    inside = coord in rng
                except Exception:  # noqa: BLE001
                    inside = False
                if inside and value not in allowed:
                    entry["warning"] = f"목록에 없는 값({', '.join(allowed[:6])})"
                    break
        if explicit and not fresh and (r > max(fs.dim_row, fs.last_row, 1) or c > max(fs.dim_col, fs.last_col, 1)):
            entry["outside"] = True
        self.rep.cells.append(entry)
        if not prev_src:
            self.rep.filled += 1
        self.written[(fs.title, (r, c))] = source
        return True


def _weak(*ms: str | None) -> str:
    ms2 = [m for m in ms if m]
    return max(ms2, key=lambda m: STRENGTH.get(m, 9)) if ms2 else "exact"


# ── 시트 짝 · 덧붙이기 ─────────────────────────────────────

def _visible(wb: Any) -> list[Any]:
    return [ws for ws in wb.worksheets if ws.sheet_state == "visible"] or list(wb.worksheets)


def resolve_sheet(wb: Any, ref: Any, default: Any) -> Any:
    if ref is None or ref == "":
        return default
    sheets = list(wb.worksheets)
    if isinstance(ref, int) and not isinstance(ref, bool):
        return sheets[ref] if 0 <= ref < len(sheets) else None
    name = str(ref)
    for ws in sheets:
        if ws.title == name:
            return ws
    for ws in sheets:
        if key_exact(ws.title) == key_exact(name):
            return ws
    return None


def _copy_row_style(fs: FormSheet, src: int, dst: int, c_from: int, c_to: int) -> None:
    ws = fs.ws
    for c in range(c_from, c_to + 1):
        s = ws._cells.get((src, c))
        if s is not None and getattr(s, "has_style", False):
            ws.cell(dst, c)._style = copy(s._style)
    h = ws.row_dimensions[src].height
    if h:
        ws.row_dimensions[dst].height = h
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row == mr.max_row == src and c_from <= mr.min_col and mr.max_col <= c_to:
            ws.merge_cells(start_row=dst, start_column=mr.min_col, end_row=dst, end_column=mr.max_col)


def _append_rows(fs: FormSheet, plan: TablePlan, pending: list[int], w: Writer, lang: str) -> None:
    t, ws = plan.table, fs.ws
    lc = plan.label_col or 1
    cols = [c for c, _, _ in plan.col_of.values()]
    c_from = min([lc, *cols])
    c_to = max([lc, *cols, *[fs.bounds(fs.anchor_of(r, lc))[3] for r, _, _ in plan.row_of.values()]])
    tmpl = max((r for r, _, _ in plan.row_of.values()), default=None)
    start = fs.bottom() + 2
    title = fs.ensure_cell(start, lc)
    title.value = APPEND_TITLE["en" if lang == "en" else "ko"]
    head = ws._cells.get((plan.header_row, lc)) if plan.header_row else None
    if head is not None and head.has_style:
        title._style = copy(head._style)
    else:
        title.font = Font(bold=True)
    for k, i in enumerate(pending):
        r = start + 1 + k
        if tmpl:
            _copy_row_style(fs, tmpl, r, c_from, c_to)
        fs._refresh_merges()
        label = excel_text(t.display(t.row_labels[i], f"#{i + 1}"))
        lab_cell = fs.ensure_cell(*fs.anchor_of(r, lc))
        lab_cell.value = label
        if label.startswith("="):
            lab_cell.data_type = "s"
        for j, (c, cmatch, ctext) in plan.col_of.items():
            v = t.cells[i][j]
            if not is_blank(v):
                w.write(fs, r, c, v, source=f"{t.source}.rows[{i}]", match="appended", explicit=True, fresh=True,
                        row_label=label, column_label=ctext or t.display(t.col_labels[j], t.col_keys[j]), column_match=cmatch)
        w.rep.appended_rows.append({"sheet": fs.title, "row": r, "label": label, "source": f"{t.source}.rows[{i}]"})


# ── 잃은 요소 검사 ─────────────────────────────────────────

_PART_NAMES = {
    "shapes": "도형 · 글상자", "charts": "차트", "media": "그림 파일", "form_controls": "양식 컨트롤(확인란 · 단추)",
    "activex": "ActiveX 컨트롤", "header_footer_images": "머리글 · 바닥글 그림", "ext_validations": "확장 데이터 유효성(다른 시트 목록 등)",
    "ext_conditional": "확장 조건부 서식", "sparklines": "스파크라인", "slicers": "슬라이서", "timelines": "시간 막대",
    "pivot_tables": "피벗 테이블", "embedded_objects": "포함된 개체", "macros": "매크로(VBA)", "threaded_comments": "스레드 댓글",
    "cell_images": "셀 안 그림", "external_links": "외부 연결", "tables": "표(목록 개체)", "data_validations": "데이터 유효성",
    "conditional_formats": "조건부 서식", "comments": "메모",
}
_PART_RULES: list[tuple[str, re.Pattern[str]]] = [
    ("charts", re.compile(r"xl/charts/chart\d+\.xml$")), ("media", re.compile(r"xl/media/")),
    ("form_controls", re.compile(r"xl/ctrlProps/")), ("activex", re.compile(r"xl/activeX/activeX\d+\.xml$")),
    ("slicers", re.compile(r"xl/slicers/")), ("timelines", re.compile(r"xl/timelines/")),
    ("pivot_tables", re.compile(r"xl/pivotTables/pivotTable\d+\.xml$")), ("embedded_objects", re.compile(r"xl/embeddings/")),
    ("macros", re.compile(r"xl/vbaProject\.bin$")), ("threaded_comments", re.compile(r"xl/threadedComments/")),
    ("cell_images", re.compile(r"xl/richData/")), ("external_links", re.compile(r"xl/externalLinks/externalLink\d+\.xml$")),
    ("tables", re.compile(r"xl/tables/table\d+\.xml$")), ("comments", re.compile(r"xl/comments\d*\.xml$")),
]
_SHAPE = re.compile(rb"<(?:\w+:)?(?:sp|grpSp|cxnSp)[\s>]")


def inventory(data: bytes) -> Counter[str]:
    inv: Counter[str] = Counter()
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return inv
    with z:
        for n in z.namelist():
            for key, rx in _PART_RULES:
                if rx.match(n):
                    inv[key] += 1
            if n.startswith("xl/drawings/") and n.endswith(".xml") and "/_rels/" not in n:
                inv["shapes"] += len(_SHAPE.findall(z.read(n)))
            elif re.match(r"xl/worksheets/sheet\d+\.xml$", n):
                x = z.read(n)
                inv["ext_validations"] += x.count(b"<x14:dataValidation ") + x.count(b"<x14:dataValidation>")
                inv["ext_conditional"] += x.count(b"<x14:conditionalFormatting")
                inv["sparklines"] += x.count(b"<x14:sparklineGroup ") + x.count(b"<x14:sparklineGroup>")
                inv["header_footer_images"] += x.count(b"<legacyDrawingHF")
                inv["data_validations"] += x.count(b"<dataValidation ")
                inv["conditional_formats"] += x.count(b"<conditionalFormatting ")
    return inv


def lost_parts(before: bytes, after: bytes) -> list[dict[str, Any]]:
    a, b = inventory(before), inventory(after)
    return [{"part": k, "name": _PART_NAMES.get(k, k), "base": a[k], "output": b[k]} for k in sorted(a) if b[k] < a[k]]


# ── 진입점 ───────────────────────────────────────────────

MAX_FILLS = 5000


def validate_fills(fills: Any) -> None:
    """요청 검사(API 에서 바로 400) — 칸 주소 모양 · cell 또는 row · 번호 범위."""
    if fills is None:
        return
    if not isinstance(fills, list):
        raise FormError("VALIDATION_FAILED", "fills 는 목록이어야 합니다")
    if len(fills) > MAX_FILLS:
        raise FormError("VALIDATION_FAILED", f"fills 는 {MAX_FILLS}개까지입니다", {"count": len(fills)})
    for k, f in enumerate(fills):
        if not isinstance(f, dict):
            raise FormError("VALIDATION_FAILED", f"fills[{k}] 는 객체여야 합니다")
        cell, row, col = f.get("cell"), f.get("row"), f.get("column")
        if cell not in (None, ""):
            addr = cell.rpartition("!")[2].strip() if isinstance(cell, str) else None
            m = _CELL.match(addr or "")
            if not m or not 1 <= int(m.group(2)) <= EXCEL_MAX_ROW:
                raise FormError("VALIDATION_FAILED", f"fills[{k}].cell 주소가 잘못되었습니다(예 D7): {cell!r}", {"index": k})
            try:
                if column_index_from_string(m.group(1).upper()) > EXCEL_MAX_COL:
                    raise ValueError
            except ValueError as exc:
                raise FormError("VALIDATION_FAILED", f"fills[{k}].cell 이 엑셀 범위 밖입니다: {cell!r}", {"index": k}) from exc
            continue
        if row in (None, "") or isinstance(row, bool) or not isinstance(row, (int, str)):
            raise FormError("VALIDATION_FAILED", f"fills[{k}] 에 cell(칸 주소) 또는 row(행 번호 · 이름)가 필요합니다", {"index": k})
        if isinstance(row, int) and not 1 <= row <= EXCEL_MAX_ROW:
            raise FormError("VALIDATION_FAILED", f"fills[{k}].row 는 1 이상입니다", {"index": k})
        if isinstance(row, int) and col in (None, ""):
            raise FormError("VALIDATION_FAILED", f"fills[{k}] 행 번호에는 column(열 번호 · 글자 · 머리 글)이 필요합니다", {"index": k})
        if col is not None and (isinstance(col, bool) or not isinstance(col, (int, str)) or (isinstance(col, int) and not 1 <= col <= EXCEL_MAX_COL)):
            raise FormError("VALIDATION_FAILED", f"fills[{k}].column 은 열 번호(1부터) · 글자 · 머리 글입니다", {"index": k})


def open_workbook(data: bytes, *, macro: bool = False, data_only: bool = False) -> Any:
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise FormError("INVALID_BASE_FILE", "암호가 걸렸거나 옛 형식(.xls) 엑셀 파일입니다 — 암호를 풀거나 .xlsx 로 저장해 다시 올려 주세요")
    if data[:2] != b"PK":
        raise FormError("INVALID_BASE_FILE", "엑셀(.xlsx) 파일이 아닙니다")
    try:
        return load_workbook(io.BytesIO(data), keep_vba=macro, data_only=data_only, keep_links=True)
    except Exception as exc:
        raise FormError("INVALID_BASE_FILE", f"양식 파일을 열 수 없습니다({type(exc).__name__})") from exc


def _fill_table(fs: FormSheet, t: DocTable, opts: Opts, w: Writer, lang: str) -> None:
    rep = w.rep
    plans = [] if opts.orientation == "columns" else [plan_table(fs, t, opts, "rows")]
    nonblank = sum(1 for row in t.cells for v in row if not is_blank(v))
    strong = bool(plans) and len(plans[0].writes) * 10 >= nonblank * 6        # 곧은 방향이 60% 이상 맞으면 돌려 보지 않는다
    if opts.orientation == "columns" or (opts.orientation == "auto" and not t.has_hints and not strong):
        plans.append(plan_table(fs, t.transposed(), replace(opts, column_map={}, row_map={}), "columns"))
    plan = max(plans, key=lambda p: len(p.writes))                    # 같으면 앞(rows)
    tt = plan.table
    for i, j, r, c in plan.writes:
        _r, rmatch, rtext = plan.row_of[i]
        _c, cmatch, ctext = plan.col_of[j]
        w.write(fs, r, c, tt.cells[i][j], source=f"{t.source}.rows[{i if plan.orientation == 'rows' else j}]",
                match=_weak(rmatch, cmatch), explicit=False, row_label=rtext or tt.display(tt.row_labels[i]),
                column_label=ctext or tt.display(tt.col_labels[j], tt.col_keys[j]), row_match=rmatch, column_match=cmatch)
    cols = []
    for j in range(len(tt.col_keys)):
        got = plan.col_of.get(j)
        cols.append({"key": tt.col_keys[j], "label": tt.display(tt.col_labels[j], tt.col_keys[j]),
                     "column": get_column_letter(got[0]) if got else None, "match": got[1] if got else None})
        if not got and any(not is_blank(tt.cells[i][j]) for i in range(len(tt.row_labels))):
            miss = {"sheet": fs.title, "key": tt.col_keys[j], "label": tt.display(tt.col_labels[j], tt.col_keys[j]), "source": t.source}
            if plan.col_candidates.get(j):
                miss["candidates"] = plan.col_candidates[j]
            rep.unmatched_columns.append(miss)
    pending = [i for i in range(len(tt.row_labels)) if i not in plan.row_of and any(not is_blank(v) for v in tt.cells[i])]
    for i in pending:
        miss = {"sheet": fs.title, "label": tt.display(tt.row_labels[i], f"#{i + 1}"),
                "source": f"{t.source}.{'rows' if plan.orientation == 'rows' else 'columns'}[{i}]"}
        if plan.row_candidates.get(i):
            miss["candidates"] = plan.row_candidates[i]
        rep.unmatched_rows.append(miss)
    matched_rows = {r for r, _, _ in plan.row_of.values()}
    if plan.header_row and plan.label_col:
        value_cols = [c for c, _, _ in plan.col_of.values()]
        for (r, c), text in sorted(fs.text.items()):
            if c != plan.label_col or r <= plan.header_row or r in matched_rows:
                continue
            if all(is_placeholder(fs.value(*fs.anchor_of(r, vc))) for vc in value_cols):
                rep.form_only_rows.append({"sheet": fs.title, "row": r, "label": text})
    rep.sheets.append({"sheet": fs.title, "source": t.source, "orientation": plan.orientation, "header_row": plan.header_row,
                       "label_column": get_column_letter(plan.label_col) if plan.label_col else None, "columns": cols,
                       "rows_total": len(tt.row_labels), "rows_matched": len(plan.row_of),
                       "cells": sum(1 for x in rep.cells if x["sheet"] == fs.title and x["source"].startswith(t.source + "."))})
    if pending and opts.unmatched_rows == "append":
        if plan.orientation == "rows":
            _append_rows(fs, plan, pending, w, lang)
        else:
            rep.notes.append(f"「{fs.title}」은 제품이 행인 양식이라 못 맞춘 항목을 덧붙이지 않았습니다")


def _fill_one(wb: Any, get_fs: Callable[[Any], FormSheet], default_ws: Any, f: Any, k: int, opts: Opts, w: Writer) -> None:
    src = f"fills[{k}]"
    if not isinstance(f, dict):
        w.skip(None, None, None, "invalid_cell", src)
        return
    value = f.get("value")
    if is_blank(value):
        return
    sheet_ref, cell_ref = f.get("sheet"), f.get("cell")
    if isinstance(cell_ref, str) and "!" in cell_ref:
        sp, _, cell_ref = cell_ref.rpartition("!")
        sheet_ref = sheet_ref if sheet_ref not in (None, "") else sp.strip().strip("'")
    ws = resolve_sheet(wb, sheet_ref, default_ws)
    if ws is None:
        w.skip(None, None, None, "sheet_not_found", src, sheet=str(sheet_ref))
        return
    fs = get_fs(ws)
    note = f.get("note") if isinstance(f.get("note"), str) else None
    if isinstance(cell_ref, str) and cell_ref.strip():
        m = _CELL.match(cell_ref.strip())
        if not m:
            w.skip(fs, None, None, "invalid_cell", src, address=cell_ref)
            return
        try:
            c = column_index_from_string(m.group(1).upper())
        except ValueError:
            w.skip(fs, None, None, "out_of_bounds", src, address=cell_ref)
            return
        w.write(fs, int(m.group(2)), c, value, source=src, match="cell", explicit=True, note=note)
        return
    row, col = f.get("row"), f.get("column")
    c = cmatch = None
    if isinstance(col, int) and not isinstance(col, bool):
        c, cmatch = col, "number"
    r = rmatch = None
    label_pos: Pos | None = None
    rtext = None
    if isinstance(row, int) and not isinstance(row, bool):
        r, rmatch = row, "number"
    elif isinstance(row, str) and row.strip():
        lvl, poss = fs.index.find_one([row], lambda p: c is None or p[1] < c, fuzzy=opts.fuzzy)
        if not lvl and c is not None:
            lvl, poss = fs.index.find_one([row], lambda p: True, fuzzy=opts.fuzzy)
        if lvl == "ambiguous":
            w.rep.unmatched_rows.append({"sheet": fs.title, "label": row.strip(), "source": src, "candidates": describe(fs, poss)})
            return
        if lvl:
            label_pos = poss[0]
            r, rmatch, rtext = label_pos[0], lvl, fs.text.get(label_pos)
        elif row.strip().isdigit():
            r, rmatch = int(row.strip()), "number"
        else:
            w.rep.unmatched_rows.append({"sheet": fs.title, "label": row.strip(), "source": src})
            return
    else:
        w.skip(fs, None, None, "invalid_cell", src, address=None)
        return
    ctext = None
    if c is None and isinstance(col, str) and col.strip():
        lvl, poss = fs.index.find_one([col], lambda p: p[0] < r and p[1] != (label_pos or (0, 0))[1], fuzzy=opts.fuzzy)
        by_letter = letters_column(fs, col)
        if lvl == "ambiguous":
            w.rep.unmatched_columns.append({"sheet": fs.title, "key": None, "label": col.strip(), "source": src,
                                            "candidates": describe(fs, poss)})
            return
        if lvl:
            p = max(poss, key=lambda p: (p[0], -p[1]))                  # 행 바로 위의 머리
            c, cmatch, ctext = p[1], lvl, fs.text.get(p)
        elif by_letter:
            c, cmatch = by_letter, "letter"
        elif col.strip().isdigit():
            c, cmatch = int(col.strip()), "number"
        else:
            w.rep.unmatched_columns.append({"sheet": fs.title, "key": None, "label": col.strip(), "source": src})
            return
    elif c is None:
        if label_pos is None:
            w.skip(fs, r, None, "column_required", src)
            return
        c, cmatch = fs.bounds(label_pos)[3] + 1, "next_to_label"
    explicit = rmatch == "number" and cmatch in ("number", "letter")
    w.write(fs, r, c, value, source=src, match=_weak(rmatch, cmatch), explicit=explicit, note=note, row_label=rtext,
            column_label=ctext, row_match=rmatch, column_match=cmatch)


def fill_form(base: bytes, *, document: dict[str, Any] | None = None, fills: list[Any] | None = None,
              options: dict[str, Any] | None = None, lang: str = "ko", design: dict[str, Any] | None = None,
              confidential: bool = False, macro: bool = False) -> tuple[bytes, dict[str, Any], list[str]]:
    """고객 양식 바이트 → (채운 바이트, 보고, 경고). 문서 · fills 어느 쪽도 고객 파일의 다른 칸은 건드리지 않는다."""
    opts = Opts.parse(options)
    wb = open_workbook(base, macro=macro)
    try:
        sheets = _visible(wb)
        if not sheets:
            raise FormError("INVALID_BASE_FILE", "양식에 워크시트가 없습니다")
        values_wb: list[Any] = []

        def cached(title: str, r: int, c: int) -> Any:
            if not values_wb:
                values_wb.append(open_workbook(base, data_only=True))
            try:
                return values_wb[0][title].cell(r, c).value
            except KeyError:
                return None

        rep = Report()
        w = Writer(opts, lang, rep)
        fss: dict[str, FormSheet] = {}

        def get_fs(ws: Any) -> FormSheet:
            if ws.title not in fss:
                fss[ws.title] = FormSheet(ws, cached)
                if fss[ws.title].truncated:
                    rep.notes.append(f"「{ws.title}」이 커서 앞 {MAX_SCAN_ROWS}행 · {MAX_SCAN_COLS}열에서만 이름을 찾았습니다")
            return fss[ws.title]

        default_ws = resolve_sheet(wb, opts.sheet, sheets[0])
        if default_ws is None:
            raise FormError("VALIDATION_FAILED", f"양식에 시트가 없습니다: {opts.sheet}", {"sheets": [s.title for s in wb.worksheets]})

        doc_sheets = (document or {}).get("sheets") or []
        if not isinstance(doc_sheets, list):
            raise FormError("VALIDATION_FAILED", "document.sheets 는 목록이어야 합니다")
        claimed: set[str] = set()
        extra: list[tuple[int, dict[str, Any]]] = []
        for si, spec in enumerate(doc_sheets):
            if not isinstance(spec, dict):
                raise FormError("VALIDATION_FAILED", f"sheets[{si}] 형식이 잘못되었습니다")
            so = Opts.parse(spec.get("form"), opts, first=si == 0)
            ref = (spec.get("form") or {}).get("sheet") if isinstance(spec.get("form"), dict) else None
            if ref is None and si == 0:
                ref = opts.sheet
            target = None
            if ref is not None:
                target = resolve_sheet(wb, ref, None)
                if target is None:
                    raise FormError("VALIDATION_FAILED", f"양식에 시트가 없습니다: {ref}", {"sheets": [s.title for s in wb.worksheets]})
            else:
                name = _plain(localize(spec.get("name"), "ko"))
                en = _plain(localize(spec.get("name"), "en"))
                for ws in wb.worksheets:
                    if ws.title not in claimed and any(n and key_exact(ws.title) == key_exact(n) for n in (name, en)):
                        target = ws
                        break
                if target is None and si == 0:
                    target = next((ws for ws in sheets if ws.title not in claimed), None)
            if target is None:
                extra.append((si, spec))
                continue
            claimed.add(target.title)
            _fill_table(get_fs(target), doc_table(spec, si, so), so, w, lang)
        for si, spec in extra:
            name = _plain(localize(spec.get("name"), lang)) or f"sheets[{si}]"
            if opts.extra_sheets == "append":
                used = {s.lower() for s in wb.sheetnames}
                ws = render_sheet(wb, spec, si, theme=Theme.from_design(design), lang=lang, confidential=confidential, used=used)
                rep.appended_sheets.append(ws.title)
            else:
                rep.unmatched_sheets.append(name)

        for k, f in enumerate(fills or []):
            _fill_one(wb, get_fs, default_ws, f, k, opts, w)

        wb.template = False
        try:
            wb.calculation.fullCalcOnLoad = True                         # 수식은 열 때 다시 계산
        except AttributeError:  # pragma: no cover
            pass
        buf = io.BytesIO()
        wb.save(buf)
        out = buf.getvalue()
    finally:
        arc = getattr(wb, "vba_archive", None)   # keep_vba 사본 zip — 닫지 않으면 GC 때 경고
        if arc is not None:
            with contextlib.suppress(Exception):
                arc.close()
    rep.lost = lost_parts(base, out)
    return out, rep.to_dict(), rep.warnings()
