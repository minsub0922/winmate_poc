"""API — 템플릿 카탈로그 · 썸네일 · 내보내기(201 · 202) · 렌더 · 마스터. files 는 in-process 스텁."""
from __future__ import annotations

import io
import json
import os
import zipfile

import pytest
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from winmate_common import testing
from winmate_export.templates.catalog import catalog
from winmate_export.templates.samples import sample_slots

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _deck(n: int = 2, img: str | None = None) -> dict:
    codes = ["MS-B", "VP-F3", "P3-C", "CB-A", "SS-FB-D", "MGI-D"]
    slides = []
    for i in range(n):
        t = catalog().get(codes[i % len(codes)])
        slides.append({"template_code": t["code"], "slots": sample_slots(t, img, title=f"슬라이드 {i + 1}"),
                       "notes": f"노트 {i + 1}", "sources": [{"label": "KB", "url": "https://kb"}]})
    return {"title": "A 커피 디지털 메뉴보드", "cover": {"title": "A 커피 디지털 메뉴보드", "customer": "A 커피", "date": "2026-10-06"},
            "slides": slides}


# ── 템플릿 ───────────────────────────────────────────────

async def test_list_templates(app):
    async with testing.api_client(app) as c:
        r = await c.get("/v1/templates")
        assert r.status_code == 200
        body = r.json()
        assert body["total"] == len([t for t in catalog().templates if t["status"] != "internal"])
        item = body["items"][0]
        assert {"code", "display_code", "name", "sheet_role", "kind", "thumb_url", "data_shape", "slot_schema"} <= set(item)
        assert item["thumb_url"].startswith("/api/export/v1/templates/")
        r = await c.get("/v1/templates", params={"role": "MS", "industry": "FB", "include_slots": "false"})
        items = r.json()["items"]
        assert items[0]["industry"] == "FB" and items[0]["industry_code"] == "FB"
        assert all(i["sheet_role"] == "MS" for i in items) and "slot_schema" not in items[0]
        r = await c.get("/v1/templates", params={"codes": "VP-F·3,C01,VP-B", "include_slots": "false"})
        assert sorted(i["code"] for i in r.json()["items"]) == ["C01", "VP-B3", "VP-F3"]
        r = await c.get("/v1/templates", params={"proposal_type": "quickwin", "section": "vp", "limit": 5})
        b = r.json()
        assert len(b["items"]) == 5 and b["next_cursor"]
        r2 = await c.get("/v1/templates", params={"proposal_type": "quickwin", "section": "vp", "limit": 5, "cursor": b["next_cursor"]})
        assert {i["code"] for i in r2.json()["items"]}.isdisjoint({i["code"] for i in b["items"]})
        r = await c.get("/v1/templates/stats")
        st = r.json()
        assert st["total"] == len(catalog().templates) and st["by_role"]["VP"] >= 30 and st["catalog_version"]


async def test_get_template(app):
    async with testing.api_client(app) as c:
        r = await c.get("/v1/templates/VP-F·3")
        assert r.status_code == 200
        t = r.json()
        assert t["code"] == "VP-F3" and t["display_code"] == "VP-F·3"
        assert t["slots"] and t["boxes"] and t["source"]["artifact"].startswith("https://claude.ai/artifact/")
        assert (await c.get("/v1/templates/VP-F", params={"n": 4})).json()["code"] == "VP-F4"
        r = await c.get("/v1/templates/NOPE-9")
        assert r.status_code == 404 and r.json()["error"]["code"] == "TEMPLATE_NOT_FOUND"


def _drop_none(v):
    if isinstance(v, dict):
        return {k: _drop_none(x) for k, x in v.items() if x is not None}
    if isinstance(v, list):
        return [_drop_none(x) for x in v]
    return v


async def test_get_template_every_catalog_code(app):
    """카탈로그의 모든 코드(저장 코드 · 표시 코드)와 별칭(n 마다)이 상세 200 이고, 칸 · 상자 · 원본이 카탈로그 그대로.

    회귀(2026-10-07): 목록 기본값(표 머리 headers 등) · 숫자 기본값(GN-*.recommended) · 보드 없는 원본(VP-<업종>-A..C)이
    응답 모델(str | None)에 막혀 128종이 500(ResponseValidationError) 이었다.
    """
    cat = catalog()
    failed: list[tuple[str, str]] = []
    seen = 0
    async with testing.api_client(app) as c:
        async def get(code: str, **params):
            try:
                r = await c.get(f"/v1/templates/{code}", params=params or None)
            except Exception as exc:  # noqa: BLE001 — 응답 검증 오류는 앱 예외로 올라온다
                failed.append((code, type(exc).__name__))
                return None
            if r.status_code != 200:
                failed.append((code, str(r.status_code)))
                return None
            return r.json()

        for t in cat.templates:
            for code in dict.fromkeys([t["code"], t.get("display_code") or t["code"]]):
                d = await get(code)
                seen += 1
                if d is None:
                    continue
                assert d["code"] == t["code"], code
                assert d["slots"] == _drop_none(t["slots"]), code          # 기본값 목록 · 숫자 그대로(문자열로 바꾸지 않음)
                assert d["boxes"] == _drop_none(t["boxes"]), code
                src = t.get("source") or {}
                for k, v in _drop_none(src).items():
                    assert d["source"][k] == v, (code, k)
                if src and not src.get("board"):
                    assert "board" not in d["source"] and d["source"].get("note"), code
                assert set(d["example_slots"]) <= {s["key"] for s in t["slots"]}, code
        for alias, spec in cat.aliases.items():
            d = await get(alias)
            if d is not None:
                assert d["code"] == cat.get(spec["default"])["code"], alias
            for n, target in (spec.get("options") or {}).items():
                d = await get(alias, n=int(n))
                if d is not None:
                    assert d["code"] == cat.get(target)["code"], (alias, n)
    assert not failed, f"{len(failed)}개 실패: {failed[:20]}"
    assert seen >= len(cat.templates)
    # 대표 모양 — 목록 기본값 · 숫자 기본값 · 보드 없는 원본
    cb = {s["key"]: s for s in cat.get("CB-B")["slots"]}
    assert isinstance(cb["headers"]["default"], list)
    assert any(isinstance(s.get("default"), int) for s in cat.get("GN-A")["slots"])
    assert any(not (t.get("source") or {}).get("board") for t in cat.templates)


@pytest.mark.parametrize("w,size", [(None, (120, 68)), (240, (240, 135)), (480, (480, 270))])
async def test_thumbnail(app, w, size):
    async with testing.api_client(app) as c:
        r = await c.get("/v1/templates/MS-B/thumbnail.png", params={"w": w} if w else None)
        assert r.status_code == 200 and r.headers["content-type"] == "image/png"
        assert Image.open(io.BytesIO(r.content)).size == size
        etag = r.headers["etag"]
        r2 = await c.get("/v1/templates/MS-B/thumbnail.png", params={"w": w} if w else None, headers={"If-None-Match": etag})
        assert r2.status_code == 304
        assert (await c.get("/v1/templates/MS-B/thumbnail.png", params={"w": 5000})).status_code == 422
        assert (await c.get("/v1/templates/NOPE/thumbnail.png")).status_code == 404


async def test_board_image(app):
    """레이아웃 브라우저 — 원본 보드 그림(docs/templates/_rendered). 보드 없는(제작 중) 템플릿 · 없는 코드는 404."""
    async with testing.api_client(app) as c:
        r = await c.get("/v1/templates/MS-A/board.jpg")
        assert r.status_code == 200 and r.headers["content-type"] == "image/jpeg"
        assert Image.open(io.BytesIO(r.content)).size == (1280, 720)
        assert (await c.get("/v1/templates/MS-A/board.jpg", headers={"If-None-Match": r.headers["etag"]})).status_code == 304
        assert (await c.get("/v1/templates/NOPE/board.jpg")).status_code == 404
        lst = (await c.get("/v1/templates", params={"status": "in_production", "limit": 1, "include_slots": False})).json()["items"]
        if lst:
            r = await c.get(f"/v1/templates/{lst[0]['code']}/board.jpg")
            assert r.status_code == 404 and r.json()["error"]["code"] == "BOARD_NOT_RENDERED"


# ── 내보내기 ─────────────────────────────────────────────

async def test_export_pptx_sync(app, files, png):
    img = files.put("store.png", png(), "image/png", meta={"generation": {"is_generated": True}})["id"]
    logo = files.put("logo.png", png(300, 80, (0, 0, 0)), "image/png")["id"]
    doc = _deck(3, img)
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "pptx", "filename": "A커피_제안서.pptx", "document": doc,
                                              "design": {"brand_hex": "#E4002B", "logo_file_id": logo},
                                              "confidential": True, "project_id": "prj_1", "source_ref": "proposal:prp_1"})
        assert r.status_code == 201, r.text
        body = r.json()
        assert body["status"] == "done" and body["export_id"].startswith("exp_")
        f = body["file"]
        assert f["name"] == "A커피_제안서.pptx" and f["mime"] == PPTX_MIME and f["url"].endswith("/content")
        assert body["template_codes"] == ["C03", "MS-B", "VP-F3", "P3-C"]          # 표지 그림 없음 → C03
        saved = files.saved[-1]
        assert saved["source"] == "export" and saved["confidential"] is True and saved["project_id"] == "prj_1"
        assert saved["meta"]["export_id"] == body["export_id"] and saved["meta"]["template_codes"][0] == "C03"
        prs = Presentation(io.BytesIO(files.by_id(f["id"])))
        assert len(prs.slides) == 4
        assert "AI 생성 이미지" in "".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
        r = await c.get(f"/v1/exports/{body['export_id']}")
        rec = r.json()
        assert rec["status"] == "done" and rec["file"]["id"] == f["id"] and rec["format"] == "pptx"
        r = await c.get("/v1/exports", params={"project_id": "prj_1"})
        assert [i["export_id"] for i in r.json()["items"]] == [body["export_id"]]
        assert (await c.get("/v1/exports/exp_nope")).status_code == 404


async def test_export_both_languages_files(app, files):
    doc = {"title": "보고", "slides": [{"template_code": "CB-A", "slots": {"title": {"ko": "제목 [확인 필요]", "en": "Title [확인 필요]"}}}]}
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "pptx", "filename": "deck", "document": doc, "language": "ko_en"})
        assert r.status_code == 201, r.text
        names = [f["name"] for f in r.json()["files"]]
        assert names == ["deck_KO.pptx", "deck_EN.pptx"]
        en = Presentation(io.BytesIO(files.by_id(r.json()["files"][1]["id"])))
        text = "".join(sh.text_frame.text for sh in en.slides[0].shapes if sh.has_text_frame)
        assert "Title [TBD]" in text
        r = await c.post("/v1/exports", json={"format": "pptx", "document": doc, "language": "both", "bilingual": "slides"})
        assert len(r.json()["files"]) == 1 and r.json()["slide_count"] == 2


async def test_export_from_document_file(app, files):
    doc = {"slides": [{"kind": "sheet", "template_code": "MS-A", "content": {"title": "문서 파일에서"}, "footnotes": ["KB"]}],
           "design": {"brand_hex": "#0a7d3b"}, "lang": "en", "tbd_label": "[TBC]"}
    fid = files.put("render_doc.json", json.dumps(doc).encode(), "application/json")["id"]
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "pptx", "document_file_id": fid, "filename": "from_file"})
        assert r.status_code == 201, r.text
        assert r.json()["file"]["name"] == "from_file.pptx"
        r = await c.post("/v1/exports", json={"format": "pptx", "document_file_id": "file_01JNOPENOPENOPENOPENOPENOP"})
        assert r.status_code == 404 and r.json()["error"]["code"] == "FILE_NOT_FOUND"


async def test_export_xlsx_docx_pdf_zip(app, files, png):
    img = files.put("cut.png", png(), "image/png")["id"]
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "xlsx", "filename": "수량표", "document": {
            "sheets": [{"name": "제품", "columns": [{"key": "m", "label": "모델"}, {"key": "q", "label": "수량"}],
                        "rows": [{"m": "QM55C", "q": 3}]}]}})
        assert r.status_code == 201, r.text
        xf = r.json()["file"]
        assert xf["name"] == "수량표.xlsx" and load_workbook(io.BytesIO(files.by_id(xf["id"]))).sheetnames == ["제품"]

        report = {"title": "장면 스크립트", "sections": [{"heading": "장면 1", "paragraphs": ["아침 오픈"], "images": [img]}]}
        r = await c.post("/v1/exports", json={"format": "docx", "document": report})
        assert r.status_code == 201, r.text
        assert r.json()["file"]["name"] == "장면 스크립트.docx"

        r = await c.post("/v1/exports", json={"format": "pdf", "document": {"report": report, "page_size": "Letter"}})
        assert r.status_code == 201, r.text
        assert files.by_id(r.json()["file"]["id"])[:5] == b"%PDF-"

        r = await c.post("/v1/exports", json={"format": "zip", "filename": "조감도", "document": {"entries": [
            {"file_id": img, "path": "images/cut.png"}, {"path": "sources.json", "json": {"a": 1}},
            {"path": "수량표.xlsx", "format": "xlsx", "document": {"sheets": [{"name": "s", "rows": [["a", 1]]}]}}]}})
        assert r.status_code == 201, r.text
        z = zipfile.ZipFile(io.BytesIO(files.by_id(r.json()["file"]["id"])))
        assert z.namelist() == ["images/cut.png", "sources.json", "수량표.xlsx"]

        r = await c.post("/v1/exports", json={"format": "zip", "document": {"entries": [{"file_id": "file_01JNOPENOPENOPENOPENOPENOP", "path": "x"}]}})
        assert r.status_code == 404 and r.json()["error"]["code"] == "FILE_NOT_FOUND"
        rec = (await c.get("/v1/exports")).json()["items"]
        assert any(i["status"] == "failed" for i in rec)


async def test_validation_errors(app):
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "pptx", "document": {"slides": [{"template_code": "ZZ-9", "slots": {}}]}})
        assert r.status_code == 400 and r.json()["error"]["code"] == "TEMPLATE_NOT_FOUND"
        assert r.json()["error"]["details"]["codes"] == ["ZZ-9"]
        r = await c.post("/v1/exports", json={"format": "pptx", "document": {"slides": [{"slots": {"title": "코드 없음"}}]}})
        assert r.json()["error"]["code"] == "TEMPLATE_REQUIRED"
        r = await c.post("/v1/exports", json={"format": "pptx", "document": {"slides": []}})
        assert r.status_code == 400 and r.json()["error"]["code"] == "VALIDATION_FAILED"
        r = await c.post("/v1/exports", json={"format": "xlsx", "document": {"sheets": []}})
        assert r.status_code == 400
        r = await c.post("/v1/exports", json={"format": "pptx"})
        assert r.status_code == 400
        r = await c.post("/v1/exports", json={"format": "gif", "document": {}})
        assert r.status_code == 422


async def test_soffice_formats_return_501_without_converter(app, files, no_soffice):
    from winmate_export.render import soffice

    assert not soffice.available() and soffice.status().reason == "off"
    pptx_id = files.put("deck.pptx", b"PK fake", PPTX_MIME)["id"]
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "pdf", "document": _deck(1)})
        assert r.status_code == 501 and r.json()["error"]["code"] == "PDF_CONVERTER_UNAVAILABLE"
        err = r.json()["error"]
        assert err["details"]["env"] == "SOFFICE_PATH" and err["details"]["reason"] == "off" and err["details"]["fix"]
        assert "SOFFICE_PATH=off" in err["message"]
        r = await c.post("/v1/exports", json={"format": "pdf", "from_file_id": pptx_id})
        assert r.status_code == 501
        r = await c.post("/v1/renders", json={"file_id": pptx_id})
        assert r.status_code == 501 and r.json()["error"]["code"] == "PDF_CONVERTER_UNAVAILABLE"
        r = await c.post("/v1/renders", json={"document": _deck(1)})
        assert r.status_code == 501
        info = (await c.get("/v1/info")).json()
        assert info["pdf_converter"] is False and info["pdf_converter_reason"] == "off" and "form_fill" in info["features"]


async def test_render_pdf_pages_without_converter(app, files, no_soffice):
    from winmate_export.render.pdf_render import build_pdf
    from winmate_export.worker import HANDLERS

    pdf, _ = build_pdf({"title": "두 쪽", "sections": [{"heading": "1"}, {"heading": "2", "page_break": True}]})
    fid = files.put("report.pdf", pdf, "application/pdf")["id"]
    async with testing.api_client(app) as c:
        r = await c.post("/v1/renders", json={"file_id": fid, "width": 640})
        assert r.status_code == 202, r.text
        rid = r.json()["render_id"]
        assert await testing.drain_jobs("export", HANDLERS) == 1
        rec = (await c.get(f"/v1/renders/{rid}")).json()
        assert rec["status"] == "done" and len(rec["pages"]) == 2
        png = files.by_id(rec["pages"][0]["png_file_id"])
        assert Image.open(io.BytesIO(png)).size[0] == 640


async def test_slow_export_goes_to_job(app, files, monkeypatch):
    from winmate_export.worker import HANDLERS

    monkeypatch.setenv("EXPORT_SYNC_MAX_SLIDES", "2")
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "pptx", "filename": "큰덱", "document": _deck(4)})
        assert r.status_code == 202, r.text
        acc = r.json()
        assert acc["status"] == "queued" and acc["job_id"].startswith("job_")
        assert (await c.get(f"/v1/exports/{acc['export_id']}")).json()["status"] == "queued"
        assert await testing.drain_jobs("export", HANDLERS) == 1
        rec = (await c.get(f"/v1/exports/{acc['export_id']}")).json()
        assert rec["status"] == "done" and rec["job_id"] == acc["job_id"]
        assert len(Presentation(io.BytesIO(files.by_id(rec["file"]["id"]))).slides) == 5
        job = await __import__("winmate_common.jobs", fromlist=["jobs"]).jobs().get(acc["job_id"])
        assert job.status == "succeeded" and job.result["file_id"] == rec["file"]["id"]
        # async=true 면 작아도 잡
        r = await c.post("/v1/exports", json={"format": "xlsx", "async": True, "document": {"sheets": [{"name": "a", "rows": [[1]]}]}})
        assert r.status_code == 202


async def test_masters(app, files):
    from pptx import Presentation as P

    prs = P()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    buf = io.BytesIO()
    prs.save(buf)
    data = buf.getvalue()
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for it in src.infolist():
            d = src.read(it.filename)
            if it.filename == "[Content_Types].xml":
                d = d.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            z.writestr(it, d)
    potx = files.put("고객사 템플릿.potx", out.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.template")["id"]
    bad = files.put("x.potx", b"nope", "application/octet-stream")["id"]
    async with testing.api_client(app) as c:
        r = await c.get("/v1/masters")
        assert [m["master_id"] for m in r.json()["items"]] == ["samsung_b2b", "retail_fnb", "simple_white"]
        r = await c.post("/v1/masters", json={"file_id": potx})
        assert r.status_code == 201, r.text
        m = r.json()
        assert m["master_id"].startswith("mst_") and m["name"] == "고객사 템플릿" and m["aspect"] == "16:9" and m["layouts"]
        assert len((await c.get("/v1/masters")).json()["items"]) == 4
        r = await c.post("/v1/masters", json={"file_id": bad})
        assert r.status_code == 400 and r.json()["error"]["code"] == "INVALID_MASTER"
        r = await c.post("/v1/exports", json={"format": "pptx", "master_id": m["master_id"], "document": _deck(1)})
        assert r.status_code == 201, r.text
        prs2 = Presentation(io.BytesIO(files.by_id(r.json()["file"]["id"])))
        assert len(prs2.slides) == 2 and prs2.slides[0].slide_layout.name.lower().startswith("blank")
        r = await c.post("/v1/exports", json={"format": "pptx", "master_id": "mst_nope", "document": _deck(1)})
        assert r.status_code == 404 and r.json()["error"]["code"] == "MASTER_NOT_FOUND"
        r = await c.post("/v1/exports", json={"format": "pptx", "document": _deck(1), "design": {"master_id": "simple_white"}})
        assert r.status_code == 201 and r.json()["template_codes"][0] == "C03"


@pytest.mark.skipif(not os.environ.get("SOFFICE_PATH"), reason="LibreOffice(SOFFICE_PATH) 없음")
async def test_pdf_deck_with_libreoffice(app, files, monkeypatch):  # pragma: no cover — 환경에 따라
    from winmate_export.worker import HANDLERS

    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "pdf", "document": _deck(2)})
        assert r.status_code == 202
        await testing.drain_jobs("export", HANDLERS)
        rec = (await c.get(f"/v1/exports/{r.json()['export_id']}")).json()
        assert rec["status"] == "done" and files.by_id(rec["file"]["id"])[:5] == b"%PDF-"


@pytest.mark.skipif(not os.environ.get("SOFFICE_PATH"), reason="LibreOffice(SOFFICE_PATH) 없음")
async def test_render_deck_with_libreoffice(app, files):  # pragma: no cover — 환경에 따라
    from winmate_export.worker import HANDLERS

    doc = _deck(2)
    for i, s in enumerate(doc["slides"]):
        s["sheet_id"] = f"sht_{i + 1}"
    async with testing.api_client(app) as c:
        r = await c.post("/v1/renders", json={"document": doc, "width": 640, "sheet_ids": ["sht_2"]})
        assert r.status_code == 202, r.text
        await testing.drain_jobs("export", HANDLERS)
        rec = (await c.get(f"/v1/renders/{r.json()['render_id']}")).json()
        assert rec["status"] == "done", rec
        assert [p["sheet_id"] for p in rec["pages"]] == ["sht_2"] and rec["pages"][0]["index"] == 2
