"""XLSX · DOCX · PDF(reportlab) · ZIP 만들기."""
from __future__ import annotations

import io
import json
import zipfile

import pypdfium2 as pdfium
import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from winmate_export.render.docx_render import build_docx
from winmate_export.render.pdf_render import build_pdf
from winmate_export.render.pptx_render import Asset, DictAssets
from winmate_export.render.xlsx_render import XlsxError, build_xlsx
from winmate_export.render.zip_build import ZipError, build_zip, safe_path

SHEETS = {
    "title": "수량표",
    "sheets": [{
        "name": "제품/수량[1]", "title": "1층 로비 제품 수량", "subtitle": "가격 열 없음",
        "columns": [{"key": "space", "label": "공간", "width": 14}, {"key": "model", "label": "모델코드"},
                    {"key": "qty", "label": "수량", "format": "int"}, {"key": "basis", "label": "수량 근거"}],
        "rows": [
            {"space": "로비", "model": "QM55C", "qty": 4, "basis": "규칙"},
            {"space": "카운터", "model": {"value": "QB43C", "bold": True, "comment": "대체 모델 가능"}, "qty": 2, "basis": "[확인 필요]"},
            ["창가", "QH55C", 1200, {"value": "=SUM(A1)", "fill": "#eaeefb"}],
        ],
        "merges": ["A1:D1"], "notes": ["* 수량은 도면 기준"], "sources": [{"label": "KB 제품 DB", "url": "https://kb"}],
        "autofilter": True,
    }, {"name": "제품/수량[1]", "rows": [{"a": 1, "b": "x"}]}],
}

REPORT = {
    "title": "경쟁사 분석 리포트", "subtitle": "사내용 · 고객 제출 금지", "meta": ["2026-10-06", "Winmate"],
    "sections": [
        {"heading": "요약", "level": 1, "paragraphs": ["시장은 성장 중이다 [확인 필요].", {"text": "굵은 문단", "bold": True}],
         "bullets": ["첫째", "둘째"], "numbered": ["하나", "둘"]},
        {"heading": "비교표", "level": 2,
         "table": {"columns": ["항목", "삼성", "경쟁사 A"], "rows": [["밝기", "700nit", "500nit"], ["가격", "[확인 필요]", "—"]],
                   "highlight_rows": [0], "caption": "표 1. 주요 비교"},
         "images": ["file_IMG", {"file_id": "file_AI", "caption": "현장 시안"}], "sources": [{"label": "출처 A", "url": "https://a"}]},
        {"heading": "부록", "level": 1, "page_break": True, "paragraphs": [{"ko": "한국어", "en": "English [확정 필요]"}]},
    ],
}


def _png(w=400, h=300, color=(30, 60, 160)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return buf.getvalue()


ASSETS = DictAssets({"file_IMG": Asset(_png(), "image/png", {}),
                     "file_AI": Asset(_png(color=(160, 40, 40)), "image/png", {"meta": {"is_generated": True}})})


def test_xlsx():
    data = build_xlsx(SHEETS, design={"brand_hex": "#e4002b"}, confidential=True)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["제품 수량 1", "제품 수량 1 (2)"]
    ws = wb.worksheets[0]
    assert ws["A1"].value == "1층 로비 제품 수량"
    assert "A1:D1" in [str(m) for m in ws.merged_cells.ranges]
    header = [c.value for c in ws[4]]
    assert header == ["공간", "모델코드", "수량", "수량 근거"]
    assert ws["A4"].fill.fgColor.rgb.endswith("E4002B")              # 포인트 색 머리 행
    assert ws.freeze_panes == "A5"
    assert ws["B6"].comment is not None and ws["B6"].font.b
    assert ws["D6"].value == "[확인 필요]" and ws["D6"].fill.fgColor.rgb.endswith("FFF2A8")
    assert ws["C7"].value == 1200 and ws["C7"].number_format == "#,##0"
    assert ws["D7"].value == "'=SUM(A1)"                              # 수식은 값으로
    assert ws.auto_filter.ref == "A4:D7"
    assert any("출처" in str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert ws.oddHeader.right.text == "CONFIDENTIAL"
    en = load_workbook(io.BytesIO(build_xlsx(SHEETS, lang="en"))).worksheets[0]
    assert en["D6"].value == "[TBD]"
    with pytest.raises(XlsxError):
        build_xlsx({"sheets": []})


def test_xlsx_strips_control_characters_and_caps_length():
    """모델 · PDF 추출 글의 제어 문자는 openpyxl 이 오류를 내 500 이 되던 것 — 빼고 넣는다. 칸 한도 32,767자."""
    data = build_xlsx({"sheets": [{"name": "시트\x07", "rows": [{"a": "벨\x07소리", "b": {"value": "x", "comment": "메\x01모"}, "c": "y" * 40000}]}]})
    ws = load_workbook(io.BytesIO(data)).worksheets[0]
    assert ws.title == "시트" and ws["A2"].value == "벨소리" and ws["B2"].comment.text == "메모" and len(ws["C2"].value) == 32767


def test_docx():
    data, warnings = build_docx(REPORT, assets=ASSETS, confidential=True)
    assert not warnings
    d = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "경쟁사 분석 리포트" in text and "요약" in text and "비교표" in text
    assert "AI 생성 이미지" in text                                   # 생성 이미지 캡션
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "700nit"
    assert len(d.inline_shapes) == 2
    assert "CONFIDENTIAL" in d.sections[0].header.paragraphs[0].text
    xml = d.element.xml
    assert 'w:highlight w:val="yellow"' in xml
    en, _ = build_docx(REPORT, lang="en", assets=ASSETS)
    d_en = Document(io.BytesIO(en))
    t_en = "\n".join(p.text for p in d_en.paragraphs) + "\n".join(c.text for r in d_en.tables[0].rows for c in r.cells)
    assert "[TBD]" in t_en and "[확인 필요]" not in t_en and "English [TBD]" in t_en


@pytest.mark.parametrize("page_size,expect", [("A4", (595, 842)), ("Letter", (612, 792))])
def test_pdf_report(page_size, expect):
    data, warnings = build_pdf({"report": REPORT, "page_size": page_size}, assets=ASSETS)
    assert data[:5] == b"%PDF-" and not warnings
    pdf = pdfium.PdfDocument(data)
    try:
        assert len(pdf) >= 2                                        # page_break
        w, h = pdf[0].get_size()
        assert (round(w), round(h)) == expect
        assert b"HYGothic-Medium" in data
    finally:
        pdf.close()


def test_pdf_landscape_and_missing_image():
    doc = {"title": "가로", "orientation": "landscape", "sections": [{"heading": "그림", "images": ["file_NOPE"]}]}
    data, warnings = build_pdf(doc, assets=DictAssets())
    pdf = pdfium.PdfDocument(data)
    w, h = pdf[0].get_size()
    pdf.close()
    assert w > h and warnings == ["image missing file_NOPE"]


def test_zip():
    files = {"file_A": b"\x89PNG fake", "file_B": b"hello"}
    doc = {"entries": [
        {"file_id": "file_A", "path": "images/cut_01.png"},
        {"file_id": "file_B", "path": "images/cut_01.png"},           # 겹치면 (2)
        {"path": "sources.json", "json": {"cuts": [1, 2]}},
        {"path": "readme.txt", "text": "조감도"},
        {"path": "수량표.xlsx", "format": "xlsx", "document": {"sheets": [{"name": "제품", "rows": [["a", 1]]}]}},
    ]}

    def nested(fmt, d):
        assert fmt == "xlsx"
        return build_xlsx(d)

    data = build_zip(doc, files=files, nested=nested)
    z = zipfile.ZipFile(io.BytesIO(data))
    assert z.namelist() == ["images/cut_01.png", "images/cut_01 (2).png", "sources.json", "readme.txt", "수량표.xlsx"]
    assert json.loads(z.read("sources.json")) == {"cuts": [1, 2]}
    assert load_workbook(io.BytesIO(z.read("수량표.xlsx"))).sheetnames == ["제품"]
    for bad in ("/etc/passwd", "../x.txt", "a/../../x", "C:/x"):
        with pytest.raises(ZipError):
            safe_path(bad)
    with pytest.raises(ZipError):
        build_zip({"entries": [{"file_id": "file_X", "path": "x"}]}, files={})
