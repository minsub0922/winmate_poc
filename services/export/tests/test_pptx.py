"""PPTX 렌더러 — 모든 템플릿이 열리는 PPTX 가 되는지, 포인트 색 · 마스터 · 영어 · 표시 · AI 라벨."""
from __future__ import annotations

import io
import re
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from winmate_export.render.pptx_render import (
    Asset,
    DeckRenderer,
    DictAssets,
    open_master,
)
from winmate_export.render.theme import Theme
from winmate_export.templates.catalog import catalog
from winmate_export.templates.samples import sample_slots

IMG, AI_IMG, LOGO = "file_IMG", "file_AIIMG", "file_LOGO"


def png_bytes(w: int = 640, h: int = 400, color: tuple[int, int, int] = (40, 80, 200)) -> bytes:
    im = Image.new("RGB", (w, h), color)
    ImageDraw.Draw(im).rectangle([w // 4, h // 4, 3 * w // 4, 3 * h // 4], fill=(240, 200, 60))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def _assets() -> DictAssets:
    return DictAssets({
        IMG: Asset(png_bytes(), "image/png", {"meta": {}}),
        AI_IMG: Asset(png_bytes(800, 500, (200, 60, 60)), "image/png", {"meta": {"generation": {"is_generated": True, "rights": "generated"}}}),
        LOGO: Asset(png_bytes(300, 80, (20, 20, 20)), "image/png", {}),
    })


def _texts(slide) -> str:
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    out.append(cell.text)
    return "\n".join(out)


def _pictures(slide) -> int:
    return sum(1 for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)


def _render(doc, **kw):
    design = kw.pop("design", {})
    r = DeckRenderer(theme=Theme.from_design(design), assets=_assets(), design=design, **kw)
    data = r.render(doc)
    return r, Presentation(io.BytesIO(data)), data


def test_every_template_renders():
    templates = [t for t in catalog().templates]
    doc = {"title": "전체", "slides": [{"template_code": t["code"], "slots": sample_slots(t, IMG, LOGO, title=f"T:{t['code']}")}
                                     for t in templates]}
    r, prs, _ = _render(doc, design={"logo_file_id": LOGO})
    assert len(prs.slides) == len(templates)
    assert prs.slide_width == 12192000 and prs.slide_height == 6858000      # 13.333 × 7.5in
    assert not [w for w in r.warnings if "Error" in w or "failed" in w], r.warnings[:10]
    for t, slide in zip(templates, prs.slides):
        text = _texts(slide)
        if any(s["key"] == "title" for s in t["slots"]):
            assert f"T:{t['code']}" in text, t["code"]
        has_img = any(s["type"] == "image" for s in t["slots"]) or any(
            f["type"] == "image" for s in t["slots"] for f in s.get("fields", []))
        if has_img:
            assert _pictures(slide) >= 1, t["code"]


def test_brand_colour_applied():
    t = catalog().get("MS-A")
    doc = {"slides": [{"template_code": "MS-A", "slots": sample_slots(t, IMG)}]}
    _, _, data = _render(doc, design={"brand_hex": "#E4002B"})
    xml = zipfile.ZipFile(io.BytesIO(data)).read("ppt/slides/slide1.xml").decode()
    assert "E4002B" in xml.upper()
    assert "1428A0" not in xml.upper()          # 기본 삼성 블루가 남지 않는다


def test_markers_highlighted_and_english_tbd():
    doc = {"slides": [{"template_code": "CB-A", "slots": {"title": "점유율 [확인 필요] 확보",
                                                          "summary": {"ko": "수치 [확정 필요]", "en": "Figure [확정 필요]"}},
                       "confirm": ["점유율 근거"]}]}
    _, prs, data = _render(doc)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("ppt/slides/slide1.xml").decode()
    assert "<a:highlight>" in xml and "[확인 필요]" in xml
    _, prs_en, data_en = _render(doc, lang="en")
    text = _texts(prs_en.slides[0])
    assert "[TBD]" in text and "[확인 필요]" not in text and "[확정 필요]" not in text
    assert "Figure [TBD]" in text
    notes = prs_en.slides[0].notes_slide.notes_text_frame.text
    assert "[TBD] 점유율 근거" in notes
    _, prs_lbl, _ = _render(doc, lang="en", tbd_label="[To be confirmed]")
    assert "[To be confirmed]" in _texts(prs_lbl.slides[0])


def test_tbd_mode_notes_moves_text():
    doc = {"slides": [{"template_code": "CB-A", "slots": {"title": "제목", "summary": "매출 [확인 필요] 증가"}}]}
    _, prs, _ = _render(doc, tbd_mode="notes")
    s = prs.slides[0]
    assert "[확인 필요]" not in _texts(s)
    assert "매출" in s.notes_slide.notes_text_frame.text and "[확인 필요]" in s.notes_slide.notes_text_frame.text


def test_ai_label_and_sources_and_notes():
    t = catalog().get("C01")
    doc = {"cover": None, "slides": [
        {"template_code": "BV-A", "slots": {"title": "조감도", "image": {"file_id": AI_IMG}}, "notes": "발표 메모",
         "sources": [{"label": "KB 리포트", "url": "https://example.com/r"}]},
        {"template_code": "C01", "slots": sample_slots(t, IMG), "sources": ["표지 출처"]},
    ]}
    _, prs, _ = _render(doc)
    s1, s2 = prs.slides
    assert "AI 생성 이미지" in _texts(s1)
    assert "출처: KB 리포트" in _texts(s1)
    assert "발표 메모" in s1.notes_slide.notes_text_frame.text
    # 출처 상자가 없는 표지는 노트로
    assert "표지 출처" in s2.notes_slide.notes_text_frame.text
    _, prs_en, _ = _render(doc, lang="en")
    assert "AI-generated image" in _texts(prs_en.slides[0])


def test_cover_inserted_and_kind_defaults():
    doc = {"title": "제안", "cover": {"title": "A 커피 디지털 메뉴보드", "customer": "A 커피", "date": "2026-10-06"},
           "slides": [{"kind": "toc", "slots": {"title": "목차", "items": [{"no": "01", "title": "시장"}]}},
                      {"kind": "divider", "slots": {"no": "01", "title": "시장"}},
                      {"template_code": "MS-B", "slots": {"title": "성장"}},
                      {"kind": "closing", "slots": {"title": "감사합니다"}}]}
    r, prs, _ = _render(doc)
    assert r.template_codes == ["C03", "C04", "C05", "MS-B", "C07"]
    assert "A 커피 디지털 메뉴보드" in _texts(prs.slides[0])
    # 바닥글 기본값 = 고객사 · 제목
    assert "A 커피 · A 커피 디지털 메뉴보드" in _texts(prs.slides[3])


def test_bilingual_slides_and_inline():
    doc = {"slides": [{"template_code": "CB-A", "slots": {"title": {"ko": "한국어 제목", "en": "English title"}}}]}
    _, prs, _ = _render(doc, lang="both", bilingual="slides")
    assert len(prs.slides) == 2
    assert "한국어 제목" in _texts(prs.slides[0]) and "English title" in _texts(prs.slides[1])
    _, prs, _ = _render(doc, lang="both", bilingual="inline")
    assert len(prs.slides) == 1
    t = _texts(prs.slides[0])
    assert "한국어 제목" in t and "English title" in t


def test_text_fits_by_truncation():
    long = "가" * 400
    doc = {"slides": [{"template_code": "CB-A", "slots": {"title": long}}]}
    _, prs, _ = _render(doc)
    title_shape = next(sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.startswith("가"))
    assert len(title_shape.text_frame.text) <= 47 and title_shape.text_frame.text.endswith("…")


def _potx_bytes() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "기존 슬라이드"
    buf = io.BytesIO()
    prs.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            z.writestr(item, data)
    return out.getvalue()


def test_master_potx_and_pptx():
    potx = _potx_bytes()
    assert b"template.main+xml" in zipfile.ZipFile(io.BytesIO(potx)).read("[Content_Types].xml")
    prs = open_master(potx)
    assert len(prs.slides) == 0
    doc = {"slides": [{"template_code": "MS-B", "slots": {"title": "마스터 위"}}]}
    r, out, data = _render(doc, design={"_master_bytes": potx})
    assert len(out.slides) == 1
    assert "마스터 위" in _texts(out.slides[0])
    assert out.slides[0].slide_layout.name.lower().startswith("blank")
    ct = zipfile.ZipFile(io.BytesIO(data)).read("[Content_Types].xml")
    assert b"presentation.main+xml" in ct and b"template.main+xml" not in ct


def test_invalid_master():
    import pytest
    from winmate_export.render.pptx_render import RenderError

    with pytest.raises(RenderError) as ei:
        open_master(b"not a zip")
    assert ei.value.code == "INVALID_MASTER"


def test_fonts_and_east_asian_typeface():
    doc = {"slides": [{"template_code": "CB-A", "slots": {"title": "글꼴"}}]}
    _, _, data = _render(doc)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("ppt/slides/slide1.xml").decode()
    assert re.search(r'<a:ea typeface="Noto Sans KR"', xml)
