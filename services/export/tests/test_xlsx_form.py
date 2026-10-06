"""고객사 XLSX 양식 채우기 — 맞추기(이름 · 머리 · 방향 · 추론) · 쓰기 규칙 · 보존 · API(base_file_id · fills · form · form_fill).

양식은 openpyxl 로 만든 견본(사양 회신서 · 여러 줄 머리 · 제품이 행인 견적서 · 「이름 | 값」 · 보호 시트 · 매크로 · 도형).
"""
from __future__ import annotations

import io
import os
import re
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XImage
from openpyxl.styles import Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image
from winmate_common import testing
from winmate_export.render.xlsx_form import (
    FormError,
    LabelIndex,
    Opts,
    _keys,
    fill_form,
    inventory,
    open_workbook,
    validate_fills,
    variants,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSM = "application/vnd.ms-excel.sheet.macroEnabled.12"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin")


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _open(data: bytes, **kw):
    return load_workbook(io.BytesIO(data), **kw)


def _rezip(data: bytes, *, patch: dict[str, object] | None = None, add: dict[str, bytes] | None = None) -> bytes:
    """zip 안 부품 고치기(patch: 이름 → 함수(bytes) → bytes) · 더하기."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            d = zin.read(n)
            if patch and n in patch:
                d = patch[n](d)  # type: ignore[operator]
            zout.writestr(n, d)
        for n, d in (add or {}).items():
            zout.writestr(n, d)
    return out.getvalue()


def _png() -> io.BytesIO:
    b = io.BytesIO()
    Image.new("RGB", (60, 30), (200, 0, 0)).save(b, "PNG")
    b.seek(0)
    return b


# ── 견본 양식 ──────────────────────────────────────────────

def reply_form() -> bytes:
    """세로 사양 회신서 — 제목 병합 · 머리 칠 · 요구/제안/비고 열 · 수식 합계 · 병합 값 칸 · 메모 · 그림 · 목록 유효성."""
    wb = Workbook()
    ws = wb.active
    ws.title = "사양 회신서"
    ws["A1"] = "제품 사양 회신서"
    ws.merge_cells("A1:E1")
    ws["A1"].font = Font(size=16, bold=True)
    for c, v in zip("ABCDE", ["No", "항목", "요구 사양", "제안 사양", "비고"]):
        ws[f"{c}3"] = v
        ws[f"{c}3"].fill = HEAD_FILL
    data = [(1, "화면 크기", '55" 이상', None), (2, "밝기 (cd/m²)", "500 이상", None), (3, "소비전력 [W]", "", "-"),
            (4, "수량", 10, None), (5, "단가", None, None), (6, "합계", None, "=D7*D8"), (7, "Remarks", None, None)]
    for i, (n, lab, req, val) in enumerate(data):
        r = 4 + i
        ws.cell(r, 1, n)
        ws.cell(r, 2, lab)
        ws.cell(r, 3, req)
        ws.cell(r, 4, val)
    ws["D7"].number_format = "#,##0"
    ws.merge_cells("D10:E10")
    ws.column_dimensions["B"].width = 28
    ws["B4"].comment = Comment("고객 메모", "고객")
    dv = DataValidation(type="list", formula1='"충족,미충족"')
    ws.add_data_validation(dv)
    dv.add("E4:E9")
    ws.add_image(XImage(_png()), "G1")
    return _bytes(wb)


SPEC_DOC = {"sheets": [{"name": "Spec", "columns": [{"key": "item", "label": "항목"}, {"key": "m1", "label": "QM55C"}],
                        "rows": [{"item": "화면 크기", "m1": '55"'}, {"item": "밝기(cd/m2)", "m1": 500}, {"item": "소비 전력", "m1": "150 W"},
                                 {"item": "수량", "m1": "1,200"}, {"item": "합계", "m1": 999}, {"item": "무게", "m1": "20 kg"},
                                 {"item": "Remarks", "m1": "[확정 필요]"}]}]}


def multi_header_form() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A3"] = "항목"
    ws.merge_cells("A3:A4")
    ws["B3"] = "제안 제품"
    ws.merge_cells("B3:D3")
    for c, v in zip("BCD", ["QM55C", "QM65C", "QM75C"]):
        ws[f"{c}4"] = v
    for i, lab in enumerate(["화면 크기", "밝기", "무게"]):
        ws.cell(5 + i, 1, lab)
    return _bytes(wb)


def quote_form() -> bytes:
    """제품이 행인 견적서 — 금액 · 합계는 수식."""
    wb = Workbook()
    ws = wb.active
    ws.title = "견적"
    for c, v in zip("ABCDE", ["No", "모델명", "수량", "단가", "금액"]):
        ws[f"{c}3"] = v
    ws["A4"], ws["B4"], ws["E4"] = 1, "QM55C", "=C4*D4"
    ws["A5"], ws["B5"], ws["E5"] = 2, "QM65C", "=C5*D5"
    ws["A6"] = "합계"
    ws.merge_cells("A6:D6")
    ws["E6"] = "=SUM(E4:E5)"
    return _bytes(wb)


# ── 이름 정규화 ────────────────────────────────────────────

def test_label_keys_and_variants():
    k = _keys
    assert k("밝기 (cd/m²)")[2] == k("밝기(cd/m2)")[2] == k("밝기")[0]          # base: 단위 괄호 무시
    assert k("밝기 (cd/m²)")[1] == k("밝기(cd/m2)")[1]                           # normalized: ² → 2 · 공백 · 기호
    assert k("1. 화면 크기")[1] == k("화면크기")[1] == k("① 화면 크기 :")[1] == k("가. 화면 크기")[1] == k("- 화면 크기")[1]
    assert k("Brightness*")[0] == k("brightness")[0]                              # 끝의 필수 표시 · 대소문자
    assert k("2.5kg")[1] != k("5kg")[1] and k("4K 해상도")[1] == "4k해상도"          # 숫자는 번호 매김으로 지우지 않는다
    assert k("최대 소비전력")[2] != k("소비전력")[2]
    assert set(variants("밝기 / Brightness")) >= {"밝기", "Brightness"}
    assert set(variants("소비 전력\nPower consumption")) >= {"소비 전력", "Power consumption"}
    assert set(variants("화면 크기 (Screen size)")) >= {"화면 크기", "Screen size"}
    assert variants("소비전력 (W)") == ["소비전력 (W)"]                             # 단위는 따로 보지 않는다


def test_label_index_levels_fuzzy_and_ambiguity():
    idx = LabelIndex()
    for i, t in enumerate(["밝기 (cd/m²)", "Brigthness", "Power consumption max", "Power consumption min", "최대 소비전력",
                           "소비 전력 [W]"]):
        idx.add((i + 1, 1), t)
    def anywhere(p):
        return True

    assert idx.find(["소비전력"], anywhere) == ("base", [(6, 1)])            # 대괄호 단위 무시
    assert idx.find(["밝기(cd/m2)"], anywhere) == ("normalized", [(1, 1)])
    assert idx.find(["밝기"], anywhere) == ("normalized", [(1, 1)])          # 「밝기 (cd/m²)」의 바깥 이름(통째 exact 보다 약하게)
    assert idx.find(["밝기"], lambda p: p[0] > 1) == (None, [])
    assert idx.find_one(["Brightness"], anywhere, fuzzy=False) == (None, [])
    assert idx.find_one(["Brightness"], anywhere, fuzzy=True) == ("fuzzy", [(2, 1)])
    assert idx.find_one(["Power consumption mix"], anywhere, fuzzy=True) == (None, [])   # 비슷한 후보 둘 — 고르지 않는다
    assert idx.find_one(["소비전력"], lambda p: p[0] != 6, fuzzy=True) == (None, [])     # 「최대 소비전력」과 맞추지 않는다


# ── 표 맞추기 ──────────────────────────────────────────────

def test_reply_form_fills_inferred_column_and_keeps_everything_else():
    out, rep, warns = fill_form(reply_form(), document=SPEC_DOC)
    wb = _open(out)
    s = wb["사양 회신서"]
    # 값: 제안 사양(D) 열 — 숫자는 숫자로, '-' 자리표시는 덮고, 수식 칸은 건너뛰고, 병합 칸은 머리칸에
    assert [s[f"D{r}"].value for r in range(4, 11)] == ['55"', 500, "150 W", 1200, None, "=D7*D8", "[확정 필요]"]
    assert s["D10"].fill.fgColor.rgb.endswith("FFF2A8") and s["D10"].font.b                  # 자리표시 노란 칠
    # 양식은 그대로: 요구 사양 · 제목 병합 · 머리 칠 · 열 너비 · 서식 · 메모 · 그림 · 유효성 · 수식
    assert s["C4"].value == '55" 이상' and s["E4"].value is None
    assert {str(m) for m in s.merged_cells.ranges} == {"A1:E1", "D10:E10"}
    assert s["C3"].fill.fgColor.rgb.endswith("DDEBF7") and s["A1"].font.sz == 16
    assert s.column_dimensions["B"].width == 28 and s["D7"].number_format == "#,##0"
    assert s["B4"].comment.text == "고객 메모" and len(s._images) == 1
    assert len(s.data_validations.dataValidation) == 1
    assert wb.calculation.fullCalcOnLoad is True
    # 보고
    sh = rep["sheets"][0]
    assert (sh["header_row"], sh["label_column"], sh["orientation"]) == (3, "B", "rows")
    assert sh["columns"] == [{"key": "m1", "label": "QM55C", "column": "D", "match": "inferred"}]
    assert rep["filled"] == 5 and sh["rows_matched"] == 6
    assert [x["label"] for x in rep["unmatched_rows"]] == ["무게"]
    assert [(x["cell"], x["reason"]) for x in rep["skipped"]] == [("D9", "formula")]
    assert rep["form_only_rows"] == [{"sheet": "사양 회신서", "row": 8, "label": "단가"}]
    assert rep["lost"] == []
    cells = {c["cell"]: c for c in rep["cells"]}
    assert cells["D5"]["row_match"] == "normalized" and cells["D6"]["row_match"] == "base"
    assert any("무게" in w for w in warns) and any("수식 1" in w for w in warns)


def test_multi_level_header_and_unknown_product_column():
    doc = {"sheets": [{"name": "x", "columns": [{"key": "item", "label": "항목"}, {"key": "a", "label": "QM55C"},
                                                {"key": "b", "label": "QM65C"}, {"key": "c", "label": "QM85C"}],
                       "rows": [{"item": "화면 크기", "a": 55, "b": 65, "c": 85}, {"item": "밝기", "a": 500, "b": 500, "c": 700}]}]}
    out, rep, _ = fill_form(multi_header_form(), document=doc)
    s = _open(out).active
    assert [[s.cell(r, c).value for c in range(2, 5)] for r in (5, 6, 7)] == [[55, 65, None], [500, 500, None], [None, None, None]]
    assert rep["sheets"][0]["header_row"] == 4 and rep["sheets"][0]["label_column"] == "A"
    assert [(x["key"], x["label"]) for x in rep["unmatched_columns"]] == [("c", "QM85C")]
    assert rep["form_only_rows"] == [{"sheet": "Sheet", "row": 7, "label": "무게"}]


def test_transposed_quote_form_keeps_formulas():
    doc = {"sheets": [{"name": "spec", "columns": [{"key": "item", "label": "항목"}, {"key": "QM55C", "label": "QM55C"},
                                                   {"key": "QM65C", "label": "QM65C"}],
                       "rows": [{"item": "수량", "QM55C": 2, "QM65C": 3}, {"item": "단가", "QM55C": 1500000, "QM65C": "1,900,000"}]}]}
    out, rep, _ = fill_form(quote_form(), document=doc)
    s = _open(out)["견적"]
    assert [[s.cell(r, c).value for c in range(3, 6)] for r in (4, 5)] == [[2, 1500000, "=C4*D4"], [3, 1900000, "=C5*D5"]]
    assert s["E6"].value == "=SUM(E4:E5)" and rep["sheets"][0]["orientation"] == "columns" and rep["filled"] == 4
    # 방향을 rows 로 못 박으면 맞는 칸이 없다
    _, rep2, _ = fill_form(quote_form(), document=doc, options={"orientation": "rows"})
    assert rep2["filled"] == 0 and rep2["sheets"][0]["orientation"] == "rows"


def test_ambiguous_qualified_labels_are_not_guessed():
    wb = Workbook()
    ws = wb.active
    ws.append(["항목", "제안"])
    ws.append(["소비 전력 (최대)", None])
    ws.append(["소비 전력 (대기)", None])
    doc = {"sheets": [{"name": "s", "columns": [{"key": "i", "label": "항목"}, {"key": "v", "label": "제안"}],
                       "rows": [{"i": "소비전력(정격)", "v": "120W"}, {"i": "소비 전력 (대기)", "v": "0.5W"}]}]}
    out, rep, _ = fill_form(_bytes(wb), document=doc)
    s = _open(out).active
    assert (s["B2"].value, s["B3"].value) == (None, "0.5W")
    assert rep["unmatched_rows"] == [{"sheet": "Sheet", "label": "소비전력(정격)", "source": "sheets[0].rows[0]",
                                      "candidates": ["A2 소비 전력 (최대)", "A3 소비 전력 (대기)"]}]
    _, rep, _ = fill_form(_bytes(wb), fills=[{"row": "소비전력", "value": 1}])
    assert rep["filled"] == 0 and len(rep["unmatched_rows"][0]["candidates"]) == 2


def test_duplicate_labels_follow_document_order():
    wb = Workbook()
    ws = wb.active
    for r, (a, b) in enumerate([("항목", "제안"), ("디스플레이", None), ("밝기", None), ("비고", None), ("전원", None),
                                ("소비 전력", None), ("비고", None)], start=1):
        ws.cell(r, 1, a)
        ws.cell(r, 2, b)
    doc = {"sheets": [{"name": "s", "columns": [{"key": "i", "label": "항목"}, {"key": "v", "label": "제안"}],
                       "rows": [{"i": "밝기", "v": 500}, {"i": "비고", "v": "A"}, {"i": "소비 전력", "v": "150W"}, {"i": "비고", "v": "B"}]}]}
    s = _open(fill_form(_bytes(wb), document=doc)[0]).active
    assert [s.cell(r, 2).value for r in range(3, 8)] == [500, "A", None, "150W", "B"]


def test_hints_column_map_row_map_and_form_row():
    wb = Workbook()
    ws = wb.active
    for c, v in zip("ABCD", ["구분", "사양", "제안1", "제안2"]):
        ws[f"{c}2"] = v
    ws["A3"], ws["A4"], ws["A5"] = "Brightness", "Power", "Weight"
    doc = {"sheets": [{"name": "s", "columns": [{"key": "no", "label": "No"}, {"key": "item", "label": "항목"},
                                                {"key": "a", "label": "QM55C"}, {"key": "b", "label": "QM65C", "form_column": "D"}],
                       "rows": [{"no": 1, "item": "Luminance", "a": 500, "b": 700}, {"no": 2, "item": "전력", "a": 100, "b": 120, "_form_row": 4},
                                {"no": 3, "item": "Weight", "a": "20kg", "b": "25kg"}]}]}
    opts = {"header_row": 2, "column_map": {"a": "제안1"}, "row_map": {"Luminance": "Brightness"}}
    out, rep, _ = fill_form(_bytes(wb), document=doc, options=opts)
    s = _open(out).active
    assert [[s.cell(r, c).value for c in (3, 4)] for r in (3, 4, 5)] == [[500, 700], [100, 120], ["20kg", "25kg"]]
    cols = {c["key"]: c for c in rep["sheets"][0]["columns"]}
    assert cols["a"]["match"] == cols["b"]["match"] == "hint" and cols["a"]["column"] == "C"
    assert "no" in cols and cols["no"]["column"] is None          # 번호 열(No)은 행 이름 열로 고르지 않는다(글이 든 첫 열 = 항목)
    assert rep["sheets"][0]["orientation"] == "rows"
    # 열 번호 · 행 번호 힌트 · label_key
    out, rep, _ = fill_form(_bytes(wb), document=doc, options={"label_key": "item", "header_row": 2, "label_column": "A",
                                                                "column_map": {"a": 3, "b": "D"}, "row_map": {"Luminance": 3}})
    assert _open(out).active["C3"].value == 500 and rep["sheets"][0]["label_column"] == "A"
    with pytest.raises(FormError) as e:
        fill_form(_bytes(wb), document=doc, options={"label_key": "nope"})
    assert e.value.code == "VALIDATION_FAILED"


def test_overwrite_policy_and_formulas():
    wb = Workbook()
    ws = wb.active
    ws.append(["항목", "제안"])
    ws.append(["밝기", "기존 값"])
    ws.append(["전력", "=1+1"])
    ws.append(["무게", "(기재)"])
    base = _bytes(wb)
    doc = {"sheets": [{"name": "s", "columns": [{"key": "i", "label": "항목"}, {"key": "v", "label": "제안"}],
                       "rows": [{"i": "밝기", "v": 600}, {"i": "전력", "v": 100}, {"i": "무게", "v": "20kg"}]}]}
    out, rep, _ = fill_form(base, document=doc)
    s = _open(out).active
    assert (s["B2"].value, s["B3"].value, s["B4"].value) == ("기존 값", "=1+1", "20kg")
    assert sorted((x["cell"], x["reason"], x.get("current")) for x in rep["skipped"]) == [("B2", "not_empty", "기존 값"), ("B3", "formula", None)]
    out, rep, _ = fill_form(base, document=doc, options={"overwrite": "always"})
    assert _open(out).active["B2"].value == 600 and {c["cell"]: c.get("replaced") for c in rep["cells"]}["B2"] == "기존 값"
    out, rep, _ = fill_form(base, document=doc, options={"overwrite": "always", "overwrite_formulas": True})
    assert _open(out).active["B3"].value == 100 and not rep["skipped"]


def test_append_unmatched_rows_and_extra_sheets():
    wb = Workbook()
    ws = wb.active
    ws.title = "회신"
    for c, v in zip("ABC", ["항목", "제안 사양", "충족 여부"]):
        ws[f"{c}2"] = v
        ws[f"{c}2"].fill = HEAD_FILL
    for i, lab in enumerate(["밝기", "화면 크기"]):
        ws.cell(3 + i, 1, lab)
        for c in range(1, 4):
            ws.cell(3 + i, c).border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    dv = DataValidation(type="list", formula1='"충족,미충족"')
    ws.add_data_validation(dv)
    dv.add("C3:C10")
    ws["A7"], ws["B7"] = "작성일:", "2026-10-07"
    doc = {"sheets": [{"name": "회신", "columns": [{"key": "i", "label": "항목"}, {"key": "v", "label": "제안 사양"},
                                                 {"key": "ok", "label": "충족 여부"}],
                       "rows": [{"i": "밝기", "v": 500, "ok": "충족"}, {"i": "화면 크기", "v": "55", "ok": "Yes"}, {"i": "무게", "v": "20 kg", "ok": "충족"}]},
                      {"name": "메모 · 출처", "columns": [{"key": "a", "label": "출처"}], "rows": [{"a": "KB"}]}]}
    out, rep, warns = fill_form(_bytes(wb), document=doc, options={"unmatched_rows": "append", "extra_sheets": "append"})
    w2 = _open(out)
    s = w2["회신"]
    assert [s.cell(9, c).value for c in (1, 2, 3)] == ["추가 항목", None, None]                  # 양식 맨 아래(작성일 + 2)
    assert [s.cell(10, c).value for c in (1, 2, 3)] == ["무게", "20 kg", "충족"]
    assert s["A10"].border.top.style == "thin" and s["A9"].fill.fgColor.rgb.endswith("DDEBF7")   # 맞춘 행 · 머리 서식 복사
    assert w2.sheetnames == ["회신", "메모 · 출처"] and rep["appended_sheets"] == ["메모 · 출처"]
    assert rep["appended_rows"] == [{"sheet": "회신", "row": 10, "label": "무게", "source": "sheets[0].rows[2]"}]
    assert {c["cell"]: c.get("warning") for c in rep["cells"]}["C4"].startswith("목록에 없는 값")
    assert any("덧붙였습니다" in w for w in warns) and any("드롭다운" in w for w in warns)
    _, rep, warns = fill_form(_bytes(wb), document=doc)
    assert rep["unmatched_sheets"] == ["메모 · 출처"] and rep["appended_rows"] == [] and any("짝이 없는" in w for w in warns)


def test_bilingual_labels_and_english_values():
    wb = Workbook()
    ws = wb.active
    ws.append(["항목 / Item", "제안 / Proposal"])
    ws.append(["밝기 / Brightness", None])
    ws.append(["소비 전력\nPower consumption", None])
    ws.append(["화면 크기 (Screen size)", None])
    doc = {"sheets": [{"name": "s", "columns": [{"key": "i", "label": {"ko": "항목", "en": "Item"}},
                                                {"key": "v", "label": {"ko": "제안", "en": "Proposal"}}],
                       "rows": [{"i": {"ko": "밝기", "en": "Brightness"}, "v": {"ko": "500 니트", "en": "500 nit"}},
                                {"i": "Power consumption", "v": "[확인 필요]"}, {"i": "Screen size", "v": 55}]}]}
    s = _open(fill_form(_bytes(wb), document=doc, lang="en")[0]).active
    assert (s["B2"].value, s["B3"].value, s["B4"].value) == ("500 nit", "[TBD]", 55)
    assert s["B3"].fill.fgColor.rgb.endswith("FFF2A8")
    s = _open(fill_form(_bytes(wb), document=doc, lang="ko", options={"highlight_marks": False})[0]).active
    assert (s["B2"].value, s["B3"].value) == ("500 니트", "[확인 필요]") and s["B3"].fill.fgColor.rgb in ("00000000", None)


def test_fuzzy_match_reported_and_can_be_disabled():
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Proposal"])
    ws.append(["Brigthness", None])
    ws.append(["Contrast ratio", None])
    doc = {"sheets": [{"name": "s", "columns": [{"key": "i", "label": "Item"}, {"key": "v", "label": "Proposal"}],
                       "rows": [{"i": "Brightness", "v": 500}, {"i": "Contrast ratio", "v": "4000:1"}]}]}
    out, rep, warns = fill_form(_bytes(wb), document=doc)
    assert _open(out).active["B2"].value == 500
    assert {c["cell"]: c["match"] for c in rep["cells"]} == {"B2": "fuzzy", "B3": "exact"}
    assert any("비슷한 이름" in w for w in warns)
    out, rep, _ = fill_form(_bytes(wb), document=doc, options={"fuzzy": False})
    assert _open(out).active["B2"].value is None and [x["label"] for x in rep["unmatched_rows"]] == ["Brightness"]


# ── fills ─────────────────────────────────────────────────

def test_fills_by_address_label_and_header():
    wb = Workbook()
    ws = wb.active
    ws.title = "기본 정보"
    for i, lab in enumerate(["업체명", "대표자 :", "제안 모델"]):
        ws.cell(1 + i, 1, lab)
    ws["A5"], ws["B5"], ws["C5"] = "구분", "요구", "응답"
    ws["A6"] = "납기"
    ws.merge_cells("B8:C8")
    ws.merge_cells("B10:B11")
    ws2 = wb.create_sheet("별지")
    ws2["A1"] = "비고"
    fills = [
        {"row": "업체명", "value": "삼성전자"},                                  # 이름 → 오른쪽 칸
        {"row": "대표자", "value": "=홍길동"},                                    # 끝의 ':' 무시, '=' 는 글로
        {"cell": "B3", "value": "QM55C", "note": "모델 확인"},                    # 주소
        {"row": "납기", "column": "응답", "value": "4주"},                         # 이름 × 머리
        {"row": 2, "column": "B", "value": "덮기"},                               # 번호 × 글자 = 덮는다
        {"cell": "C8", "value": "병합 안쪽"},                                     # 같은 행 병합 → 머리칸 B8
        {"cell": "B11", "value": "다른 행 병합"},                                 # 다른 행에 걸친 병합 → 건너뜀
        {"cell": "H40", "value": "밖"},                                           # 양식 범위 밖 — 쓰고 알린다
        {"sheet": "별지", "row": "비고", "value": "별지 메모"},
        {"sheet": 1, "cell": "C1", "value": "번호로"},
        {"sheet": "없는 시트", "cell": "A1", "value": 1},
        {"row": "없는 행", "value": 1},
        {"row": "납기", "column": "없는 머리", "value": 1},
        {"row": "업체명", "value": None},                                         # null 은 쓰지 않는다
    ]
    out, rep, warns = fill_form(_bytes(wb), fills=fills)
    w2 = _open(out)
    s = w2["기본 정보"]
    assert (s["B1"].value, s["B2"].value, s["B2"].data_type, s["B3"].value, s["C6"].value) == ("삼성전자", "덮기", "s", "QM55C", "4주")
    assert s["B3"].comment.text == "모델 확인" and s["B8"].value == "병합 안쪽" and s["B11"].value is None and s["H40"].value == "밖"
    assert w2["별지"]["B1"].value == "별지 메모" and w2["별지"]["C1"].value == "번호로"
    cells = {(c["sheet"], c["cell"], c["source"]): c for c in rep["cells"]}
    assert cells[("기본 정보", "B1", "fills[0]")]["match"] == "next_to_label"
    assert cells[("기본 정보", "C6", "fills[3]")]["match"] == "exact"
    assert cells[("기본 정보", "B8", "fills[5]")]["redirected_from"] == "C8"
    assert cells[("기본 정보", "H40", "fills[7]")]["outside"] is True
    assert {(x["reason"], x["source"]) for x in rep["skipped"]} == {("merged", "fills[6]"), ("sheet_not_found", "fills[10]")}
    assert [x["label"] for x in rep["unmatched_rows"]] == ["없는 행"] and [x["label"] for x in rep["unmatched_columns"]] == ["없는 머리"]
    assert any("범위 밖" in w for w in warns) and any("병합 머리칸" in w for w in warns)
    # fills[1] 이 쓴 B2 를 fills[4](번호 × 글자)가 덮었다 — 칸 수는 한 번
    assert rep["filled"] == len({(c["sheet"], c["cell"]) for c in rep["cells"]})


def test_protected_sheet_writes_only_unlocked_cells():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = "업체명", "제안 모델"
    ws.protection.sheet = True
    ws["B1"].protection = Protection(locked=False)
    out, rep, _ = fill_form(_bytes(wb), fills=[{"row": "업체명", "value": "삼성"}, {"row": "제안 모델", "value": "QM55C"}])
    s = _open(out).active
    assert (s["B1"].value, s["B2"].value) == ("삼성", None) and s.protection.sheet
    assert [(x["cell"], x["reason"]) for x in rep["skipped"]] == [("B2", "locked")]


def test_column_styled_input_cells_inherit_style_and_unlock():
    """보호 시트에서 입력 열(B)을 열 서식으로만 풀어 둔 양식 — 빈 칸에 써도 그 열 서식(테두리 · 숫자 서식 · 잠금 해제)을 잇는다."""
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = "업체명", "단가", "비고"
    ws.protection.sheet = True
    col = ws.column_dimensions["B"]
    col.protection = Protection(locked=False)
    col.border = Border(bottom=THIN)
    col.number_format = "#,##0"
    base = _bytes(wb)
    assert "B2" not in {c.coordinate for row in _open(base).active.iter_rows() for c in row if c.value is not None}
    out, rep, _ = fill_form(base, fills=[{"row": "업체명", "value": "삼성"}, {"row": "단가", "value": "1,500,000"},
                                         {"row": "비고", "value": None}])
    s = _open(out).active
    assert (s["B1"].value, s["B2"].value) == ("삼성", 1500000) and rep["skipped"] == []
    assert s["B2"].number_format == "#,##0" and s["B2"].border.bottom.style == "thin" and s["B2"].protection.locked is False
    assert s["B3"].value is None and s["B3"].has_style is False          # 쓰지 않은 칸은 만들지 않는다


def test_value_conversion_dates_numbers_and_control_chars():
    wb = Workbook()
    ws = wb.active
    for i, lab in enumerate(["납기", "코드", "수량", "메모", "금액"]):
        ws.cell(1 + i, 1, lab)
    ws["B1"].number_format = "yyyy-mm-dd"
    ws["B5"].number_format = "@"                                           # 텍스트 서식 칸은 숫자로 바꾸지 않는다
    out, _, _ = fill_form(_bytes(wb), fills=[{"row": "납기", "value": "2026-11-30"}, {"row": "코드", "value": "007"},
                                             {"row": "수량", "value": "1,200.5"}, {"row": "메모", "value": "줄\x0b바꿈", "note": "n\x01ote"},
                                             {"row": "금액", "value": "1,500"}])
    s = _open(out).active
    assert s["B1"].value.date().isoformat() == "2026-11-30" and s["B1"].number_format == "yyyy-mm-dd"
    assert (s["B2"].value, s["B3"].value, s["B4"].value, s["B4"].comment.text, s["B5"].value) == ("007", 1200.5, "줄바꿈", "note", "1,500")


# ── 보존 · 형식 ────────────────────────────────────────────

_TEXTBOX = (b'<twoCellAnchor><from><col>5</col><colOff>0</colOff><row>1</row><rowOff>0</rowOff></from>'
            b'<to><col>7</col><colOff>0</colOff><row>3</row><rowOff>0</rowOff></to>'
            b'<sp macro="" textlink=""><nvSpPr><cNvPr id="9" name="TextBox 1"/><cNvSpPr txBox="1"/></nvSpPr>'
            b'<spPr><a:prstGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" prst="rect"/></spPr></sp>'
            b'<clientData/></twoCellAnchor>')


def test_lost_parts_are_reported_not_silent():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = "항목", "밝기"
    ws.add_image(XImage(_png()), "D2")
    base = _rezip(_bytes(wb), patch={"xl/drawings/drawing1.xml": lambda d: d.replace(b"</wsDr>", _TEXTBOX + b"</wsDr>")})
    assert inventory(base)["shapes"] == 1
    out, rep, warns = fill_form(base, fills=[{"row": "밝기", "value": 500}])
    assert _open(out).active["B2"].value == 500 and len(_open(out).active._images) == 1   # 그림은 남고
    assert rep["lost"] == [{"part": "shapes", "name": "도형 · 글상자", "base": 1, "output": 0}]  # 글상자는 잃었다고 알린다
    assert any("도형" in w for w in warns)


def test_formula_labels_use_cached_values():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["A2"], ws["A3"] = "항목", "제안", '="밝"&"기"', "무게"
    cached = "밝기".encode()
    base = _rezip(_bytes(wb), patch={"xl/worksheets/sheet1.xml": lambda d: re.sub(
        rb'<c r="A2"([^>]*)><f>(.*?)</f><v\s*/?>(</v>)?</c>',
        lambda m: b'<c r="A2"' + m.group(1) + b' t="str"><f>' + m.group(2) + b"</f><v>" + cached + b"</v></c>", d)})
    doc = {"sheets": [{"name": "s", "columns": [{"key": "i", "label": "항목"}, {"key": "v", "label": "제안"}],
                       "rows": [{"i": "밝기", "v": 500}, {"i": "무게", "v": "20kg"}]}]}
    out, rep, _ = fill_form(base, document=doc)
    s = _open(out).active
    assert (s["B2"].value, s["B3"].value, s["A2"].value) == (500, "20kg", '="밝"&"기"') and rep["unmatched_rows"] == []


def _xlsm() -> bytes:
    """매크로 통합 문서 견본 — vbaProject.bin(가짜) · 매크로 콘텐츠 형식 · 관계."""
    wb = Workbook()
    wb.active["A1"] = "업체명"
    return _rezip(_bytes(wb), add={"xl/vbaProject.bin": b"\xd0\xcf\x11\xe0fake-vba"}, patch={
        "[Content_Types].xml": lambda d: d.replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            b"application/vnd.ms-excel.sheet.macroEnabled.main+xml").replace(
            b"</Types>", b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>'),
        "xl/_rels/workbook.xml.rels": lambda d: d.replace(b"</Relationships>", (
            b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
            b'Target="vbaProject.bin"/></Relationships>')),
    })


def test_macro_workbook_keeps_vba():
    out, rep, _ = fill_form(_xlsm(), fills=[{"row": "업체명", "value": "삼성"}], macro=True)
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        assert "xl/vbaProject.bin" in z.namelist() and b"macroEnabled" in z.read("[Content_Types].xml")
    assert rep["lost"] == [] and _open(out).active["B1"].value == "삼성"


def test_invalid_inputs():
    with pytest.raises(FormError) as e:
        open_workbook(b"\xd0\xcf\x11\xe0" + b"\0" * 100)
    assert e.value.code == "INVALID_BASE_FILE" and "암호" in e.value.message
    with pytest.raises(FormError):
        open_workbook(b"%PDF-1.7")
    with pytest.raises(FormError):
        open_workbook(b"PK\x03\x04broken")
    for bad in ({"orientation": "diagonal"}, {"overwrite": "sometimes"}, {"header_row": 0}, {"label_column": "No such"},
                {"column_map": ["D"]}, "rows"):
        with pytest.raises(FormError):
            Opts.parse(bad)
    for bad in ({"value": 1}, [{"cell": "D0", "value": 1}], [{"cell": "XFE1", "value": 1}], [{"row": 3, "value": 1}],
                [{"row": True, "column": 1, "value": 1}], ["D7"], [{"row": "밝기", "column": 0, "value": 1}]):
        with pytest.raises(FormError):
            validate_fills(bad)
    validate_fills([{"cell": "'시트 1'!$D$7", "value": 1}, {"row": "밝기", "value": 1}, {"row": 3, "column": "D", "value": 1}])
    with pytest.raises(FormError):
        fill_form(reply_form(), document={"sheets": [{"name": "x", "form": {"sheet": "없는 시트"}, "rows": [["a", 1]]}]})


# ── API ───────────────────────────────────────────────────

async def test_api_form_fill_from_document(app, files):
    base = files.put("A사_사양회신서.xlsx", reply_form(), XLSX, confidential=True, project_id="prj_01JTESTTESTTESTTESTTESTTEST")["id"]
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "xlsx", "base_file_id": base, "document": SPEC_DOC, "source_ref": "spec:sp_x"})
        assert r.status_code == 201, r.text
        body = r.json()
        assert body["file"]["name"] == "A사_사양회신서_작성.xlsx" and body["file"]["mime"] == XLSX
        ff = body["form_fill"]
        assert ff["base_file_id"] == base and ff["base_name"] == "A사_사양회신서.xlsx" and ff["filled"] == 5
        assert ff["sheets"][0]["columns"][0]["column"] == "D" and [x["label"] for x in ff["unmatched_rows"]] == ["무게"]
        assert any("무게" in w for w in body["warnings"])
        saved = files.saved[-1]
        assert saved["confidential"] is True and saved["project_id"] == "prj_01JTESTTESTTESTTESTTESTTEST"  # 원본 양식에서 이어받음
        assert saved["meta"]["base_file_id"] == base and saved["meta"]["source_ref"] == "spec:sp_x"
        s = _open(files.by_id(body["file"]["id"]))["사양 회신서"]
        assert s["D5"].value == 500 and s["C5"].value == "500 이상"
        rec = (await c.get(f"/v1/exports/{body['export_id']}")).json()
        assert rec["status"] == "done" and rec["mode"] == "form" and rec["form_fill"]["filled"] == 5 and rec["confidential"] is True


async def test_api_fills_only_alias_and_both_languages(app, files):
    wb = Workbook()
    wb.active.append(["업체명", None])
    wb.active.append(["밝기 / Brightness", None])
    base = files.put("form.xlsx", _bytes(wb), XLSX)["id"]
    fills = [{"row": "업체명", "value": {"ko": "삼성전자", "en": "Samsung Electronics"}}, {"row": "Brightness", "value": "[확인 필요]"}]
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "xlsx", "customer_template_file_id": base, "fills": fills,
                                               "language": "both", "filename": "회신"})
        assert r.status_code == 201, r.text
        names = [f["name"] for f in r.json()["files"]]
        assert names == ["회신_KO.xlsx", "회신_EN.xlsx"]
        ko, en = (_open(files.by_id(f["id"])).active for f in r.json()["files"])
        assert (ko["B1"].value, ko["B2"].value) == ("삼성전자", "[확인 필요]")
        assert (en["B1"].value, en["B2"].value) == ("Samsung Electronics", "[TBD]")
        assert r.json()["form_fill"]["filled"] == 2 and saved_conf(files) is False


def saved_conf(files) -> bool:
    return files.saved[-1]["confidential"]


async def test_api_form_fill_errors(app, files, no_soffice):
    png = files.put("logo.png", b"\x89PNG\r\n\x1a\n" + b"0" * 20, "image/png")["id"]
    xls = files.put("old.xls", b"\xd0\xcf\x11\xe0" + b"0" * 100, "application/vnd.ms-excel")["id"]
    broken = files.put("broken.xlsx", b"PK\x03\x04garbage", XLSX)["id"]
    good = files.put("form.xlsx", reply_form(), XLSX)["id"]
    doc = SPEC_DOC
    async with testing.api_client(app) as c:
        async def post(body):
            return await c.post("/v1/exports", json={"format": "xlsx", **body})

        r = await post({"base_file_id": "file_01JNOPENOPENOPENOPENOPENOP", "document": doc})
        assert r.status_code == 404 and r.json()["error"]["code"] == "FILE_NOT_FOUND"
        r = await post({"base_file_id": png, "document": doc})
        assert r.status_code == 400 and r.json()["error"]["code"] == "INVALID_BASE_FILE"
        r = await post({"fills": [{"cell": "A1", "value": 1}]})
        assert r.status_code == 400 and "base_file_id" in r.json()["error"]["message"]
        r = await c.post("/v1/exports", json={"format": "pptx", "base_file_id": good, "document": _deck_doc()})
        assert r.status_code == 400 and "xlsx" in r.json()["error"]["message"]
        r = await post({"base_file_id": good, "fills": [{"cell": "ZZZZ1", "value": 1}]})
        assert r.status_code == 400 and r.json()["error"]["code"] == "VALIDATION_FAILED"
        r = await post({"base_file_id": good, "document": doc, "form": {"orientation": "diagonal"}})
        assert r.status_code == 422
        r = await post({"base_file_id": good})
        assert r.status_code == 400 and r.json()["error"]["code"] == "VALIDATION_FAILED"
        r = await post({"base_file_id": xls, "document": doc})
        assert r.status_code == 501 and r.json()["error"]["code"] == "PDF_CONVERTER_UNAVAILABLE"
        assert "xls" in r.json()["error"]["message"] and r.json()["error"]["details"]["reason"] == "off"
        r = await post({"base_file_id": broken, "document": doc})
        assert r.status_code == 400 and r.json()["error"]["code"] == "INVALID_BASE_FILE"
        failed = [i for i in (await c.get("/v1/exports")).json()["items"] if i["status"] == "failed"]
        assert failed and failed[0]["error"]["code"] == "INVALID_BASE_FILE"


def _deck_doc() -> dict:
    return {"slides": [{"template_code": "MS-A", "slots": {"title": "t"}}]}


async def test_api_form_fill_async_job_and_xlsm(app, files):
    from winmate_common.jobs import jobs
    from winmate_export.worker import HANDLERS

    base = files.put("견적양식.xlsm", _xlsm(), XLSM)["id"]
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "xlsx", "base_file_id": base, "async": True,
                                               "fills": [{"row": "업체명", "value": "삼성"}]})
        assert r.status_code == 202, r.text
        acc = r.json()
        assert await testing.drain_jobs("export", HANDLERS) == 1
        rec = (await c.get(f"/v1/exports/{acc['export_id']}")).json()
        assert rec["status"] == "done" and rec["file"]["name"] == "견적양식_작성.xlsm" and rec["file"]["mime"] == XLSM
        assert rec["form_fill"]["filled"] == 1 and rec["form_fill"]["cells"][0]["cell"] == "B1"
        job = await jobs().get(acc["job_id"])
        assert job.status == "succeeded" and job.result["form_fill"]["filled"] == 1 and "cells" not in job.result["form_fill"]
        with zipfile.ZipFile(io.BytesIO(files.by_id(rec["file"]["id"]))) as z:
            assert "xl/vbaProject.bin" in z.namelist()


@pytest.mark.skipif(not os.environ.get("SOFFICE_PATH"), reason="LibreOffice(SOFFICE_PATH) 없음")
async def test_api_form_fill_old_xls_via_libreoffice(app, files):
    from winmate_export.render import soffice
    from winmate_export.worker import HANDLERS

    xls = soffice.convert(reply_form(), "xlsx", "xls")
    base = files.put("old_form.xls", xls, "application/vnd.ms-excel")["id"]
    async with testing.api_client(app) as c:
        r = await c.post("/v1/exports", json={"format": "xlsx", "base_file_id": base, "document": SPEC_DOC})
        assert r.status_code == 202, r.text                                   # LibreOffice 가 필요한 일은 잡으로
        assert await testing.drain_jobs("export", HANDLERS) == 1
        rec = (await c.get(f"/v1/exports/{r.json()['export_id']}")).json()
        assert rec["status"] == "done", rec.get("error")
        assert rec["form_fill"]["converted_from"] == "xls" and rec["form_fill"]["filled"] >= 4
        assert rec["file"]["name"].endswith(".xlsx") and any(".xls" in w for w in rec["warnings"])
