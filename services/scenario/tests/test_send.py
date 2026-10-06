"""SC5 보내기 · 내보내기(AC49–53) — 시트 구성 · 제안서 묶음(handoff · ProposalHandoff) · 사용 등록 · DOCX · ZIP · workspace 색인."""
from __future__ import annotations

import io
from typing import Any

from sc_flow import QM55C, board, by_no, drain, generate, get, pick, scenes
from sc_world import GANGNAM
from winmate_common import testing


async def plan_of(client: Any, sid: str) -> dict[str, Any]:
    r = await client.get(f"/v1/scenarios/{sid}/sheet-plan")
    assert r.status_code == 200, r.text
    return r.json()


async def ws_item(world: Any, item_id: str) -> dict[str, Any]:
    async with testing.api_client(world.apps["workspace"]) as ws:
        r = await ws.get(f"/v1/items/{item_id}")
        assert r.status_code == 200, r.text
        return r.json()


async def board_with_image(client: Any) -> dict[str, Any]:
    """AC49 Given — 장면 4 · 솔루션 2 · 제품 2 · 이미지 1(장면 2) · [00] 1(장면 3)."""
    sc = await board(client)
    s2 = by_no(await scenes(client, sc["id"]), 2)
    r = await client.post(f"/v1/scenes/{s2['id']}/image:attach", json={"image_ref": "kb:image:img_counter", "source": "picked"})
    assert r.status_code == 200, r.text
    return sc


async def test_ac49_sheet_plan_with(client: Any, world: Any) -> None:
    sc = await board_with_image(client)
    plan = await plan_of(client, sc["id"])
    assert [s["code"] for s in plan["sheets"]] == ["VM-A", "SS-A", "SS-B"]
    vm, ssa, ssb = plan["sheets"]
    assert vm["kind"] == "map" and vm["detail"]["rows"] == ["매장 카운터", "본사 운영실"] and len(vm["detail"]["cols"]) == 2
    assert vm["detail"]["grid"] == "공간 2 × 솔루션 2 격자"
    assert ssa["title"] == "매장 카운터" and ssa["scene_nos"] == [1, 2, 4] and ssa["images"] == {"have": 1, "total": 3}
    assert ssb["title"] == "본사 운영실" and ssb["scene_nos"] == [3] and ssb["confirm_count"] == 1
    assert [s["n"] for s in plan["sheets"]] == [1, 2, 3]
    assert plan["carry"] == {"scenes": 4, "solutions": 2, "products": 2, "images": 1, "confirm": 1}
    assert plan["images_missing"] == 3 and plan["first_missing_scene_id"] == by_no(await scenes(client, sc["id"]), 1)["id"]
    # 시트 구성은 저장해 두고 다시 묻지 않는다(공간 정규화 LLM 1번)
    n = len(world.ai.calls("sc.space_keys"))
    await plan_of(client, sc["id"])
    assert len(world.ai.calls("sc.space_keys")) == n


async def test_ac50_map_sheet_rules(client: Any, world: Any) -> None:
    # WITHOUT · 공간 2 · 시각 있는 시간대 → VM-D(하루 타임라인 × 공간)
    sc = await board(client, type_="without")
    plan = await plan_of(client, sc["id"])
    assert plan["sheets"][0]["code"] == "VM-D"
    assert plan["sheets"][0]["detail"]["cols"] == ["07:00", "11:30", "14:00", "21:00"]
    # 공간 1 이면 맵 시트가 없다
    s3 = by_no(await scenes(client, sc["id"]), 3)
    assert (await client.delete(f"/v1/scenes/{s3['id']}")).status_code == 204
    plan = await plan_of(client, sc["id"])
    assert all(s["kind"] == "space" for s in plan["sheets"])
    # 조감도 연결 + 존 위치 → VM-C
    order = ["bez_1", "bez_2", "bez_3", "bez_4"]
    r = await client.post("/v1/scenarios:from-birdseye", json={"birdseye_id": GANGNAM, "zone_ids": order, "axis": "방문객 동선", "order": order,
                                                               "keep_link": True, "type": "with"})
    sid = r.json()["scenario_id"]
    await drain()
    await pick(client, sid, ["magicinfo"], [QM55C])
    await client.post(f"/v1/scenarios/{sid}:route-generate")
    await generate(client, sid)
    plan = await plan_of(client, sid)
    assert plan["sheets"][0]["code"] == "VM-C"
    assert plan["spaces"][:2] == ["쇼윈도 · 전면 유리창", "후면 미디어월"]  # 존 이름이 공간이 된다


async def test_ac51_handoff_usage_marks_in_proposal(client: Any, world: Any) -> None:
    sc = await board_with_image(client)
    sid = sc["id"]
    doc = await get(client, sid)
    assert doc["status"] == "done" and doc["version"] >= 1
    v = doc["version"]
    # 제안서(N1)가 저장 버전으로 묶음을 읽는다
    r = await client.get(f"/v1/scenarios/{sid}/proposal-handoff", params={"type": "solution", "section": "spaceScenario", "version": v})
    assert r.status_code == 200, r.text
    ph = r.json()
    assert ph["source"] == {"feature": "scenario", "ref_id": sid, "version": v, "title": doc["title"], "updated_at": ph["source"]["updated_at"],
                            "route": f"/scenario/{sid}/result"}
    assert ph["target"] == {"proposal_type": "solution", "section_key": "spaceScenario"}
    assert [i["key"] for i in ph["items"]] == ["VM-A:1", "SS-A:2", "SS-B:3"]
    assert [i["template_hint"]["code"] for i in ph["items"]] == ["VM-A", "SS-A", "SS-B"]
    assert ph["items"][2]["status"] == "warn" and ph["items"][2]["status_label"] == "[확정 필요] 1건"
    assert ph["items"][1]["repeat_key"] == {"kind": "space", "ref": "매장 카운터", "label": "매장 카운터"}
    assert [f["placeholder"] for f in ph["facts"]] == ["[00]"] and ph["facts"][0]["source"] == {"kind": "scenario", "scene_no": 3}
    # 솔루션 섹션(SXS) — 솔루션마다 1항목
    r = await client.get(f"/v1/scenarios/{sid}/proposal-handoff", params={"section": "solution"})
    assert [i["key"] for i in r.json()["items"]] == ["sxs:magicinfo", "sxs:smartthings_pro"]
    # §8 묶음
    r = await client.get(f"/v1/scenarios/{sid}/handoff", params={"version": v})
    ho = r.json()
    assert ho["version"] == v and len(ho["scenes"]) == 4 and [s["code"] for s in ho["sheet_plan"]] == ["VM-A", "SS-A", "SS-B"]
    assert ho["confirm_items"] == [{"scene_no": 3, "token": "[00]", "near_text": ho["confirm_items"][0]["near_text"]}]
    assert [r_["name"] for r_ in ho["roles"]][:2] == ["점장", "손님"]
    r = await client.get(f"/v1/scenarios/{sid}/handoff", params={"version": 99})
    assert r.status_code == 404
    # 넣은 뒤 사용 등록 → SC0 「제안서에 사용 중」 · workspace 색인
    r = await client.post(f"/v1/scenarios/{sid}/usages", json={"service": "proposal", "ref": "prp_test1", "label": "A 커피 제안서", "version": v})
    assert r.status_code == 201, r.text
    rows = (await client.get("/v1/scenarios", params={"status": "done", "q": doc["title"]})).json()["items"]
    row = next(x for x in rows if x["id"] == sid)
    assert row["in_proposal"] is True and row["status_sub"] == "제안서에 사용 중"
    item = await ws_item(world, sid)
    assert item["meta"]["in_proposal"] is True and item["meta"]["status"]["sub"] == "제안서에 사용 중"
    # 빼면 다시 「제안서에 아직 안 넣음」
    r = await client.delete(f"/v1/scenarios/{sid}/usages/proposal/prp_test1")
    assert r.status_code == 204
    assert (await get(client, sid))["in_proposal"] is False


async def test_ac52_docx_export_and_zip_rule(client: Any, world: Any) -> None:
    sc = await board(client)
    sid = sc["id"]
    # 이미지 0장 → ZIP 비활성(서버도 막는다)
    r = await client.post(f"/v1/scenarios/{sid}/exports", json={"kind": "zip"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "NO_IMAGES"
    r = await client.post(f"/v1/scenarios/{sid}/exports", json={"kind": "docx"})
    assert r.status_code == 202, r.text
    acc = r.json()
    assert acc["job_id"] and acc["export_id"].startswith("sce_")
    await drain()
    r = await client.get(f"/v1/scenarios/{sid}/exports/{acc['export_id']}")
    out = r.json()
    assert out["status"] == "done" and out["file_id"] and out["url"].endswith("/content?download=1"), out
    assert out["file_name"].endswith(".docx")
    async with testing.api_client(world.apps["files"]) as fs:
        r = await fs.get(f"/v1/files/{out['file_id']}/content")
        assert r.status_code == 200
        data = r.content
    import docx  # python-docx(export 의존성)

    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    items = await scenes(client, sid)
    for s in items:
        assert f"장면 {s['no']} · {s['title']}" in text
        assert s["story"][:20] in text
    cells = ["\t".join(c.text for c in row.cells) for t in d.tables for row in t.rows]
    joined = "\n".join(cells)
    assert "시각\t07:00" in joined and "솔루션 동작\tMagicINFO · 전원 스케줄" in joined and "제품\tQM55C ×3" in joined
    assert "확정 필요" in joined  # 장면 3 [00]
    assert "역할\t한 줄 소개\t원하는 것\t불편한 점" in joined  # 페르소나 표
    assert any(line.startswith("점장\t") for line in cells)
    # 만들기 요청이 export 로 갔고(문서 모양) 기밀 표시는 고객 없음 → false
    assert "등장인물 · 페르소나" in text


async def test_ac53_workspace_register(client: Any, world: Any) -> None:
    r = await client.post("/v1/scenarios", json={"type": "with"})
    sid = r.json()["id"]
    item = await ws_item(world, sid)
    assert item["feature"] == "SC" and item["route"] == f"/scenario/{sid}/type" and item["status"] == "draft"
    r = await client.patch(f"/v1/scenarios/{sid}", json={"step": 2})
    assert r.status_code == 200, r.text
    assert (await ws_item(world, sid))["route"] == f"/scenario/{sid}/input"
    sc = await board(client)
    item = await ws_item(world, sc["id"])
    assert item["feature"] == "SC" and item["route"] == f"/scenario/{sc['id']}/result" and item["status"] == "done"
    assert item["meta"]["scenes"] == 4 and item["meta"]["type"] == "with"
    assert item["summary"] and "장면 4" in item["summary"]
