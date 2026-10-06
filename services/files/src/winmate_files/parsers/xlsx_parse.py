"""XLSX(.xlsm · .xltx 포함) — 시트마다 값(수식은 마지막 계산값), 병합 범위, 숨김 여부. 시트 = 쪽.

openpyxl read_only(메모리 적게)로 읽고, 병합 범위는 시트 XML 의 <mergeCell> 을 직접 읽는다.
행 · 열 위치는 A1 기준 그대로 둔다(앞쪽 빈 행 · 열을 지우지 않는다 — 고객 양식 칸 대응용).
"""
from __future__ import annotations

import datetime as dt
import io
import logging
import math
import posixpath
import re
import zipfile
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..docprops import xml, xlsx_sheets
from ..textutil import clean
from .common import Doc, ParseContext, ParseError
from .ooxml import props

log = logging.getLogger("winmate.files.xlsx")

MAX_ROWS = 2000
MAX_COLS = 256
BLOCK_ROWS = 200
_MERGE = re.compile(rb'<(?:\w+:)?mergeCell\s+ref="([A-Z]+\d+(?::[A-Z]+\d+)?)"')
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return str(v)
        return int(v) if v.is_integer() and abs(v) < 2**53 else v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, dt.datetime):
        if v.hour == v.minute == v.second == 0 and not v.microsecond:
            return v.date().isoformat()
        return v.isoformat(timespec="seconds")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    s = clean(str(v)) if not isinstance(v, str) else v.replace("\r\n", "\n").replace("\x00", "")
    return s


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return str(v).replace("\t", " ").replace("\n", " ")


def _sheet_paths(zf: zipfile.ZipFile) -> dict[str, str]:
    """관계 id → 시트 XML 경로."""
    try:
        rels = xml(zf.read("xl/_rels/workbook.xml.rels"))
    except KeyError:
        return {}
    out = {}
    if rels is None:
        return out
    for el in rels.iter(f"{{{_REL_NS}}}Relationship"):
        target = el.get("Target") or ""
        if target.startswith("/"):
            path = target.lstrip("/")
        else:
            path = posixpath.normpath(posixpath.join("xl", target))
        out[el.get("Id") or ""] = path
    return out


def parse_xlsx(data: bytes, ctx: ParseContext) -> dict[str, Any]:
    doc = Doc("xlsx", ctx)
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = xlsx_sheets(zf)
        paths = _sheet_paths(zf)
    except Exception as exc:  # noqa: BLE001
        raise ParseError("PARSE_FAILED", f"XLSX 를 열 수 없어요: {exc}") from exc
    p = props(data)
    doc.meta.update({k: v for k, v in p.items() if v})
    merged_by_name: dict[str, list[str]] = {}
    hidden_by_name: dict[str, bool] = {}
    for info in infos:
        hidden_by_name[info["name"]] = info["state"] in ("hidden", "veryHidden")
        path = paths.get(info.get("rid") or "")
        if path:
            try:
                raw = zf.read(path)
                merged_by_name[info["name"]] = [m.decode("ascii") for m in _MERGE.findall(raw)][:5000]
            except KeyError:
                pass
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseError("PARSE_FAILED", f"XLSX 를 열 수 없어요: {exc}") from exc
    sheets: list[dict[str, Any]] = []
    try:
        for idx, ws in enumerate(wb.worksheets, start=1):
            name = ws.title
            rows: list[list[Any]] = []
            truncated = False
            wide = False
            try:
                for r_i, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True)):
                    if r_i >= MAX_ROWS:
                        truncated = True
                        break
                    vals = list(row)
                    if len(vals) > MAX_COLS:
                        wide = wide or any(v is not None for v in vals[MAX_COLS:])
                        vals = vals[:MAX_COLS]
                    rows.append([_cell(v) for v in vals])
            except Exception as exc:  # noqa: BLE001
                doc.warn(f"sheet_error:{name}")
                log.info("sheet %s failed: %s", name, exc)
            # 뒤쪽 빈 칸 · 빈 행 정리
            for r in rows:
                while r and (r[-1] is None or r[-1] == ""):
                    r.pop()
            while rows and not rows[-1]:
                rows.pop()
            ncols = max((len(r) for r in rows), default=0)
            dims = f"A1:{get_column_letter(ncols)}{len(rows)}" if rows and ncols else None
            if truncated:
                doc.warn(f"rows_truncated:{name}:{MAX_ROWS}")
            if wide:
                doc.warn(f"cols_truncated:{name}:{MAX_COLS}")
            sheet = {"name": name, "rows": rows, "dims": dims, "merged": merged_by_name.get(name, []),
                     "hidden": hidden_by_name.get(name, False), "row_count": len(rows), "truncated": truncated}
            sheets.append(sheet)
            lines = ["\t".join(_cell_str(v) for v in r).rstrip("\t") for r in rows]
            text_lines = [ln for ln in lines if ln.strip()]
            page = doc.page(idx, text=clean("\n".join(text_lines)), title=name)
            if hidden_by_name.get(name):
                page["hidden"] = True
            page["blocks"] = [{"type": "title", "text": name, "level": 1}]
            if text_lines:
                block_text = "\n".join(text_lines[:BLOCK_ROWS])
                if len(text_lines) > BLOCK_ROWS:
                    block_text += "\n…"
                page["blocks"].append({"type": "table", "text": block_text})
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    chartsheets = [i["name"] for i in infos if i["name"] not in {s["name"] for s in sheets}]
    if chartsheets:
        doc.meta["chartsheets"] = chartsheets
    doc.sheets = sheets
    doc.title = doc.meta.get("title") or (sheets[0]["name"] if sheets else None)
    return doc.finish()

