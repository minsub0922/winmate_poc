"""PPTX(.potx · .ppsx · .pptm 포함) — 슬라이드마다 제목 · 본문 · 표 · 차트 데이터 · SmartArt 글자 · 그림 · 노트 · 레이아웃 · 구역.

그림은 자식 파일(source=derived)로 뽑고 sha256 으로 중복을 없앤다. 그룹 도형은 좌표 변환을 따라 실제 위치(bbox)를 계산한다.
"""
from __future__ import annotations

import io
import logging
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from ..docprops import slide_size
from ..textutil import clean, is_bullet
from .common import Doc, ParseContext, ParseError, norm_bbox, table_text
from .ooxml import SKIP_IMAGE_TYPES, as_main_package, image_ext, props

log = logging.getLogger("winmate.files.pptx")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
R_EMBED = f"{{{NS['r']}}}embed"
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_SKIP_PH = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}
_TITLE_PH = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_EMU_PT = 12700


def _open(data: bytes) -> Any:
    try:
        return Presentation(io.BytesIO(data))
    except ValueError:
        return Presentation(io.BytesIO(as_main_package(data, "pptx")))


class _Xf:
    """그룹 좌표계 → 슬라이드 좌표계."""

    def __init__(self, ox: float = 0, oy: float = 0, sx: float = 1, sy: float = 1):
        self.ox, self.oy, self.sx, self.sy = ox, oy, sx, sy

    def apply(self, x: float, y: float) -> tuple[float, float]:
        return self.ox + x * self.sx, self.oy + y * self.sy

    def child(self, grp: Any) -> "_Xf":
        xfrm = grp._element.find("p:grpSpPr/a:xfrm", NS)
        if xfrm is None:
            return self
        off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
        choff, chext = xfrm.find("a:chOff", NS), xfrm.find("a:chExt", NS)
        if off is None or ext is None or choff is None or chext is None:
            return self
        try:
            gx, gy = float(off.get("x")), float(off.get("y"))
            gw, gh = float(ext.get("cx")), float(ext.get("cy"))
            cx, cy = float(choff.get("x")), float(choff.get("y"))
            cw, ch = float(chext.get("cx")) or 1.0, float(chext.get("cy")) or 1.0
        except (TypeError, ValueError):
            return self
        ssx, ssy = gw / cw, gh / ch
        # 자식 좌표 c → 그룹 좌표 g = gx + (c - cx) * ssx → 바깥 변환 적용
        ox, oy = self.apply(gx - cx * ssx, gy - cy * ssy)
        return _Xf(ox, oy, self.sx * ssx, self.sy * ssy)


def _bbox(shape: Any, xf: _Xf, sw: int, sh: int) -> tuple[list[float] | None, float, float]:
    try:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
    except Exception:  # noqa: BLE001
        return None, 0.0, 0.0
    if left is None or top is None or width is None or height is None:
        return None, 0.0, 0.0
    x0, y0 = xf.apply(float(left), float(top))
    x1, y1 = xf.apply(float(left) + float(width), float(top) + float(height))
    return norm_bbox(x0, y0, x1, y1, sw, sh), y0, x0


def _para_is_list(p: Any) -> bool:
    ppr = p._p.find("a:pPr", NS)
    if ppr is not None:
        if ppr.find("a:buNone", NS) is not None:
            return False
        if ppr.find("a:buChar", NS) is not None or ppr.find("a:buAutoNum", NS) is not None:
            return True
    return False


def _text_frame(shape: Any) -> tuple[str, float | None, bool]:
    """(글, 가장 큰 글자 크기 pt, 목록 여부)."""
    paras = []
    max_size = None
    listy = False
    levels = False
    for p in shape.text_frame.paragraphs:
        t = clean("".join(r.text for r in p.runs) if p.runs else p.text).replace("\v", "\n")
        for r in p.runs:
            if r.font.size is not None:
                pt = r.font.size.pt
                max_size = pt if max_size is None else max(max_size, pt)
        if not t:
            continue
        if _para_is_list(p):
            listy = True
        if (p.level or 0) > 0:
            levels = True
        paras.append(t)
    return "\n".join(paras), max_size, listy or levels


def _blips(el: etree._Element) -> list[str]:
    out = []
    for blip in el.iter(f"{{{NS['a']}}}blip"):
        rid = blip.get(R_EMBED)
        if rid:
            out.append(rid)
    return out


def _alt(shape: Any) -> str:
    for el in shape._element.iter(f"{{{NS['p']}}}cNvPr"):
        d = el.get("descr") or el.get("title")
        if d:
            return clean(d)[:300]
        break
    return ""


def _chart_rows(chart: Any) -> tuple[str, list[list[str]]]:
    title = ""
    try:
        if chart.has_title:
            title = clean(chart.chart_title.text_frame.text)
    except Exception:  # noqa: BLE001
        pass
    rows: list[list[str]] = []
    try:
        plot = chart.plots[0]
        cats = [clean(str(c)) for c in plot.categories]
        series = list(plot.series)
        rows.append([""] + [clean(s.name or "") for s in series])
        values = [list(s.values) for s in series]
        for i, cat in enumerate(cats):
            row = [cat]
            for vals in values:
                v = vals[i] if i < len(vals) else None
                row.append("" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)))
            rows.append(row)
    except Exception:  # noqa: BLE001
        pass
    return title, rows


def _smartart_text(shape: Any, part: Any) -> str:
    rel = shape._element.find(".//dgm:relIds", NS)
    if rel is None:
        return ""
    rid = rel.get(f"{{{NS['r']}}}dm")
    try:
        dpart = part.related_part(rid)
        root = etree.fromstring(dpart.blob)
    except Exception:  # noqa: BLE001
        return ""
    texts = []
    for pt in root.iter(f"{{{NS['dgm']}}}pt"):
        if pt.get("type") not in (None, "node"):
            continue
        t = "".join(x.text or "" for x in pt.iter(f"{{{NS['a']}}}t")).strip()
        if t:
            texts.append(clean(t))
    return "\n".join(texts)


def _sections(prs: Any) -> dict[str, str]:
    """슬라이드 id → 구역 이름(PowerPoint 구역)."""
    out: dict[str, str] = {}
    root = prs.part._element
    for sec in root.iter(f"{{{NS['p14']}}}section"):
        name = sec.get("name") or ""
        for sid in sec.iter(f"{{{NS['p14']}}}sldId"):
            if sid.get("id"):
                out[sid.get("id")] = name
    return out


class _SlideWalker:
    """슬라이드 하나의 도형을 돌며 블록 · 표 · 그림을 모은다."""

    def __init__(self, doc: Doc, slide: Any, idx: int, page: dict[str, Any], sw: int, sh: int):
        self.doc, self.slide, self.idx, self.page = doc, slide, idx, page
        self.sw, self.sh = sw, sh
        self.items: list[tuple[int, int, float, int, dict[str, Any]]] = []  # (우선, 세로 띠, 왼쪽, 순서, 블록)
        self.title_text: str | None = None
        self.skipped_vector = 0

    def add(self, block: dict[str, Any], top: float, left: float, priority: int = 1) -> None:
        self.items.append((priority, round(top / max(self.sh, 1) * 40), left, len(self.items), block))

    def walk(self, shapes: Any, xf: _Xf, depth: int = 0) -> None:
        for shape in shapes:
            try:
                st = shape.shape_type
            except Exception:  # noqa: BLE001
                st = None
            if st == MSO_SHAPE_TYPE.GROUP and depth < 10:
                self.walk(shape.shapes, xf.child(shape), depth + 1)
                continue
            self.shape(shape, xf)

    def shape(self, shape: Any, xf: _Xf) -> None:
        bbox, top, left = _bbox(shape, xf, self.sw, self.sh)
        ph_type = None
        if shape.is_placeholder:
            try:
                ph_type = shape.placeholder_format.type
            except Exception:  # noqa: BLE001
                ph_type = None
        if ph_type in _SKIP_PH:
            return
        # 표
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [["" if c.is_spanned else " ".join(clean(c.text).split()) for c in r.cells] for r in shape.table.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                self.page["tables"].append(rows)
                self.add({"type": "table", "text": table_text(rows), "bbox": bbox}, top, left)
            return
        # 차트 → 표(분류 × 계열)
        if getattr(shape, "has_chart", False) and shape.has_chart:
            title, rows = _chart_rows(shape.chart)
            if rows:
                self.page["tables"].append(rows)
            text = (f"[차트] {title}\n" if title else "[차트]\n") + table_text(rows)
            self.add({"type": "table", "text": text.strip(), "bbox": bbox}, top, left)
            return
        # SmartArt
        gd = shape._element.find(".//a:graphicData", NS)
        if gd is not None and gd.get("uri") == DIAGRAM_URI:
            t = _smartart_text(shape, self.slide.part)
            if t:
                self.add({"type": "list", "text": t, "bbox": bbox}, top, left)
            return
        # 그림(그림 도형 · 그림 채우기 · 그림 자리표시)
        for rid in _blips(shape._element)[:4]:
            self.picture(shape, rid, bbox, top, left)
        # 글
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text, size, listy = _text_frame(shape)
            if not text:
                return
            if ph_type in _TITLE_PH and self.title_text is None:
                self.title_text = text.replace("\n", " ")
                self.add({"type": "title", "text": text, "bbox": bbox, "font_size": size, "level": 1}, top, left, priority=0)
                return
            kind = "body"
            if ph_type == PP_PLACEHOLDER.SUBTITLE:
                kind = "heading"
            elif listy or (ph_type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT) and text.count("\n") >= 1) \
                    or all(is_bullet(x) for x in text.split("\n")):
                kind = "list"
            block: dict[str, Any] = {"type": kind, "text": text, "bbox": bbox, "font_size": size}
            if kind == "heading":
                block["level"] = 2
            self.add(block, top, left)

    def picture(self, shape: Any, rid: str, bbox: list[float] | None, top: float, left: float) -> None:
        try:
            ipart = self.slide.part.related_part(rid)
        except Exception:  # noqa: BLE001
            return
        ct = getattr(ipart, "content_type", "") or ""
        if not ct.startswith("image/"):
            return
        if ct in SKIP_IMAGE_TYPES:
            self.skipped_vector += 1
            return
        alt = _alt(shape)
        name = f"slide{self.idx}_{(shape.name or 'image').replace('/', '_')[:40]}.{image_ext(ct, str(ipart.partname))}"
        ref = self.doc.add_child(ipart.blob, name, ct, {"from": "pptx", "alt": alt or None}, page=self.idx)
        blk: dict[str, Any] = {"type": "image", "text": alt, "bbox": bbox}
        if ref:
            blk["ref"] = ref
            self.page["image_refs"].append(ref)
            self.page["images"].append({"ref": ref, "bbox": bbox})
        self.add(blk, top, left)

    def blocks(self) -> list[dict[str, Any]]:
        out = [it[4] for it in sorted(self.items, key=lambda it: it[:4])]
        for b in out:
            if b.get("font_size") is None:
                b.pop("font_size", None)
        return out


def parse_pptx(data: bytes, ctx: ParseContext) -> dict[str, Any]:
    doc = Doc("pptx", ctx)
    try:
        prs = _open(data)
    except Exception as exc:  # noqa: BLE001
        raise ParseError("PARSE_FAILED", f"PPTX 를 열 수 없어요: {exc}") from exc
    sw = int(prs.slide_width or 9144000)
    sh = int(prs.slide_height or 6858000)
    p = props(data)
    doc.meta.update({k: v for k, v in p.items() if v})
    doc.meta["slide_size"] = slide_size(sw, sh)
    sections = _sections(prs)
    skipped_vector = 0
    slide_ids = [el.get("id") for el in prs.part._element.findall("p:sldIdLst/p:sldId", NS)]

    for idx, slide in enumerate(prs.slides, start=1):
        try:
            layout = slide.slide_layout.name
        except Exception:  # noqa: BLE001
            layout = None
        page = doc.page(idx, size=[round(sw / _EMU_PT, 2), round(sh / _EMU_PT, 2)], layout=layout or None)
        if slide._element.get("show") == "0":
            page["hidden"] = True
        sid = slide_ids[idx - 1] if idx - 1 < len(slide_ids) else None
        if sid and sid in sections:
            page["section"] = sections[sid]
        walker = _SlideWalker(doc, slide, idx, page, sw, sh)
        try:
            walker.walk(slide.shapes, _Xf())
        except Exception as exc:  # noqa: BLE001
            doc.warn(f"slide_error:{idx}")
            log.info("slide %s failed: %s", idx, exc)
        skipped_vector += walker.skipped_vector
        blocks = walker.blocks()
        title_text = walker.title_text
        if title_text is None:
            # 제목 자리표시가 없으면 위쪽에서 가장 큰 글
            cands = [b for b in blocks if b["type"] in ("body", "heading") and b.get("bbox") and b["bbox"][1] < 0.35
                     and len(b["text"]) <= 120]
            if cands:
                best = max(cands, key=lambda b: (b.get("font_size") or 0, -b["bbox"][1]))
                title_text = best["text"].replace("\n", " ")
        page["blocks"] = blocks
        page["title"] = title_text[:200] if title_text else None
        page["text"] = "\n".join(b["text"] for b in blocks if b["text"] and b["type"] != "image")
        if slide.has_notes_slide:
            try:
                tf = slide.notes_slide.notes_text_frame
                notes = clean(tf.text) if tf is not None else ""
            except Exception:  # noqa: BLE001
                notes = ""
            if notes:
                page["notes"] = notes
    if skipped_vector:
        doc.warn(f"vector_images_skipped:{skipped_vector}")
    first_title = next((p["title"] for p in doc.pages if p.get("title")), None)
    doc.title = doc.meta.get("title") or first_title
    return doc.finish()
