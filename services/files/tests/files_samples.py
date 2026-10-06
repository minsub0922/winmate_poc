"""테스트용 샘플 파일을 코드로 만든다(네트워크 · 외부 파일 없음).

PDF(reportlab, 한글 CID 글꼴 + 표 + 목차) · 스캔 PDF · 암호 PDF · PPTX · DOCX · XLSX · EML · MSG(OLE 직접 작성) ·
PNG/JPEG(EXIF 회전) · HEIC · HWPX · cp949 텍스트.
"""
from __future__ import annotations

import datetime as dt
import io
import struct
import zipfile
from email.message import EmailMessage
from email.utils import format_datetime

from PIL import Image

KST = dt.timezone(dt.timedelta(hours=9))


# ── 이미지 ─────────────────────────────────────────────────────

def png(w: int = 320, h: int = 200, color: tuple[int, int, int] = (20, 40, 160)) -> bytes:
    img = Image.new("RGB", (w, h), color)
    for x in range(0, w, 40):
        img.paste((240, 240, 240), (x, 0, x + 20, h))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def noisy_png(w: int = 64, h: int = 64, seed: int = 1) -> bytes:
    """압축이 잘 안 되는 그림(크기가 4KB 를 넘게)."""
    import random

    img = Image.frombytes("RGB", (w, h), random.Random(seed).randbytes(w * h * 3))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def jpeg_rotated(w: int = 400, h: int = 200, orientation: int = 6) -> bytes:
    """가로 400 × 세로 200 픽셀 + EXIF 회전 6(시계 90°) → 보이는 크기는 200 × 400."""
    img = Image.new("RGB", (w, h), (200, 60, 30))
    exif = Image.Exif()
    exif[0x0112] = orientation
    exif[0x010F] = "Samsung"
    exif[0x0110] = "Galaxy S25"
    exif_ifd = exif.get_ifd(0x8769)
    exif_ifd[0x9003] = "2025:11:04 10:30:00"
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif.tobytes(), quality=90)
    return buf.getvalue()


def heic(w: int = 64, h: int = 48) -> bytes:
    import pillow_heif

    pillow_heif.register_heif_opener()
    img = Image.new("RGB", (w, h), (30, 160, 60))
    buf = io.BytesIO()
    img.save(buf, format="HEIF", quality=80)
    return buf.getvalue()


# ── PDF ────────────────────────────────────────────────────────

def _korean_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    try:
        pdfmetrics.getFont("HYSMyeongJo-Medium")
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    return "HYSMyeongJo-Medium"


def pdf_rfp() -> bytes:
    """2쪽 RFP: 1쪽 제목(20pt) · 본문 · 표 · 그림, 2쪽 '2. 요구사항' 제목 · 목록. 목차(북마크) 포함."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Image as RLImage
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib.styles import ParagraphStyle

    font = _korean_font()
    h1 = ParagraphStyle("h1", fontName=font, fontSize=20, leading=26, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName=font, fontSize=15, leading=20, spaceAfter=6)
    body = ParagraphStyle("b", fontName=font, fontSize=10, leading=14)
    buf = io.BytesIO()

    def on_page(canvas, doc):  # noqa: ANN001
        canvas.bookmarkPage(f"p{doc.page}")
        canvas.addOutlineEntry("개요" if doc.page == 1 else "요구사항", f"p{doc.page}", level=0)

    doc = SimpleDocTemplate(buf, pagesize=A4, title="용산 업무시설 제안요청서", author="E 자산운용")
    table = Table([["구분", "요구사항", "비고"], ["디스플레이", "55인치 이상", "필수"], ["솔루션", "MagicINFO", "선택"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTNAME", (0, 0), (-1, -1), font)]))
    img = RLImage(io.BytesIO(png(300, 180)), width=200, height=120)
    story = [
        Paragraph("용산 업무시설 제안요청서", h1),
        Paragraph("본 사업은 AI Ready 오피스 구축을 목표로 합니다. 고객사는 E 자산운용입니다.", body),
        Spacer(1, 12), table, Spacer(1, 12), img,
        PageBreak(),
        Paragraph("2. 요구사항", h2),
        Paragraph("• 사용자를 인식하고 반응하는 오피스", body),
        Paragraph("• 에너지 절감 및 정량 데이터 확보", body),
    ]
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def pdf_scanned() -> bytes:
    """1쪽: 글자 층 없이 쪽 전체 그림(스캔), 2쪽: 글자."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawImage(ImageReader(io.BytesIO(png(600, 840, (250, 250, 250)))), 0, 0, width=A4[0], height=A4[1])
    c.showPage()
    c.setFont("Helvetica", 12)
    c.drawString(72, 760, "Second page has a text layer.")
    c.showPage()
    c.save()
    return buf.getvalue()


def pdf_encrypted() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, encrypt="secret")
    c.drawString(72, 760, "secret")
    c.showPage()
    c.save()
    return buf.getvalue()


def pdf_pages(n: int) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for i in range(n):
        c.drawString(72, 760, f"Page {i + 1}")
        c.showPage()
    c.save()
    return buf.getvalue()


# ── PPTX ───────────────────────────────────────────────────────

def pptx_deck() -> bytes:
    """3장: 표지(제목 · 부제), 요구사항(본문 · 표 · 그림 · 그룹 · 노트), 시장 규모(차트 · 같은 그림 · 숨김)."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = "김하늘"
    prs.core_properties.title = "용산 AI Ready 오피스 제안"
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "용산 업무시설 AI Ready 오피스"
    s1.placeholders[1].text = "E 자산운용 · 제안지원요청서"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "요구사항 정리"
    tb = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(2))
    tb.text_frame.text = "사용자를 인식하고 반응하는 오피스"
    p = tb.text_frame.add_paragraph()
    p.text = "에너지 절감 · 정량 데이터 확보"
    p.runs[0].font.size = Pt(18)
    table = s2.shapes.add_table(3, 3, Inches(7), Inches(2), Inches(5), Inches(1.5)).table
    for r, row in enumerate([["구분", "요구", "비고"], ["디스플레이", "55인치", "필수"], ["솔루션", "MagicINFO", "선택"]]):
        for c, v in enumerate(row):
            table.cell(r, c).text = v
    pic = io.BytesIO(png(200, 120))
    s2.shapes.add_picture(pic, Inches(1), Inches(5), Inches(2), Inches(1.2))
    grp = s2.shapes.add_group_shape()
    rect = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8), Inches(5), Inches(2), Inches(1))
    rect.text = "그룹 안 글"
    s2.notes_slide.notes_text_frame.text = "발표자 노트: 성수 오피스 대비 최초 AI Ready 강조"
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "시장 규모"
    cd = CategoryChartData()
    cd.categories = ["2023", "2024", "2025"]
    cd.add_series("시장(억원)", (120, 150, 190))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), cd)
    s3.shapes.add_picture(io.BytesIO(png(200, 120)), Inches(8), Inches(2), Inches(2), Inches(1.2))  # 같은 그림(중복)
    s3._element.set("show", "0")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def as_potx(pptx: bytes) -> bytes:
    """같은 덱을 서식 파일(.potx) 형식으로."""
    src = zipfile.ZipFile(io.BytesIO(pptx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            blob = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                blob = blob.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            dst.writestr(info, blob)
    return out.getvalue()


# ── DOCX ───────────────────────────────────────────────────────

def docx_minutes() -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches

    d = Document()
    d.core_properties.author = "김하늘"
    d.core_properties.title = "고객 미팅 회의록"
    d.add_heading("고객 미팅 회의록", 0)
    d.add_heading("1. 개요", level=1)
    d.add_paragraph("일시: 2025년 11월 4일, 장소: E 자산운용 본사")
    d.add_paragraph("대표이사: AI Ready 오피스", style="List Bullet")
    d.add_paragraph("공간컨텐츠실장: 업무환경 플랫폼", style="List Bullet")
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for r, row in enumerate([["구분", "내용", "담당"], ["디스플레이", "55인치 이상", "개발사업팀장"], ["일정", "12월 제출", ""]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    t.cell(2, 1).merge(t.cell(2, 2))
    d.add_picture(io.BytesIO(png(300, 200, (200, 50, 50))), width=Inches(2))
    p = d.add_paragraph("다음 쪽으로")
    p.runs[0].add_break(WD_BREAK.PAGE)
    d.add_heading("2. 요구사항", level=2)
    d.add_paragraph("에너지 절감 · 정량 데이터 확보")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ── XLSX ───────────────────────────────────────────────────────

def xlsx_spec() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "스펙"
    ws["A1"] = "고객사 스펙 양식"
    ws.merge_cells("A1:C1")
    ws.append(["항목", "QM55C", "QB55C"])
    ws.append(["밝기(nit)", 500, 350])
    ws.append(["무게(kg)", 17.5, 14.1])
    ws.append(["출시일", dt.date(2024, 3, 1), None])
    ws2 = wb.create_sheet("가격")
    ws2.append(["모델", "단가"])
    ws2.append(["QM55C", 1234567])
    ws2.merge_cells("A3:B4")
    ws2.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 메일 ───────────────────────────────────────────────────────

def eml(*, attach_eml: bytes | None = None) -> bytes:
    m = EmailMessage()
    m["Subject"] = "[E 자산운용] 용산 오피스 제안 관련 자료"
    m["From"] = "홍길동 <hong@e-am.co.kr>"
    m["To"] = "최민섭 <minseop@samsung.com>, 김하늘 <sky@samsung.com>"
    m["Cc"] = "개발사업팀장 <dev@e-am.co.kr>"
    m["Date"] = format_datetime(dt.datetime(2025, 11, 4, 10, 30, tzinfo=KST))
    m.set_content("안녕하세요.\n\n요청하신 에너지 사용량 자료를 첨부합니다.\n\n감사합니다.")
    m.add_alternative("<html><body><p>안녕하세요.</p><p>요청하신 <b>에너지</b> 자료를 첨부합니다.</p>"
                      "<img src='cid:logo1'></body></html>", subtype="html")
    m.get_payload()[1].add_related(png(40, 20), maintype="image", subtype="png", cid="<logo1>", disposition="inline",
                                   filename="logo.png")
    m.add_attachment(pdf_pages(1), maintype="application", subtype="pdf", filename="성수오피스_에너지사용량_2025.pdf")
    if attach_eml is not None:
        m.add_attachment(attach_eml, maintype="message", subtype="rfc822", filename="원본메일.eml")
    return bytes(m)


# ── OLE 복합 문서 쓰기(테스트 전용) · MSG ─────────────────────────

_END, _FREE, _FATSECT, _NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF


def build_cfb(streams: dict[str, bytes]) -> bytes:
    """경로 → 바이트로 v3 OLE 파일을 만든다(4096 바이트 미만은 미니 스트림)."""
    nodes: dict[str, dict] = {"": {"name": "Root Entry", "type": 5, "children": []}}
    for path in sorted(streams):
        parts = path.split("/")
        for i in range(1, len(parts)):
            sp = "/".join(parts[:i])
            if sp not in nodes:
                nodes[sp] = {"name": parts[i - 1], "type": 1, "children": []}
                nodes["/".join(parts[:i - 1])]["children"].append(sp)
        nodes[path] = {"name": parts[-1], "type": 2, "children": [], "data": streams[path]}
        nodes["/".join(parts[:-1])]["children"].append(path)
    order = [""]
    i = 0
    while i < len(order):
        order.extend(nodes[order[i]]["children"])
        i += 1
    index = {p: n for n, p in enumerate(order)}
    sectors: list[bytes] = []
    fat: list[int] = []

    def alloc(data: bytes) -> int:
        if not data:
            return _END
        start = len(sectors)
        chunks = [data[o:o + 512] for o in range(0, len(data), 512)]
        for k, ch in enumerate(chunks):
            sectors.append(ch.ljust(512, b"\x00"))
            fat.append(start + k + 1 if k < len(chunks) - 1 else _END)
        return start

    mini = bytearray()
    minifat: list[int] = []
    for p in order:
        node = nodes[p]
        if node["type"] != 2:
            continue
        data = node["data"]
        if len(data) >= 4096:
            node["start"] = alloc(data)
        elif data:
            start = len(mini) // 64
            n = (len(data) + 63) // 64
            mini += data.ljust(n * 64, b"\x00")
            minifat.extend([start + k + 1 for k in range(n - 1)] + [_END])
            node["start"] = start
        else:
            node["start"] = _END
    root_start = alloc(bytes(mini))
    minifat_bytes = b"".join(struct.pack("<I", v) for v in minifat)
    if minifat_bytes:
        minifat_bytes = minifat_bytes.ljust(((len(minifat_bytes) + 511) // 512) * 512, b"\xff")
    minifat_start = alloc(minifat_bytes) if minifat_bytes else _END
    n_minifat = len(minifat_bytes) // 512
    entries = []
    for p in order:
        node = nodes[p]
        name = node["name"].encode("utf-16-le") + b"\x00\x00"
        kids = node["children"]
        child = index[kids[0]] if kids else _NOSTREAM
        parent_kids = nodes["/".join(p.split("/")[:-1])]["children"] if p else []
        right = _NOSTREAM
        if p and p in parent_kids:
            j = parent_kids.index(p)
            right = index[parent_kids[j + 1]] if j + 1 < len(parent_kids) else _NOSTREAM
        if node["type"] == 5:
            start, size = root_start, len(mini)
        elif node["type"] == 2:
            start, size = node["start"], len(node["data"])
        else:
            start, size = 0, 0
        e = name.ljust(64, b"\x00") + struct.pack("<HBB", len(name), node["type"], 1)
        e += struct.pack("<III", _NOSTREAM, right, child) + b"\x00" * 16 + struct.pack("<I", 0) + b"\x00" * 16
        e += struct.pack("<IQ", start, size)
        entries.append(e)
    while len(entries) % 4:
        entries.append(b"\x00" * 64 + struct.pack("<HBB", 0, 0, 0) + struct.pack("<III", _NOSTREAM, _NOSTREAM, _NOSTREAM)
                       + b"\x00" * 36 + struct.pack("<IQ", 0, 0))
    dir_start = alloc(b"".join(entries))
    n_fat = 1
    while (len(sectors) + n_fat) > n_fat * 128:
        n_fat += 1
    fat_start = len(sectors)
    fat.extend([_FATSECT] * n_fat)
    fat_bytes = b"".join(struct.pack("<I", v) for v in fat).ljust(n_fat * 512, b"\xff")
    for k in range(n_fat):
        sectors.append(fat_bytes[k * 512:(k + 1) * 512])
    difat = [fat_start + k for k in range(n_fat)] + [_FREE] * (109 - n_fat)
    header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16 + struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6)
    header += b"\x00" * 6 + struct.pack("<IIIIIIIII", 0, n_fat, dir_start, 0, 4096, minifat_start, n_minifat, _END, 0)
    header += b"".join(struct.pack("<I", v) for v in difat)
    assert len(header) == 512
    return header + b"".join(sectors)


def _u16(s: str) -> bytes:
    return s.encode("utf-16-le") + b"\x00\x00"


def _filetime(t: dt.datetime) -> bytes:
    epoch = dt.datetime(1601, 1, 1, tzinfo=dt.timezone.utc)
    return struct.pack("<Q", int((t - epoch).total_seconds() * 10_000_000))


def msg() -> bytes:
    """Outlook .msg — 제목 · 보낸 사람 · 받는 사람(To 1, Cc 1) · 본문 · 첨부 2(작은 것 · 4KB 넘는 것) · 보낸 시각."""
    sent = dt.datetime(2025, 11, 5, 9, 0, tzinfo=dt.timezone.utc)
    props = b"\x00" * 32 + struct.pack("<II", (0x0039 << 16) | 0x0040, 6) + _filetime(sent)
    streams = {
        "__substg1.0_0037001F": _u16("[E 자산운용] 견적 요청 회신"),
        "__substg1.0_1000001F": _u16("견적 관련 자료 보내드립니다.\r\n\r\n대표이사 의견 포함."),
        "__substg1.0_0C1A001F": _u16("홍길동"),
        "__substg1.0_5D01001F": _u16("hong@e-am.co.kr"),
        "__substg1.0_0E04001F": _u16("최민섭"),
        "__properties_version1.0": props,
        "__recip_version1.0_#00000000/__substg1.0_3001001F": _u16("최민섭"),
        "__recip_version1.0_#00000000/__substg1.0_39FE001F": _u16("minseop@samsung.com"),
        "__recip_version1.0_#00000000/__properties_version1.0": b"\x00" * 8 + struct.pack("<II", (0x0C15 << 16) | 0x0003, 6)
        + struct.pack("<Q", 1),
        "__recip_version1.0_#00000001/__substg1.0_3001001F": _u16("김하늘"),
        "__recip_version1.0_#00000001/__substg1.0_39FE001F": _u16("sky@samsung.com"),
        "__recip_version1.0_#00000001/__properties_version1.0": b"\x00" * 8 + struct.pack("<II", (0x0C15 << 16) | 0x0003, 6)
        + struct.pack("<Q", 2),
        "__attach_version1.0_#00000000/__substg1.0_3707001F": _u16("견적.txt"),
        "__attach_version1.0_#00000000/__substg1.0_37010102": "QM55C 1대 견적\n".encode("utf-8"),
        "__attach_version1.0_#00000000/__substg1.0_370E001F": _u16("text/plain"),
        "__attach_version1.0_#00000001/__substg1.0_3707001F": _u16("현장사진.png"),
        "__attach_version1.0_#00000001/__substg1.0_37010102": noisy_png(64, 64),
        "__attach_version1.0_#00000001/__substg1.0_370E001F": _u16("image/png"),
    }
    assert len(streams["__attach_version1.0_#00000001/__substg1.0_37010102"]) >= 4096
    return build_cfb(streams)


# ── 그 밖 ──────────────────────────────────────────────────────

def hwpx() -> bytes:
    ns = "http://www.hancom.co.kr/hwpml/2011/paragraph"
    section = (f'<?xml version="1.0" encoding="UTF-8"?><hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
               f'xmlns:hp="{ns}"><hp:p><hp:run><hp:t>제안요청서(HWPX)</hp:t></hp:run></hp:p>'
               f'<hp:p><hp:run><hp:t>사업명: 용산 AI Ready 오피스</hp:t></hp:run></hp:p></hs:sec>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip")
        z.writestr("Contents/section0.xml", section)
        z.writestr("version.xml", "<v/>")
    return out.getvalue()


def text_cp949() -> bytes:
    return "고객 미팅 메모\n\n제작자 의견: 설계 단계 스펙인이 목표.\n\n- 에너지 절감\n- 정량 데이터\n".encode("cp949")
