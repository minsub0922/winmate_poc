"""requirements 테스트 도구 — 견본 파일 · 호출 기록 · ai-tools 대역 · 업로드 · 잡 처리."""
from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

from fastapi import Request
from fastapi.responses import JSONResponse
from winmate_common import testing
from winmate_common.ids import new_id

ROOT = Path(__file__).resolve().parents[3]
MOCKS = ROOT / "mocks" / "ai-tools"

DEMO_RFP_SLIDES = [
    ("용산 업무시설 재개발 AI Ready 오피스 제안 요청", "발주처: E 자산운용"),
    ("사업 개요", "용산 업무시설 재개발 사업의 오피스 공간 제안을 요청합니다."),
    ("대표이사 요구", "사용자를 인식하고 반응하는 'AI Ready' 오피스를 구현\n'최초 AI Ready' 공간으로 알릴 수 있어야 함"),
    ("공간컨텐츠실 요구", "오피스를 단순 공간이 아닌 업무환경 플랫폼으로"),
]
DEMO_MEMO = """E 자산운용 미팅 메모 (11월 4일)
참석: 대표이사, 공간컨텐츠실장, 개발사업팀장
- 대표이사: 에너지 절감 효과를 정량 데이터로 확보하고 싶다
- 공간컨텐츠실장: 예측하고 반응하는 공간(Connecting-AI)
- 공간컨텐츠실장: 건물 가치와 임대 선호도를 높이는 것
- 개발사업팀장: AI 인프라를 설계 단계에 미리 반영

내부 메모
설계 단계 스펙인이 목표
성수 오피스 대비 '최초 AI Ready'를 강조
"""
DEMO_REPLY = (
    "안녕하세요, E 자산운용 공간컨텐츠실입니다.\n\n"
    "1. 업무환경 플랫폼은 입주사 앱, 공용 공간 예약, 방문객 안내를 생각하고 있습니다.\n"
    "2. 비교용으로 성수 오피스 에너지 사용량 자료를 첨부합니다.\n"
    "3. 제안 발표에는 대표이사님과 투자심의위원 2명이 참석합니다.\n\n감사합니다."
)


def pptx_bytes(slides: list[tuple[str, str]] = DEMO_RFP_SLIDES) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for title, body in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        _ = Inches
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def pdf_bytes(text: str = "성수 오피스 에너지 사용량 2025") -> bytes:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setFont("HYGothic-Medium", 12)
    c.drawString(72, 720, text)
    c.save()
    return buf.getvalue()


class Recorder:
    """ASGI 감싸개 — 서비스로 간 요청(메서드 · 경로 · JSON 본문)을 기록한다."""

    def __init__(self, app: Any, name: str):
        self.app = app
        self.name = name
        self.calls: list[dict[str, Any]] = []

    async def __call__(self, scope: dict[str, Any], receive: Any, send: Any) -> None:
        if scope["type"] != "http":
            await self.app(scope, receive, send)
            return
        chunks: list[bytes] = []
        more = True
        while more:
            msg = await receive()
            chunks.append(msg.get("body", b""))
            more = msg.get("more_body", False)
        body = b"".join(chunks)
        parsed: Any = None
        try:
            parsed = json.loads(body) if body else None
        except ValueError:
            parsed = None
        self.calls.append({"method": scope["method"], "path": scope["path"], "json": parsed})
        sent = False

        async def replay() -> dict[str, Any]:
            nonlocal sent
            if not sent:
                sent = True
                return {"type": "http.request", "body": body, "more_body": False}
            return {"type": "http.disconnect"}

        await self.app(scope, replay, send)

    def find(self, method: str, prefix: str) -> list[dict[str, Any]]:
        return [c for c in self.calls if c["method"] == method and c["path"].startswith(prefix)]


class AiStub:
    """ai-tools 대역 — task 별 응답을 바꿀 수 있다(없으면 mocks/ai-tools/<task>.json). 모든 요청을 기록."""

    def __init__(self) -> None:
        self.overrides: dict[str, Any] = {}
        self.calls: list[dict[str, Any]] = []
        self.app = testing.stub_app("ai-tools")
        app = self.app

        @app.post("/v1/llm/chat")
        async def chat(request: Request) -> JSONResponse:
            body = await request.json()
            self.calls.append(body)
            spec = self._spec(body["task"], body)
            if "error" in spec:
                e = spec["error"]
                return JSONResponse({"error": {"code": e["code"], "message": e.get("message", e["code"]), "details": {}}},
                                    status_code=e["status"])
            return JSONResponse({"call_id": new_id("call"), "provider": "stub", "model": "stub", "content": spec.get("content", ""),
                                 "json": spec.get("json"), "tool_calls": [], "finish_reason": "stop",
                                 "usage": {"input_tokens": 1, "output_tokens": 1}, "latency_ms": 1, "fallback": "none"})

        @app.post("/v1/i2t/analyze")
        async def i2t(request: Request) -> JSONResponse:
            body = await request.json()
            self.calls.append(body)
            return JSONResponse({"call_id": new_id("call"), "provider": "stub", "model": "stub", "content": "", "json": None,
                                 "boxes": [], "usage": {"input_tokens": 1, "output_tokens": 1}, "latency_ms": 1, "warnings": [],
                                 "fallback": "none"})

    def _spec(self, task: str, body: dict[str, Any]) -> dict[str, Any]:
        o = self.overrides.get(task)
        if callable(o):
            o = o(body)
        if o is not None:
            return o
        p = MOCKS / f"{task}.json"
        if p.is_file():
            data = json.loads(p.read_text(encoding="utf-8"))
            return data["responses"][0] if "responses" in data else data
        return {"json": {}}

    def tasks(self) -> list[str]:
        return [c.get("task") for c in self.calls]


async def upload(name: str, data: bytes, mime: str) -> str:
    """files 서비스(in-process)에 올리고 file_id."""
    from winmate_common.client import ServiceClient

    files = {"file": (name, data, mime)}
    meta = await ServiceClient("files").request("POST", "/v1/files", files=files, data={"confidential": "true"})
    return meta["id"]


async def drain() -> int:
    from winmate_requirements.worker import HANDLERS

    return await testing.drain_jobs("requirements", HANDLERS)


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
